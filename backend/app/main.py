"""OpenAtlas HTTP API — single FastAPI app, all routes defined inline for MVP."""
from __future__ import annotations

import io
import asyncio
import hashlib
import json
import os
import re
import subprocess
import sys
import uuid as _uuid
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
import mimetypes

from fastapi import Depends, FastAPI, File, Header, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select, or_
from sqlalchemy.orm import Session

from app.core.config import CORS_ORIGINS, HERMES_HOME, OPENATLAS_HOME, OPENATLAS_HERMES_AGENT_ROOT
from app.core.encryption import decrypt, encrypt
from app.core.security import decode_token, hash_password, issue_token, verify_password
from app.db.models import (
    AuditLog,
    CanvasEvent,
    CollaborationTemplate,
    ContextInjection,
    DigitalEmployee,
    EmployeeStatus,
    FileAsset,
    HermesRuntime,
    Job,
    JobStatus,
    MemoryBinding,
    MemoryEntry,
    MessageRecord,
    OrganizationUnit,
    RolePermissionOverride,
    RuntimeStatus,
    RuntimeType,
    Scope,
    SessionRecord,
    SessionRun,
    SkillBinding,
    SkillPackage,
    SkillRun,
    Tenant,
    TenantStatus,
    TaskArtifact,
    User,
    UserRole,
    WorkflowCheckpoint,
    WorkflowNodeRun,
    WorkflowRun,
    WorkflowRunFork,
    WorkflowStepEvent,
)
from app.db.session import SessionLocal, get_db, init_db
from app.services import hermes_client
from app.services.memory_resolver import resolve_effective_memories, build_context_block


# ── App ─────────────────────────────────────────────────────────────────────
app = FastAPI(title="OpenAtlas Backend", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS + ["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


_PENDING_HERMES_RUNS: dict[str, dict[str, Any]] = {}
_RUN_RECONCILE_TASKS: dict[str, asyncio.Task] = {}
_JOB_SCHEDULER_TASK: asyncio.Task | None = None
_RUNNING_JOB_IDS: set[str] = set()


def _remember_hermes_run(
    run_id: str,
    *,
    tenant_id: str,
    user_id: str,
    session_id: str,
    employee_id: str,
    hermes_session_id: str,
) -> None:
    if not run_id:
        return
    now = datetime.now(timezone.utc)
    expired_before = now - timedelta(hours=2)
    for old_run_id, meta in list(_PENDING_HERMES_RUNS.items()):
        if meta.get("created_at", now) < expired_before:
            _PENDING_HERMES_RUNS.pop(old_run_id, None)
    _PENDING_HERMES_RUNS[run_id] = {
        "tenant_id": tenant_id,
        "user_id": user_id,
        "session_id": session_id,
        "employee_id": employee_id,
        "hermes_session_id": hermes_session_id,
        "created_at": now,
    }


def _schedule_detached_run_reconcile(run_id: str) -> None:
    """Start a lightweight watcher that imports late Hermes output/artifacts.

    The UI stream may detach for long-running tasks, but Hermes can still finish
    and write files. This watcher keeps the OpenAtlas session/artifact tables in
    sync without requiring the user to reopen history manually.
    """
    if not run_id or run_id in _RUN_RECONCILE_TASKS:
        return
    try:
        task = asyncio.create_task(_watch_detached_run(run_id))
    except RuntimeError:
        return
    _RUN_RECONCILE_TASKS[run_id] = task
    task.add_done_callback(lambda _t: _RUN_RECONCILE_TASKS.pop(run_id, None))


async def _watch_detached_run(run_id: str) -> None:
    meta = _PENDING_HERMES_RUNS.get(run_id)
    if not meta:
        return
    # Keep this bounded. A front-end detach is not a license to poll forever.
    attempts = int(os.environ.get("OPENATLAS_RUN_RECONCILE_ATTEMPTS", "60"))
    interval = float(os.environ.get("OPENATLAS_RUN_RECONCILE_INTERVAL_SECONDS", "10"))
    for _ in range(max(1, attempts)):
        await asyncio.sleep(max(2.0, interval))
        with SessionLocal() as db:
            rec = db.get(SessionRecord, meta.get("session_id"))
            if not rec or rec.tenant_id != meta.get("tenant_id") or rec.archived:
                return
            target = await hermes_client.resolve_target(db, rec.tenant_id)
            emp = db.get(DigitalEmployee, meta.get("employee_id")) if meta.get("employee_id") else None
            imported = await _reconcile_hermes_session_transcript(
                db,
                target=target,
                rec=rec,
                tenant_id=rec.tenant_id,
                user_id=rec.user_id,
                employee_id=meta.get("employee_id") or rec.employee_id,
                speaker_name=emp.display_name if emp else "",
            )
            if imported:
                db.add(CanvasEvent(
                    tenant_id=rec.tenant_id,
                    user_id=rec.user_id,
                    session_id=rec.id,
                    employee_id=meta.get("employee_id") or rec.employee_id,
                    event_type="runtime.reconciled",
                    node_id=f"employee-{meta.get('employee_id') or rec.employee_id or ''}",
                    payload=json.dumps({"hermes_run_id": run_id, "imported": imported}, ensure_ascii=False),
                ))
                db.commit()
                return


def _get_authorized_hermes_run(run_id: str, p: Principal) -> dict[str, Any]:
    meta = _PENDING_HERMES_RUNS.get(run_id)
    if not meta:
        raise HTTPException(404, "hermes run not found or expired")
    if meta.get("tenant_id") != p.tenant.id or meta.get("user_id") != p.user.id:
        raise HTTPException(404, "hermes run not found")
    return meta


@app.on_event("startup")
def _startup() -> None:
    from app.db.session import SessionLocal
    from app.core.encryption import migrate_plaintext_keys
    from app.services.supervisor import start_supervisor
    init_db()
    # Phase 3.1 — one-shot encryption migration of legacy plaintext keys
    with SessionLocal() as db:
        n = migrate_plaintext_keys(db)
        if n:
            import logging
            logging.getLogger("openatlas.startup").info(
                "encrypted %d legacy plaintext api_key rows", n
            )
    # Phase 3.3 — start the gateway supervisor background task
    start_supervisor()


@app.on_event("startup")
async def _startup_job_scheduler() -> None:
    _start_job_scheduler()


# ── Auth dependency ─────────────────────────────────────────────────────────
class Principal:
    def __init__(self, user: User, tenant: Tenant):
        self.user = user
        self.tenant = tenant


def get_principal(
    authorization: str | None = Header(default=None),
    db: Session = Depends(get_db),
) -> Principal:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(401, "missing bearer token")
    token = authorization.split(" ", 1)[1]
    payload = decode_token(token)
    if not payload:
        raise HTTPException(401, "invalid or expired token")
    user = db.get(User, payload["sub"])
    if not user or not user.is_active:
        raise HTTPException(401, "user not active")
    tenant = db.get(Tenant, payload["tenant_id"])
    if not tenant:
        raise HTTPException(401, "tenant missing")
    return Principal(user=user, tenant=tenant)


# ── Phase 2/3: Multi-tenant admin deps ───────────────────────────────────
def require_system_admin(p: Principal = Depends(get_principal)) -> Principal:
    if p.user.role.value != "system_admin":
        raise HTTPException(403, "system_admin role required")
    return p


def require_tenant_or_system_admin(
    tid: str, p: Principal = Depends(get_principal),
) -> Principal:
    """Allow either system_admin (any tenant) OR tenant_admin whose
    tenant_id matches the path. Used for per-tenant runtime lifecycle.
    """
    if p.user.role.value == "system_admin":
        from app.db.models import Tenant
        target = db_query_tenant(tid)
        if target:
            p.tenant = target
        return p
    if p.user.role.value == "tenant_admin" and p.tenant.id == tid:
        return p
    raise HTTPException(403, "tenant_admin can only access own tenant")


def require_identity_admin(p: Principal = Depends(get_principal)) -> Principal:
    if p.user.role.value not in {"system_admin", "tenant_admin"}:
        raise HTTPException(403, "admin role required")
    return p


def db_query_tenant(tid: str):
    from app.db.models import Tenant
    from app.db.session import SessionLocal
    with SessionLocal() as db:
        return db.get(Tenant, tid)


def audit(
    db: Session, *, principal: Principal | None, action: str, resource_type: str,
    resource_id: str, request: Request | None = None, extra: dict | None = None,
) -> None:
    db.add(
        AuditLog(
            tenant_id=principal.tenant.id if principal else None,
            user_id=principal.user.id if principal else None,
            action=action,
            resource_type=resource_type,
            resource_id=resource_id,
            ip=(request.client.host if request and request.client else ""),
            user_agent=(request.headers.get("user-agent", "") if request else ""),
            extra=json.dumps(extra or {}),
        )
    )


# ── Pydantic schemas ────────────────────────────────────────────────────────
class LoginIn(BaseModel):
    email: str
    password: str


class LoginOut(BaseModel):
    access_token: str
    user: dict
    tenant: dict


class EmployeeIn(BaseModel):
    display_name: str
    description: str = ""
    avatar: str = "A"
    model: str = "hermes-agent"
    provider: str = "hermes"
    temperature: float = 0.7
    max_tokens: int = 2048
    system_prompt: str = ""
    toolsets: list[str] = Field(default_factory=list)
    initial_skill_ids: list[str] = Field(default_factory=list)


class EmployeePatch(BaseModel):
    display_name: str | None = None
    description: str | None = None
    avatar: str | None = None
    model: str | None = None
    provider: str | None = None
    temperature: float | None = None
    max_tokens: int | None = None
    system_prompt: str | None = None
    toolsets: list[str] | None = None


class SessionCreateIn(BaseModel):
    employee_id: str | None = None
    title: str = ""
    # Bug 3/4/11 (2026-06-06): 群聊接力员工 UUID 列表.
    # 之前 createGroupConversation 在 api.ts:525-526 只调 createSession 不传 relays,
    # 导致群聊变单聊.  现在后端落库 SessionRecord.employee_id = 主员工, 接力员
    # 工 ID 存到 SessionRecord.participant_ids (json list), chat stream 串行 dispatch.
    participant_ids: list[str] | None = None


class ChatIn(BaseModel):
    # P3.12 (2026-06-07) Bug 7: 附件 id 列表, 用于 <file_context> 注入
    attachment_ids: list[str] = Field(default_factory=list)
    # 保留 message 字段 (str 强类型, 拒绝 number shim, Bug 群聊 3.4.2)
    message: str = Field(..., min_length=1, max_length=32000)
    # P3.12: 群聊 relay 员工列表 (P3.11 已加透传, 这里冗余声明以便 /chat/stream 直接收)
    relay_employee_ids: list[str] = Field(default_factory=list)
    primary_employee_id: str | None = None  # 群聊时主员工 id, 跟 session.employee_id 一致时可省
    reasoning_effort: str | None = Field(
        default=None,
        pattern="^(none|minimal|low|medium|high|xhigh)$",
    )


class RunApprovalIn(BaseModel):
    choice: str = Field(default="deny", pattern="^(once|session|always|deny|approve|allow|approved)$")
    resolve_all: bool = False
    approval_id: str | None = None


class RunStopIn(BaseModel):
    reason: str = Field(default="user_requested", max_length=256)


class WorkflowNodeActionIn(BaseModel):
    action: str = Field(default="continue", pattern="^(retry|continue)$")
    message: str = Field(default="", max_length=4000)


class WorkflowCheckpointResumeIn(BaseModel):
    mode: str = Field(default="fork_resume", pattern="^(fork_resume|prompt_only)$")
    message: str = Field(default="", max_length=4000)


class WorkflowStepActionIn(BaseModel):
    action: str = Field(default="retry", pattern="^(retry|skip)$")
    message: str = Field(default="", max_length=4000)


class ArtifactPatchIn(BaseModel):
    name: str | None = Field(default=None, max_length=255)
    status: str | None = Field(default=None, max_length=24)


class SessionForkIn(BaseModel):
    title: str = Field(default="", max_length=255)


class SessionPatchIn(BaseModel):
    title: str | None = Field(default=None, max_length=255)
    pinned: bool | None = None
    workspace: str | None = Field(default=None, max_length=128)
    model_override: str | None = Field(default=None, max_length=128)


class CanvasStateIn(BaseModel):
    nodes: list[dict[str, Any]] = Field(default_factory=list)
    edges: list[dict[str, Any]] = Field(default_factory=list)
    version: int = 1


class TemplateSaveIn(BaseModel):
    name: str = Field(min_length=1, max_length=128)
    description: str = ""
    category: str = Field(default="general", max_length=64)
    visibility: str = Field(default="private", max_length=16)
    strategy: str = Field(default="relay", max_length=32)
    failure_strategy: str = Field(default="continue_with_next_employee", max_length=64)
    default_prompt: str = ""
    output_type: str = Field(default="markdown", max_length=64)


class CanvasEventIn(BaseModel):
    employee_id: str | None = None
    event_type: str = Field(default="canvas.event", max_length=64)
    node_id: str = Field(default="", max_length=128)
    edge_id: str = Field(default="", max_length=128)
    payload: dict[str, Any] = Field(default_factory=dict)


class TemplateUseIn(BaseModel):
    title: str = ""


class TaskStatusPatchIn(BaseModel):
    task_status: str = Field(max_length=24)


class SessionResumeIn(BaseModel):
    message: str = Field(default="", max_length=32000)


class MemoryIn(BaseModel):
    scope: Scope
    title: str
    content: str
    employee_id: str | None = None
    tags: list[str] = Field(default_factory=list)
    priority: int = 50
    visibility: str = "private"
    mutable: bool = True


class MemoryPatch(BaseModel):
    title: str | None = None
    content: str | None = None
    tags: list[str] | None = None
    priority: int | None = None
    visibility: str | None = None


class SkillIn(BaseModel):
    name: str
    slug: str
    description: str = ""
    category: str = "general"
    version: str = "1.0.0"
    scope: Scope = Scope.tenant
    visibility: str = "tenant"
    mutable: bool = True


class HermesSkillInstallIn(BaseModel):
    identifier: str = Field(min_length=1, max_length=1024)
    source: str = ""
    category: str = ""
    name: str = ""
    force: bool = False


class HermesSkillLifecycleIn(BaseModel):
    name: str = Field(default="", max_length=256)
    deep: bool = False
    restore: bool = False
    remove: bool = False
    sync: bool = False
    force: bool = False
    path: str = Field(default="", max_length=1024)
    url: str = Field(default="", max_length=2048)


class EmployeeToolsetsPatchIn(BaseModel):
    toolsets: list[str] = Field(default_factory=list)


class SkillBindIn(BaseModel):
    skill_id: str
    target_type: str  # "user" | "employee"
    target_id: str
    binding_mode: str = "inherited"


# ── Helpers ─────────────────────────────────────────────────────────────────
def _slugify(name: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9_-]+", "-", name.strip().lower()).strip("-")
    if s:
        return s
    digest = hashlib.sha1(name.strip().encode("utf-8")).hexdigest()[:8] if name.strip() else "default"
    return f"employee-{digest}"


def _employee_to_dict(e: DigitalEmployee) -> dict:
    return {
        "id": e.id,
        "tenant_id": e.tenant_id,
        "display_name": e.display_name,
        "profile_name": e.profile_name,
        "description": e.description,
        "avatar": e.avatar,
        "status": e.status.value if hasattr(e.status, "value") else str(e.status),
        "model": e.model,
        "provider": e.provider,
        "temperature": e.temperature,
        "max_tokens": e.max_tokens,
        "system_prompt": e.system_prompt,
        "toolsets": json.loads(e.toolsets) if e.toolsets else [],
        "created_by": e.created_by,
        "created_at": e.created_at.isoformat() if e.created_at else None,
        "updated_at": e.updated_at.isoformat() if e.updated_at else None,
    }


def _memory_to_dict(m: MemoryEntry) -> dict:
    return {
        "id": m.id,
        "scope": m.scope.value,
        "owner_user_id": m.owner_user_id,
        "employee_id": m.employee_id,
        "title": m.title,
        "content": m.content,
        "tags": [t for t in m.tags.split(",") if t] if m.tags else [],
        "priority": m.priority,
        "visibility": m.visibility,
        "mutable": m.mutable,
        "status": m.status,
        "version": m.version,
        "created_by": m.created_by,
        "created_at": m.created_at.isoformat(),
        "updated_at": m.updated_at.isoformat() if m.updated_at else None,
    }


def _skill_to_dict(s: SkillPackage, db: Session | None = None) -> dict:
    binding_count = 0
    bound_employee_count = 0
    run_count = 0
    fail_count = 0
    last_run = None
    last_error = ""
    if db is not None:
        binding_count = db.query(SkillBinding).filter(
            SkillBinding.skill_id == s.id,
            SkillBinding.enabled == True,  # noqa: E712
        ).count()
        bound_employee_count = db.query(SkillBinding).filter(
            SkillBinding.skill_id == s.id,
            SkillBinding.target_type == "employee",
            SkillBinding.enabled == True,  # noqa: E712
        ).count()
        runs = db.query(SkillRun).filter(SkillRun.skill_id == s.id).order_by(SkillRun.created_at.desc()).limit(100).all()
        run_count = len(runs)
        fail_count = sum(1 for r in runs if r.status in ("failed", "missing_in_hermes", "read_failed", "skipped_budget"))
        last_run = runs[0].created_at if runs else None
        last_error = next((r.error for r in runs if r.error), "")
    risk_level = "low"
    text = f"{s.name} {s.description} {s.category}".lower()
    if any(k in text for k in ("delete", "write", "execute", "shell", "browser", "外部", "删除", "执行")):
        risk_level = "high"
    elif any(k in text for k in ("file", "upload", "web", "api", "文件", "联网")):
        risk_level = "medium"
    return {
        "id": s.id,
        "scope": s.scope.value,
        "name": s.name,
        "slug": s.slug,
        "description": s.description,
        "ability_description": s.description or f"{s.name} capability provided by Hermes/OpenAtlas.",
        "input_example": "输入业务目标、相关文件或上下文说明。",
        "output_example": "返回结构化结论、报告、表格或可下载交付物。",
        "suitable_employees": _suggest_skill_employee_types(s),
        "risk_level": risk_level,
        "category": s.category,
        "version": s.version,
        "visibility": s.visibility,
        "mutable": s.mutable,
        "status": s.status,
        "owner_tenant_id": s.owner_tenant_id,
        "owner_user_id": s.owner_user_id,
        "source_ref": getattr(s, "source_ref", None),
        "created_by": s.created_by,
        "created_at": s.created_at.isoformat() if s.created_at else None,
        "binding_count": binding_count,
        "bound_employee_count": bound_employee_count,
        "health": {
            "run_count": run_count,
            "failure_count": fail_count,
            "failure_rate": round(fail_count / max(1, run_count), 4),
            "last_run_at": last_run.isoformat() if last_run else None,
            "last_error": last_error[:500],
        },
        "disable_impact": {
            "binding_count": binding_count,
            "employee_count": bound_employee_count,
        },
    }


def _suggest_skill_employee_types(s: SkillPackage) -> list[str]:
    text = f"{s.name} {s.slug} {s.description} {s.category}".lower()
    mapping = [
        (("doc", "pdf", "document", "文件", "文档"), "文档处理员工"),
        (("finance", "excel", "model", "财务", "表格"), "财务分析员工"),
        (("web", "search", "research", "市场", "检索"), "市场研究员工"),
        (("legal", "contract", "compliance", "法务", "合规"), "法务合规员工"),
        (("code", "python", "shell", "代码", "执行"), "技术自动化员工"),
    ]
    out = [label for keys, label in mapping if any(k in text for k in keys)]
    return out or ["通用数智员工"]


def _dedupe_skill_rows(rows: list[SkillPackage]) -> list[SkillPackage]:
    """Keep one visible row for the same scoped slug/version/source tuple."""
    def row_ts(row: SkillPackage) -> float:
        dt = row.updated_at or row.created_at
        if not dt:
            return 0.0
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.timestamp()

    by_key: dict[tuple, SkillPackage] = {}
    for s in rows:
        key = (
            s.scope.value,
            s.owner_tenant_id or "",
            s.owner_user_id or "",
            s.slug,
            s.version,
            getattr(s, "source_ref", "") or "",
        )
        prev = by_key.get(key)
        if not prev:
            by_key[key] = s
            continue
        if prev.status != "enabled" and s.status == "enabled":
            by_key[key] = s
        elif row_ts(s) > row_ts(prev):
            by_key[key] = s
    return list(by_key.values())


# P3.12 (2026-06-07) Bug 7: 附件文本提取 utility.
# 支持 PDF / DOCX / TXT / MD / CSV / XLSX. 图片先只 metadata.
_MAX_FILE_CTX_CHARS = 4000  # 注入字符上限, 超过截断
_MAX_HERMES_SKILL_CHARS = 12000
_MAX_HERMES_SKILLS_TOTAL_CHARS = 40000


def _safe_filename(name: str) -> str:
    """防止路径穿越 — strip 路径分隔符, 限制长度."""
    base = os.path.basename(name or "file")
    base = re.sub(r"[^a-zA-Z0-9._\u4e00-\u9fa5-]+", "_", base).strip("._-")
    if not base:
        base = "file"
    return base[:200]


_FILE_TTL_DAYS = int(os.environ.get("OPENATLAS_FILE_TTL_DAYS", "30") or "30")


def _guess_mime(path: str, uploaded_mime: str | None = None) -> str:
    if uploaded_mime and uploaded_mime not in ("application/octet-stream", "binary/octet-stream"):
        return uploaded_mime
    guessed, _ = mimetypes.guess_type(path)
    return guessed or uploaded_mime or "application/octet-stream"


def _extract_text_from_file(path: str, mime: str) -> str:
    """从常见文件类型提取文本. 失败返回 ''. P3.12 暂时尽力, 不要求 OCR / 视觉模型."""
    try:
        if mime == "application/pdf" or path.lower().endswith(".pdf"):
            try:
                # pdfplumber / pypdf 优先; fallback pypdf
                try:
                    import pdfplumber  # type: ignore
                    with pdfplumber.open(path) as pdf:
                        return "\n".join((p.extract_text() or "") for p in pdf.pages)
                except ImportError:
                    pass
                try:
                    from pypdf import PdfReader  # type: ignore
                    return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
                except ImportError:
                    pass
            except Exception:
                return ""
        if mime in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",) or path.lower().endswith(".docx"):
            try:
                from docx import Document  # type: ignore
                doc = Document(path)
                lines = [p.text for p in doc.paragraphs if p.text]
                for table in doc.tables:
                    for row in table.rows:
                        lines.append(" | ".join(cell.text.strip() for cell in row.cells))
                return "\n".join(lines)
            except ImportError:
                pass
            except Exception:
                pass
            try:
                import zipfile
                import xml.etree.ElementTree as ET
                with zipfile.ZipFile(path) as zf:
                    xml = zf.read("word/document.xml")
                root = ET.fromstring(xml)
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                lines: list[str] = []
                for para in root.findall(".//w:p", ns):
                    text = "".join(t.text or "" for t in para.findall(".//w:t", ns)).strip()
                    if text:
                        lines.append(text)
                return "\n".join(lines)
            except Exception:
                return ""
        if (
            mime in ("application/msword", "application/vnd.ms-word")
            or path.lower().endswith(".doc")
        ):
            try:
                proc = subprocess.run(
                    ["textutil", "-convert", "txt", "-stdout", path],
                    check=False,
                    capture_output=True,
                    text=True,
                    timeout=15,
                )
                if proc.returncode == 0 and proc.stdout.strip():
                    return proc.stdout
            except Exception:
                return ""
        lower = path.lower()
        if (
            mime in (
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                "application/vnd.ms-powerpoint",
            )
            or lower.endswith((".pptx", ".pptm", ".ppt"))
        ):
            if lower.endswith((".pptx", ".pptm")):
                try:
                    from pptx import Presentation  # type: ignore
                    prs = Presentation(path)
                    lines: list[str] = []
                    for slide_idx, slide in enumerate(prs.slides, start=1):
                        slide_lines: list[str] = []
                        for shape in slide.shapes:
                            if getattr(shape, "has_text_frame", False):
                                text = "\n".join(
                                    p.text.strip()
                                    for p in shape.text_frame.paragraphs
                                    if p.text and p.text.strip()
                                )
                                if text:
                                    slide_lines.append(text)
                            if getattr(shape, "has_table", False):
                                for row in shape.table.rows:
                                    slide_lines.append(" | ".join(cell.text.strip() for cell in row.cells))
                        if getattr(slide, "has_notes_slide", False):
                            notes = slide.notes_slide.notes_text_frame.text.strip()
                            if notes:
                                slide_lines.append(f"Notes: {notes}")
                        if slide_lines:
                            lines.append(f"# Slide {slide_idx}\n" + "\n".join(slide_lines))
                    if lines:
                        return "\n\n".join(lines)
                except ImportError:
                    pass
                except Exception:
                    pass
                try:
                    import zipfile as _zipfile
                    import xml.etree.ElementTree as ET
                    lines: list[str] = []
                    with _zipfile.ZipFile(path) as zf:
                        slide_names = sorted(
                            n for n in zf.namelist()
                            if n.startswith("ppt/slides/slide") and n.endswith(".xml")
                        )
                        for slide_idx, name in enumerate(slide_names, start=1):
                            root = ET.fromstring(zf.read(name))
                            texts = [t.text.strip() for t in root.iter() if t.tag.endswith("}t") and t.text and t.text.strip()]
                            if texts:
                                lines.append(f"# Slide {slide_idx}\n" + "\n".join(texts))
                    return "\n\n".join(lines)
                except Exception:
                    return ""
            try:
                proc = subprocess.run(
                    ["textutil", "-convert", "txt", "-stdout", path],
                    check=False,
                    capture_output=True,
                    text=True,
                    timeout=20,
                )
                if proc.returncode == 0 and proc.stdout.strip():
                    return proc.stdout
            except Exception:
                return ""
            return ""
        if mime == "text/csv" or lower.endswith(".csv"):
            try:
                import csv
                with open(path, newline='', encoding='utf-8', errors='ignore') as f:
                    rows = list(csv.reader(f))
                return "\n".join(", ".join(r) for r in rows[:200])  # 限 200 行
            except Exception:
                return ""
        if (
            mime in (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "application/vnd.ms-excel",
            )
            or lower.endswith((".xlsx", ".xlsm", ".xls"))
        ):
            try:
                from openpyxl import load_workbook  # type: ignore
                wb = load_workbook(path, read_only=True, data_only=True)
                lines = []
                for ws in wb.worksheets:
                    lines.append(f"# Sheet: {ws.title}")
                    for row in ws.iter_rows(max_row=200, values_only=True):
                        lines.append(" | ".join(str(c) if c is not None else "" for c in row))
                return "\n".join(lines)
            except ImportError:
                pass
            except Exception:
                pass
            if lower.endswith(".xls"):
                try:
                    import pandas as pd  # type: ignore
                    sheets = pd.read_excel(path, sheet_name=None, nrows=200)
                    lines = []
                    for name, df in sheets.items():
                        lines.append(f"# Sheet: {name}")
                        lines.append(df.fillna("").to_csv(index=False))
                    return "\n".join(lines)
                except Exception:
                    return ""
            return ""
        # plain text / md / json / yaml / log 等
        try:
            with open(path, encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception:
            return ""
    except Exception:
        return ""


def _build_file_context_block(file_assets: list[FileAsset]) -> str:
    """拼 <file_context>...</file_context>. 限制总字符 _MAX_FILE_CTX_CHARS."""
    if not file_assets:
        return ""
    parts: list[str] = []
    for f in file_assets:
        if not f.extracted_text:
            snippet = f"  - (no extractable text, mime={f.mime_type}, size={f.size})"
        else:
            snippet = f.extracted_text[:1500]
        parts.append(
            f"### {f.original_name}  (mime={f.mime_type}, size={f.size}, status={f.status})\n{snippet}"
        )
    body = "\n\n".join(parts)
    if len(body) > _MAX_FILE_CTX_CHARS:
        body = body[:_MAX_FILE_CTX_CHARS] + f"\n\n... (truncated at {_MAX_FILE_CTX_CHARS} chars)"
    return f"<file_context>\n{body}\n</file_context>\n\n"


def _skill_visible_to_principal(skill: SkillPackage, p: Principal) -> bool:
    if skill.scope == Scope.global_:
        return True
    if skill.scope == Scope.tenant:
        return skill.owner_tenant_id == p.tenant.id
    if skill.scope == Scope.user:
        return skill.owner_user_id == p.user.id
    if skill.scope == Scope.employee:
        return skill.owner_tenant_id == p.tenant.id
    return False


def _hermes_skill_key(name: str) -> str:
    s = re.sub(r"[\s_]+", "-", (name or "").strip().lower())
    s = re.sub(r"[^a-z0-9:./-]+", "", s)
    s = re.sub(r"-{2,}", "-", s).strip("-")
    return s


def _tenant_hermes_home(db: Session, tenant_id: str) -> Path:
    row = (
        db.query(HermesRuntime)
        .filter(HermesRuntime.tenant_id == tenant_id)
        .order_by(HermesRuntime.created_at.desc())
        .first()
    )
    if row and row.hermes_home_path:
        return Path(row.hermes_home_path).expanduser().resolve()
    return HERMES_HOME


def _read_skill_md_meta(path: Path) -> tuple[dict, str]:
    content = path.read_text(encoding="utf-8", errors="ignore")
    parsed = _parse_skill_md(content)
    return parsed.get("meta") or {}, content


def _find_hermes_skill_md(hermes_home: Path, skill: SkillPackage) -> tuple[Path | None, str]:
    skills_root = (hermes_home / "skills").resolve()
    if not skills_root.exists():
        return None, f"Hermes skills directory not found: {skills_root}"
    wanted = {
        _hermes_skill_key(skill.name),
        _hermes_skill_key(skill.slug),
        _hermes_skill_key((skill.source_ref or "").removeprefix("hermes:")),
    }
    wanted.discard("")
    for md in skills_root.rglob("SKILL.md"):
        try:
            if not md.resolve().is_relative_to(skills_root):
                continue
            meta, _ = _read_skill_md_meta(md)
        except Exception:
            continue
        meta_name = str(meta.get("name") or md.parent.name)
        category_name = f"{md.parent.parent.name}:{meta_name}" if md.parent.parent != skills_root else meta_name
        candidates = {
            _hermes_skill_key(meta_name),
            _hermes_skill_key(md.parent.name),
            _hermes_skill_key(category_name),
        }
        if wanted & candidates:
            return md, ""
    return None, f"Hermes SKILL.md not found for {skill.name}"


def _build_employee_skills_block(db: Session, employee_id: str | None, hermes_home: Path) -> tuple[str, list[dict]]:
    """Load bound Skill instructions as prompt context for Gateway chat.

    Gateway session APIs do not dispatch slash commands, so OpenAtlas maps its
    employee Hermes skill bindings to installed SKILL.md instructions. OpenAtlas
    managed or forked skills without a hermes: source are injected as governed
    metadata instead of being reported as missing Hermes packages.
    """
    if not employee_id:
        return "", []
    rows = db.execute(
        select(SkillBinding, SkillPackage)
        .join(SkillPackage, SkillPackage.id == SkillBinding.skill_id)
        .where(
            SkillBinding.target_type == "employee",
            SkillBinding.target_id == employee_id,
            SkillBinding.enabled == True,  # noqa: E712
            SkillPackage.status == "enabled",
        )
        .order_by(SkillPackage.scope, SkillPackage.category, SkillPackage.name)
    ).all()
    if not rows:
        return "", []
    parts: list[str] = []
    missing: list[str] = []
    evidence: list[dict] = []
    total = 0
    for binding, skill in rows[:20]:
        source_ref = str(getattr(skill, "source_ref", "") or "")
        if not source_ref.startswith("hermes:"):
            content = (
                f"## {skill.name}\n"
                f"OpenAtlas binding_mode: {binding.binding_mode}\n"
                "Source: OpenAtlas governed skill metadata\n\n"
                f"{(skill.description or 'No detailed skill description configured.').strip()[:1200]}"
            )
            if total + len(content) > _MAX_HERMES_SKILLS_TOTAL_CHARS:
                missing.append(f"- {skill.name}: skipped because skill prompt budget is full")
                evidence.append({
                    "id": skill.id,
                    "name": skill.name,
                    "slug": skill.slug,
                    "version": skill.version,
                    "binding_id": binding.id,
                    "binding_mode": binding.binding_mode,
                    "status": "skipped_budget",
                    "source": "openatlas_metadata",
                    "summary": "skill prompt budget is full",
                })
                continue
            total += len(content)
            parts.append(content)
            evidence.append({
                "id": skill.id,
                "name": skill.name,
                "slug": skill.slug,
                "version": skill.version,
                "binding_id": binding.id,
                "binding_mode": binding.binding_mode,
                "status": "injected",
                "source": "openatlas_metadata",
                "summary": (skill.description or "OpenAtlas governed skill metadata").strip()[:360],
            })
            continue
        md_path, err = _find_hermes_skill_md(hermes_home, skill)
        if not md_path:
            missing.append(f"- {skill.name}: {err}")
            evidence.append({
                "id": skill.id,
                "name": skill.name,
                "slug": skill.slug,
                "version": skill.version,
                "binding_id": binding.id,
                "binding_mode": binding.binding_mode,
                "status": "missing_in_hermes",
                "source": "",
                "summary": err,
            })
            continue
        try:
            _, content = _read_skill_md_meta(md_path)
        except Exception as exc:
            missing.append(f"- {skill.name}: failed to read SKILL.md ({exc})")
            evidence.append({
                "id": skill.id,
                "name": skill.name,
                "slug": skill.slug,
                "version": skill.version,
                "binding_id": binding.id,
                "binding_mode": binding.binding_mode,
                "status": "read_failed",
                "source": str(md_path),
                "summary": str(exc)[:240],
            })
            continue
        content = content[:_MAX_HERMES_SKILL_CHARS]
        if total + len(content) > _MAX_HERMES_SKILLS_TOTAL_CHARS:
            missing.append(f"- {skill.name}: skipped because skill prompt budget is full")
            evidence.append({
                "id": skill.id,
                "name": skill.name,
                "slug": skill.slug,
                "version": skill.version,
                "binding_id": binding.id,
                "binding_mode": binding.binding_mode,
                "status": "skipped_budget",
                "source": str(md_path),
                "summary": "skill prompt budget is full",
            })
            continue
        total += len(content)
        evidence.append({
            "id": skill.id,
            "name": skill.name,
            "slug": skill.slug,
            "version": skill.version,
            "binding_id": binding.id,
            "binding_mode": binding.binding_mode,
            "status": "injected",
            "source": str(md_path),
            "summary": (skill.description or content[:220]).strip()[:360],
        })
        parts.append(
            f"## {skill.name}\n"
            f"OpenAtlas binding_mode: {binding.binding_mode}\n"
            f"Source: {md_path}\n\n"
            f"{content}"
        )
    if not parts and not missing:
        return "", evidence
    body = "\n\n".join(parts)
    if missing:
        body = f"{body}\n\n### Unloaded bound skills\n" + "\n".join(missing)
    return (
        "<hermes_skills>\n"
        "The following installed Hermes Skills are bound to this OpenAtlas employee. "
        "Use their SKILL.md instructions when the user's task matches them. "
        "Do not expose this block to the user.\n\n"
        f"{body}\n"
        "</hermes_skills>"
    ), evidence


def _build_employee_system_message(
    db: Session,
    *,
    tenant_id: str,
    user_id: str,
    employee_id: str | None,
    hermes_home: Path,
    file_ctx_block: str = "",
    recent_context_block: str = "",
    relay_context: str = "",
) -> tuple[str, list[dict], list[dict]]:
    """Build one isolated prompt context for a single employee turn."""
    emp = db.get(DigitalEmployee, employee_id) if employee_id else None
    eff = resolve_effective_memories(
        db, tenant_id=tenant_id, user_id=user_id, employee_id=employee_id
    )
    system_prompt_block = ""
    if emp and emp.system_prompt and emp.system_prompt.strip():
        system_prompt_block = f"<system_prompt>\n{emp.system_prompt.strip()}\n</system_prompt>"
    ctx_block = build_context_block(eff)
    if ctx_block:
        ctx_block = f"<context>\n{ctx_block}\n</context>"
    skills_block, skill_evidence = _build_employee_skills_block(db, employee_id, hermes_home)
    relay_block = ""
    if relay_context.strip():
        relay_block = (
            "<relay_context>\n"
            "Previous employees in this group chat have already responded. "
            "Use their outputs as context, do not simply repeat them, and speak as your own role.\n\n"
            f"{relay_context.strip()}\n"
            "</relay_context>"
        )
    system_message = "\n\n".join(
        part.strip()
        for part in (system_prompt_block, skills_block, file_ctx_block, ctx_block, recent_context_block, relay_block)
        if part and part.strip()
    )
    return system_message, eff, skill_evidence


def _build_recent_conversation_block(
    db: Session,
    *,
    session_id: str,
    current_user_message: str = "",
    limit: int = 10,
) -> str:
    rows = (
        db.query(MessageRecord)
        .filter_by(session_id=session_id)
        .order_by(MessageRecord.created_at.desc())
        .limit(max(1, limit))
        .all()
    )
    rows = list(reversed(rows))
    lines: list[str] = []
    for row in rows:
        content = (row.content or "").strip()
        if not content:
            continue
        role = "user" if row.role == "user" else (row.speaker_name or row.role or "assistant")
        lines.append(f"{role}: {content[:1200]}")
    if current_user_message.strip():
        lines.append(f"user 当前输入: {current_user_message.strip()[:1200]}")
    if not lines:
        return ""
    return (
        "<recent_conversation>\n"
        "This block is the authoritative short-term context for the current OpenAtlas session. "
        "For questions like 刚才/上一轮/前面/继续, answer from this current-session context first. "
        "Do not use historical session_search unless this block is insufficient.\n\n"
        + "\n\n".join(lines[-limit:]) +
        "\n</recent_conversation>"
    )


def _resolve_initial_employee_skills(
    db: Session, skill_ids: list[str], p: Principal,
) -> list[SkillPackage]:
    """Validate and de-duplicate initial employee skills.

    Creating an employee should fail loudly when the UI submits a stale,
    disabled, or invisible skill id. Silent skips make the employee look
    configured while Hermes receives no skill context.
    """
    resolved: list[SkillPackage] = []
    seen: set[str] = set()
    for sid in skill_ids:
        if not sid or sid in seen:
            continue
        seen.add(sid)
        skill = db.get(SkillPackage, sid)
        if not skill:
            raise HTTPException(400, f"skill not found: {sid}")
        if skill.status != "enabled":
            raise HTTPException(400, f"skill is not enabled: {skill.name}")
        if not _skill_visible_to_principal(skill, p):
            raise HTTPException(403, f"skill not visible to you: {skill.name}")
        resolved.append(skill)
    return resolved


async def _sync_installed_hermes_skills(db: Session, p: Principal) -> dict:
    """Mirror Gateway /v1/skills into OpenAtlas governance metadata.

    Hermes owns the real skill package under HERMES_HOME/skills. OpenAtlas keeps
    tenant visibility, enablement, and employee binding state.
    """
    target = await hermes_client.resolve_target(db, p.tenant.id)
    payload = await hermes_client.list_skills(target)
    items = payload.get("data") if isinstance(payload, dict) else []
    if not isinstance(items, list):
        items = []
    seen_refs: set[str] = set()
    created = 0
    updated = 0
    for item in items:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or "").strip()
        if not name:
            continue
        source_ref = f"hermes:{name}"
        seen_refs.add(source_ref)
        desc = str(item.get("description") or "")
        category = str(item.get("category") or "hermes")
        slug = _hermes_skill_key(name) or f"hermes-{created + updated + 1}"
        row = db.execute(
            select(SkillPackage).where(
                SkillPackage.owner_tenant_id == p.tenant.id,
                SkillPackage.source_ref == source_ref,
            )
        ).scalar_one_or_none()
        if not row:
            row = db.execute(
                select(SkillPackage).where(
                    SkillPackage.owner_tenant_id == p.tenant.id,
                    SkillPackage.name == name,
                    SkillPackage.source_ref == "",
                )
            ).scalar_one_or_none()
        if row:
            row.slug = row.slug or slug
            row.description = desc or row.description
            row.category = category or row.category
            row.source_ref = source_ref
            row.scope = Scope.tenant
            row.visibility = "tenant"
            row.mutable = False
            if row.status != "disabled":
                row.status = "enabled"
            updated += 1
        else:
            db.add(SkillPackage(
                scope=Scope.tenant,
                owner_tenant_id=p.tenant.id,
                owner_user_id=None,
                name=name,
                slug=slug,
                description=desc,
                category=category,
                version="hermes-installed",
                source_ref=source_ref,
                visibility="tenant",
                mutable=False,
                status="enabled",
                created_by=p.user.id,
            ))
            created += 1
    stale = db.execute(
        select(SkillPackage).where(
            SkillPackage.owner_tenant_id == p.tenant.id,
            SkillPackage.source_ref.like("hermes:%"),
        )
    ).scalars().all()
    disabled = 0
    for row in stale:
        if row.source_ref not in seen_refs and row.status == "enabled":
            row.status = "disabled"
            disabled += 1
    db.commit()
    return {
        "installed": len(seen_refs),
        "created": created,
        "updated": updated,
        "disabled": disabled,
        "names": sorted(ref.replace("hermes:", "", 1) for ref in seen_refs),
        "refs": sorted(seen_refs),
    }


def _find_synced_hermes_skill(db: Session, tenant_id: str, candidates: list[str]) -> SkillPackage | None:
    wanted = {
        token
        for candidate in candidates
        for token in (
            str(candidate or "").strip().lower(),
            _hermes_skill_key(str(candidate or "")),
        )
        if token
    }
    if not wanted:
        return None
    rows = db.execute(
        select(SkillPackage).where(
            SkillPackage.owner_tenant_id == tenant_id,
            SkillPackage.source_ref.like("hermes:%"),
        )
    ).scalars().all()
    for row in rows:
        fields = [
            row.name,
            row.slug,
            row.source_ref,
            (row.source_ref or "").replace("hermes:", "", 1),
        ]
        tokens = {
            token
            for value in fields
            for token in (str(value or "").strip().lower(), _hermes_skill_key(str(value or "")))
            if token
        }
        if tokens & wanted:
            return row
    return None


def _hermes_cli_env(hermes_home: Path) -> dict[str, str]:
    env = os.environ.copy()
    env["HERMES_HOME"] = str(hermes_home)
    env["OPENATLAS_HERMES_AGENT_ROOT"] = str(OPENATLAS_HERMES_AGENT_ROOT)
    existing = env.get("PYTHONPATH", "")
    root = str(OPENATLAS_HERMES_AGENT_ROOT)
    env["PYTHONPATH"] = root if not existing else f"{root}{os.pathsep}{existing}"
    env.setdefault("PYTHONUNBUFFERED", "1")
    return env


def _hermes_python_bin() -> str:
    for candidate in (
        OPENATLAS_HERMES_AGENT_ROOT / ".venv" / "bin" / "python",
        OPENATLAS_HERMES_AGENT_ROOT / ".venv" / "bin" / "python3",
    ):
        if candidate.exists() and os.access(candidate, os.X_OK):
            return str(candidate)
    return sys.executable


def _run_hermes_python_json(hermes_home: Path, code: str, args: list[str], *, timeout: int = 45) -> dict:
    proc = subprocess.run(
        [_hermes_python_bin(), "-c", code, *args],
        cwd=str(OPENATLAS_HERMES_AGENT_ROOT),
        env=_hermes_cli_env(hermes_home),
        capture_output=True,
        text=True,
        timeout=timeout,
    )
    if proc.returncode != 0:
        detail = (proc.stderr or proc.stdout or f"exit {proc.returncode}").strip()
        raise HTTPException(502, f"Hermes Skills Hub command failed: {detail[:2000]}")
    raw = (proc.stdout or "").strip()
    try:
        return json.loads(raw or "{}")
    except json.JSONDecodeError:
        raise HTTPException(502, f"Hermes Skills Hub returned non-json output: {raw[:2000]}")


def _run_hermes_skills_install(
    hermes_home: Path,
    *,
    identifier: str,
    category: str = "",
    name: str = "",
    force: bool = False,
) -> dict:
    cmd = [
        _hermes_python_bin(),
        "-c",
        (
            "import sys;"
            "from hermes_cli.main import main;"
            "sys.argv=['hermes','skills','install',*sys.argv[1:]];"
            "main()"
        ),
        identifier,
        "--yes",
    ]
    if category:
        cmd.extend(["--category", category])
    if name:
        cmd.extend(["--name", name])
    if force:
        cmd.append("--force")
    proc = subprocess.run(
        cmd,
        cwd=str(OPENATLAS_HERMES_AGENT_ROOT),
        env=_hermes_cli_env(hermes_home),
        capture_output=True,
        text=True,
        timeout=180,
    )
    output = "\n".join(part for part in (proc.stdout, proc.stderr) if part).strip()
    error_like = re.search(r"(^|\n)\s*(error|fatal|traceback):", output, re.IGNORECASE)
    missing_skill = "no skill named" in output.lower() or "not found in any source" in output.lower()
    if proc.returncode != 0 or error_like or missing_skill:
        raise HTTPException(502, f"Hermes skill install failed: {output[:3000]}")
    return {"ok": True, "output": output[-4000:]}


def _normalize_hub_install_identifier(identifier: str, source: str) -> str:
    ident = identifier.strip()
    src = source.strip().lower()
    # Hermes CLI treats a slashless value as a short display name and searches
    # by exact skill name. Some Hub sources, especially ClawHub, return
    # slashless slugs whose display name differs (e.g. name=Bazi,
    # identifier=xray-bazi). Prefixing the source keeps it on the direct fetch
    # path and preserves tenant-local install behavior.
    if ident and "/" not in ident and src in {"clawhub", "lobehub", "browse-sh"}:
        return f"{src}/{ident}"
    return ident


def _run_hermes_cli(
    hermes_home: Path,
    argv: list[str],
    *,
    timeout: int = 120,
    input_text: str | None = None,
    allow_error: bool = False,
) -> dict:
    if not argv or any(not str(part).strip() for part in argv):
        raise HTTPException(400, "invalid Hermes CLI arguments")
    env = _hermes_cli_env(hermes_home)
    for key in (
        "ALL_PROXY", "all_proxy", "HTTP_PROXY", "http_proxy",
        "HTTPS_PROXY", "https_proxy", "SOCKS_PROXY", "socks_proxy",
    ):
        env.pop(key, None)
    cmd = [
        _hermes_python_bin(),
        "-c",
        (
            "import sys;"
            "from hermes_cli.main import main;"
            "sys.argv=['hermes',*sys.argv[1:]];"
            "main()"
        ),
        *argv,
    ]
    proc = subprocess.run(
        cmd,
        cwd=str(OPENATLAS_HERMES_AGENT_ROOT),
        env=env,
        input=input_text,
        capture_output=True,
        text=True,
        timeout=timeout,
    )
    output = "\n".join(part for part in (proc.stdout, proc.stderr) if part).strip()
    if proc.returncode != 0 and not allow_error:
        raise HTTPException(502, f"Hermes CLI failed ({' '.join(argv)}): {output[:3000]}")
    return {
        "ok": proc.returncode == 0,
        "exit_code": proc.returncode,
        "argv": argv,
        "output": output[-12000:],
    }


def _require_tenant_admin(p: Principal, label: str) -> None:
    if p.user.role not in (UserRole.tenant_admin, UserRole.system_admin):
        raise HTTPException(403, f"{label} requires admin role")


def _extract_hermes_session_id(payload: dict) -> str:
    cand = payload.get("id") or payload.get("session_id")
    if not cand:
        inner = payload.get("session") or payload.get("data") or {}
        if isinstance(inner, dict):
            cand = inner.get("id") or inner.get("session_id")
    return str(cand or "")


def _toolset_risk(name: str) -> str:
    n = name.lower()
    if n in {"terminal", "file", "code_execution", "computer_use", "browser"}:
        return "high"
    if n in {"web", "image_gen", "video_gen", "vision", "memory", "skills", "delegation", "cronjob"}:
        return "medium"
    return "low"


def _normalize_toolset_items(payload: dict) -> list[dict[str, Any]]:
    raw = payload.get("data") or payload.get("items") or payload.get("toolsets") or []
    if not isinstance(raw, list):
        return []
    items: list[dict[str, Any]] = []
    for item in raw:
        if isinstance(item, str):
            name = item
            data = {"name": name}
        elif isinstance(item, dict):
            name = str(item.get("name") or item.get("id") or item.get("slug") or "")
            data = dict(item)
            data["name"] = name
        else:
            continue
        if not name:
            continue
        data["risk_level"] = data.get("risk_level") or _toolset_risk(name)
        data["requires_approval"] = data["risk_level"] == "high"
        items.append(data)
    return items


def _tail_file(path: Path, *, lines: int = 200, max_bytes: int = 200_000) -> str:
    if not path.exists() or not path.is_file():
        return ""
    with open(path, "rb") as fh:
        try:
            fh.seek(0, os.SEEK_END)
            size = fh.tell()
            fh.seek(max(0, size - max_bytes), os.SEEK_SET)
        except OSError:
            pass
        text = fh.read().decode("utf-8", errors="replace")
    return "\n".join(text.splitlines()[-lines:])


def _zip_directory(src: Path, dst: Path) -> int:
    dst.parent.mkdir(parents=True, exist_ok=True)
    total = 0
    with zipfile.ZipFile(dst, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in src.rglob("*"):
            if not path.is_file():
                continue
            if path.name in {".DS_Store"}:
                continue
            rel = path.relative_to(src)
            zf.write(path, rel.as_posix())
            try:
                total += path.stat().st_size
            except OSError:
                pass
    return total


def _file_to_dict(f: FileAsset) -> dict:
    text = f.extracted_text or ""
    compact = re.sub(r"\s+", " ", text).strip()
    summary = compact[:280] + ("..." if len(compact) > 280 else "") if compact else ""
    snippets = []
    if compact:
        window = 220
        for i in range(0, min(len(compact), window * 3), window):
            snippet = compact[i:i + window].strip()
            if snippet:
                snippets.append({
                    "index": len(snippets) + 1,
                    "text": snippet + ("..." if i + window < len(compact) else ""),
            })
    created = f.created_at
    if created and created.tzinfo is None:
        created = created.replace(tzinfo=timezone.utc)
    expires_at = created + timedelta(days=_FILE_TTL_DAYS) if created and _FILE_TTL_DAYS > 0 else None
    is_expired = bool(expires_at and datetime.now(timezone.utc) > expires_at)
    return {
        "id": f.id,
        "tenant_id": f.tenant_id,
        "user_id": f.user_id,
        "session_id": f.session_id,
        "employee_id": f.employee_id,
        "original_name": f.original_name,
        "mime_type": f.mime_type,
        "size": f.size,
        "status": f.status,
        "extracted_chars": len(f.extracted_text or ""),
        "summary": summary,
        "snippets": snippets,
        "expires_at": expires_at.isoformat() if expires_at else None,
        "is_expired": is_expired,
        "storage_path": f.storage_path,  # 仅同租户可读, 跨租户会被 403
        "created_at": f.created_at.isoformat() if f.created_at else None,
    }


def _rough_token_count(text: str) -> int:
    if not text:
        return 0
    cjk = len(re.findall(r"[\u4e00-\u9fff]", text))
    rest = max(0, len(text) - cjk)
    return max(1, int(cjk * 0.9 + rest / 4))


def _session_message_count(db: Session, sid: str, fallback: int | None = 0) -> int:
    count = db.query(MessageRecord).filter_by(session_id=sid).count()
    return count if count > 0 else int(fallback or 0)


def _session_summary(db: Session, rec: SessionRecord) -> dict:
    rows = db.query(MessageRecord).filter_by(session_id=rec.id).order_by(MessageRecord.created_at.asc()).all()
    input_chars = sum(len(r.content or "") for r in rows if r.role == "user")
    output_chars = sum(len(r.content or "") for r in rows if r.role == "assistant")
    usage_input = sum(int(getattr(r, "input_tokens", 0) or 0) for r in rows)
    usage_output = sum(int(getattr(r, "output_tokens", 0) or 0) for r in rows)
    speakers: dict[str, dict] = {}
    for r in rows:
        if r.role != "assistant":
            continue
        key = r.speaker_employee_id or r.speaker_name or "assistant"
        item = speakers.setdefault(key, {
            "employee_id": r.speaker_employee_id,
            "name": r.speaker_name or "Assistant",
            "turns": 0,
            "output_chars": 0,
        })
        item["turns"] += 1
        item["output_chars"] += len(r.content or "")
    user_turns = [r.content.strip() for r in rows if r.role == "user" and r.content.strip()]
    assistant_turns = [r.content.strip() for r in rows if r.role == "assistant" and r.content.strip()]
    brief_parts = [f"本会话累计 {len(user_turns)} 轮输入、{len(assistant_turns)} 条回复。"]
    if user_turns:
        recent_inputs = " / ".join(t[:80] for t in user_turns[-3:])
        brief_parts.append(f"主要问题: {recent_inputs}")
    if assistant_turns:
        recent_outputs = " / ".join(t[:110] for t in assistant_turns[-3:])
        brief_parts.append(f"主要结论: {recent_outputs}")
    return {
        "message_count": len(rows),
        "user_turns": len(user_turns),
        "assistant_turns": len(assistant_turns),
        "input_chars": input_chars,
        "output_chars": output_chars,
        "input_tokens": usage_input,
        "output_tokens": usage_output,
        "total_tokens": usage_input + usage_output,
        "input_tokens_est": usage_input or _rough_token_count("\n".join(r.content or "" for r in rows if r.role == "user")),
        "output_tokens_est": usage_output or _rough_token_count("\n".join(r.content or "" for r in rows if r.role == "assistant")),
        "speakers": list(speakers.values()),
        "summary": "\n".join(brief_parts) if brief_parts else "本会话还没有可总结的对话内容。",
    }


def _infer_task_status_from_assistant(text: str) -> tuple[str, str]:
    """Best-effort task closure signal from visible assistant output.

    Hermes may ask for clarification without emitting a dedicated approval/input
    event. We keep the heuristic intentionally conservative and only mark
    needs_input when the assistant explicitly asks the user for missing info.
    """
    compact = re.sub(r"\s+", " ", (text or "").strip())
    if not compact:
        return "needs_input", "assistant returned no visible content"
    ask_patterns = [
        r"请(提供|补充|确认|说明|上传|告知)",
        r"(需要|还需|缺少|无法继续|看不到).{0,24}(信息|资料|文件|上下文|参数|内容)",
        r"(which|please provide|need more|missing).{0,40}(information|context|file|details)",
    ]
    if compact.endswith(("?", "？")):
        return "needs_input", "assistant ended with a question"
    for pattern in ask_patterns:
        if re.search(pattern, compact, flags=re.IGNORECASE):
            return "needs_input", "assistant requested more input"
    return "completed", "assistant produced a final response"


APPROVAL_EVENT_NAMES = {
    "approval.request",
    "approval.requested",
    "approval.required",
    "approval_required",
    "tool.approval_required",
    "tool.requires_approval",
    "run.requires_approval",
    "run.requires_action",
    "requires_action",
}

APPROVAL_SENSITIVE_TOOLS = {
    "terminal",
    "shell",
    "bash",
    "execute_code",
    "code_execution",
    "python",
    "file",
    "file_write",
    "browser",
    "computer_use",
}


def _event_key(name: Any) -> str:
    return str(name or "").strip().lower().replace("-", "_")


def _boolish(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    if isinstance(v, str):
        return v.strip().lower() in {"1", "true", "yes", "y", "required", "pending", "waiting", "wait"}
    return bool(v)


def _extract_tool_command(data: dict[str, Any]) -> str:
    for key in ("command", "preview", "script", "code", "input", "label", "description"):
        val = data.get(key)
        if isinstance(val, str) and val.strip():
            return val.strip()
    for key in ("args", "parameters", "payload"):
        val = data.get(key)
        if isinstance(val, dict):
            nested = _extract_tool_command(val)
            if nested:
                return nested
        elif isinstance(val, str) and val.strip():
            return val.strip()
    return ""


def _approval_choices(data: dict[str, Any]) -> list[str]:
    choices = data.get("choices")
    if isinstance(choices, list):
        cleaned = [str(c) for c in choices if str(c) in {"once", "session", "always", "deny"}]
        if cleaned:
            return cleaned
    if data.get("allow_permanent") is False:
        return ["once", "session", "deny"]
    return ["once", "session", "always", "deny"]


def _normalize_approval_payload(
    ev_name: str,
    data: dict[str, Any],
    run_id: str,
    *,
    source: str = "hermes",
) -> dict[str, Any]:
    tool_name = str(data.get("tool_name") or data.get("tool") or data.get("name") or "")
    description = (
        data.get("description")
        or data.get("reason")
        or data.get("message")
        or data.get("prompt")
        or ("Hermes 工具调用需要人工确认" if source == "hermes" else "执行型工具可能正在等待人工确认")
    )
    return {
        "approval_id": data.get("approval_id") or data.get("approvalId") or data.get("id"),
        "command": _extract_tool_command(data),
        "description": str(description),
        "pattern_key": data.get("pattern_key") or data.get("risk") or ev_name,
        "pattern_keys": data.get("pattern_keys"),
        "choices": _approval_choices(data),
        "tool_name": tool_name,
        "hermes_run_id": run_id,
        "run_id": run_id,
        "source": source,
        "timeout_ms": data.get("timeout_ms") or data.get("timeoutMs") or 60000,
    }


def _is_approval_event(ev_name: str, data: dict[str, Any]) -> bool:
    key = _event_key(ev_name)
    if key in APPROVAL_EVENT_NAMES:
        return True
    if any(_boolish(data.get(k)) for k in ("approval_required", "requires_approval", "requiresApproval", "needs_approval", "needs_confirmation", "awaiting_approval")):
        return True
    status = _event_key(data.get("status") or data.get("state"))
    return status in {"waiting_approval", "pending_approval", "requires_approval", "approval_required"}


def _is_approval_sensitive_tool(tool_name: str) -> bool:
    n = _event_key(tool_name)
    return n in APPROVAL_SENSITIVE_TOOLS or any(part in n for part in ("terminal", "shell", "execute_code", "code_execution"))


def _is_stale_running_session(row: SessionRecord) -> bool:
    if (getattr(row, "task_status", "") or "") != "running" or not row.updated_at:
        return False
    updated_at = row.updated_at
    if updated_at.tzinfo is None:
        updated_at = updated_at.replace(tzinfo=timezone.utc)
    return (datetime.now(timezone.utc) - updated_at).total_seconds() > 30 * 60


def _session_health_snapshot(db: Session, rec: SessionRecord) -> dict:
    rows = db.query(MessageRecord).filter_by(session_id=rec.id).order_by(MessageRecord.created_at.asc()).all()
    artifacts = db.query(TaskArtifact).filter(
        TaskArtifact.session_id == rec.id,
        TaskArtifact.tenant_id == rec.tenant_id,
        TaskArtifact.archived == False,  # noqa: E712
    ).all()
    contexts = db.query(ContextInjection).filter(
        ContextInjection.session_id == rec.id,
        ContextInjection.tenant_id == rec.tenant_id,
    ).all()
    canvas_events = db.query(CanvasEvent).filter(
        CanvasEvent.session_id == rec.id,
        CanvasEvent.tenant_id == rec.tenant_id,
    ).order_by(CanvasEvent.created_at.desc()).limit(20).all()
    latest_user = next((r for r in reversed(rows) if r.role == "user"), None)
    latest_assistant_after_user = None
    if latest_user:
        latest_assistant_after_user = next(
            (
                r for r in reversed(rows)
                if r.role == "assistant" and r.created_at >= latest_user.created_at
            ),
            None,
        )
    pending_runs = [
        {"run_id": rid, **meta}
        for rid, meta in _PENDING_HERMES_RUNS.items()
        if meta.get("session_id") == rec.id
    ]
    issues: list[dict[str, Any]] = []
    actions: list[dict[str, Any]] = []
    status = rec.task_status or "draft"
    if status == "failed":
        issues.append({"code": "task_failed", "severity": "critical", "message": "任务已标记失败，需要重新执行或人工处理。"})
        actions.append({"key": "retry", "label": "重新进行", "kind": "task_status", "value": "running"})
    if status == "needs_input":
        issues.append({"code": "needs_input", "severity": "warning", "message": rec.task_summary or "任务需要用户补充信息。"})
        actions.append({"key": "resume", "label": "继续补充", "kind": "resume"})
    if _is_stale_running_session(rec):
        issues.append({"code": "stale_running", "severity": "critical", "message": "任务运行中但超过 30 分钟没有更新。"})
        actions.append({"key": "recover", "label": "补同步结果", "kind": "recover"})
        actions.append({"key": "inspect_runtime", "label": "检查 Runtime", "kind": "runtime"})
    if latest_user and not latest_assistant_after_user:
        issues.append({"code": "awaiting_assistant", "severity": "warning", "message": "最近一轮用户输入后还没有可见模型回复。"})
        actions.append({"key": "recover", "label": "补同步 Hermes 会话", "kind": "recover"})
    if latest_user and not contexts:
        issues.append({"code": "missing_context_trace", "severity": "warning", "message": "本会话缺少上下文注入记录，难以追溯文件、Skill、记忆来源。"})
        actions.append({"key": "retry_with_context", "label": "重新带上下文执行", "kind": "resume"})
    if rows and not artifacts:
        assistant_text = "\n\n".join(r.content or "" for r in rows if r.role == "assistant")
        if _extract_task_artifacts(assistant_text, prefix="health-probe"):
            issues.append({"code": "unregistered_artifact", "severity": "warning", "message": "回复中疑似包含交付物，但没有登记到输出物。"})
            actions.append({"key": "recover", "label": "重新抽取交付物", "kind": "recover"})
    if pending_runs:
        issues.append({"code": "pending_hermes_run", "severity": "info", "message": f"仍有 {len(pending_runs)} 个 Hermes Run 处于可追踪状态。"})
    latest_event = canvas_events[0] if canvas_events else None
    latest_run = db.query(SessionRun).filter(
        SessionRun.session_id == rec.id,
        SessionRun.tenant_id == rec.tenant_id,
    ).order_by(SessionRun.created_at.desc()).first()
    latest_workflow = db.query(WorkflowRun).filter(
        WorkflowRun.session_id == rec.id,
        WorkflowRun.tenant_id == rec.tenant_id,
    ).order_by(WorkflowRun.created_at.desc()).first()
    score = 100
    severity_penalty = {"critical": 28, "warning": 14, "info": 4}
    for issue in issues:
        score -= severity_penalty.get(issue.get("severity"), 8)
    if rows and not contexts:
        score -= 10
    if artifacts and not any((_artifact_provenance(db, a).get("context_counts") or {}) for a in artifacts[:3]):
        score -= 8
    score = max(0, min(100, score))
    health_status = "healthy"
    if any(i.get("severity") == "critical" for i in issues):
        health_status = "action_required"
    elif any(i.get("severity") == "warning" for i in issues):
        health_status = "warning"
    elif status == "running" or pending_runs:
        health_status = "running"
    return {
        "status": health_status,
        "score": score,
        "task_status": status,
        "is_stale": _is_stale_running_session(rec),
        "issues": issues,
        "recommended_actions": actions,
        "pending_runs": [{"run_id": item.get("run_id"), "created_at": item.get("created_at")} for item in pending_runs],
        "counts": {
            "messages": len(rows),
            "user_turns": sum(1 for r in rows if r.role == "user"),
            "assistant_turns": sum(1 for r in rows if r.role == "assistant"),
            "context_injections": len(contexts),
            "artifacts": len(artifacts),
            "canvas_events": db.query(CanvasEvent).filter(CanvasEvent.session_id == rec.id, CanvasEvent.tenant_id == rec.tenant_id).count(),
        },
        "latest_runtime_event": _canvas_event_to_dict(latest_event) if latest_event else None,
        "latest_run": _session_run_to_dict(latest_run) if latest_run else None,
        "latest_workflow": _workflow_run_to_dict(latest_workflow) if latest_workflow else None,
    }


def _context_to_dict(row: ContextInjection) -> dict:
    try:
        payload = json.loads(row.payload or "{}")
    except Exception:
        payload = {}
    return {
        "id": row.id,
        "session_id": row.session_id,
        "message_id": row.message_id,
        "employee_id": row.employee_id,
        "turn_index": row.turn_index,
        "kind": row.kind,
        "source_id": row.source_id,
        "name": row.name,
        "scope": row.scope,
        "status": row.status,
        "summary": row.summary,
        "payload": payload,
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }


def _artifact_provenance(db: Session, row: TaskArtifact) -> dict:
    msg = db.get(MessageRecord, row.message_id) if row.message_id else None
    run = db.get(SessionRun, row.run_id) if getattr(row, "run_id", None) else None
    extra = _json_loads_obj(getattr(row, "provenance_payload", "") or "{}", {})
    if not isinstance(extra, dict):
        extra = {}
    context_q = db.query(ContextInjection).filter(
        ContextInjection.session_id == row.session_id,
        ContextInjection.tenant_id == row.tenant_id,
    )
    if msg and msg.id:
        context_q = context_q.filter(ContextInjection.message_id == msg.id)
    context_rows = context_q.order_by(ContextInjection.created_at.desc()).limit(40).all()
    context_counts: dict[str, int] = {}
    for c in context_rows:
        context_counts[c.kind] = context_counts.get(c.kind, 0) + 1
    user_q = db.query(MessageRecord).filter(
        MessageRecord.session_id == row.session_id,
        MessageRecord.role == "user",
    )
    if row.created_at:
        user_q = user_q.filter(MessageRecord.created_at <= row.created_at)
    user_msg = user_q.order_by(MessageRecord.created_at.desc()).first()
    return {
        "source": row.source,
        "source_path": getattr(row, "source_path", "") or "",
        "run_id": getattr(row, "run_id", None),
        "hermes_run_id": run.hermes_run_id if run else "",
        "employee_id": getattr(row, "employee_id", None) or (msg.speaker_employee_id if msg else None),
        "employee_name": msg.speaker_name if msg else "",
        "message_id": row.message_id,
        "turn_index": msg.turn_index if msg else None,
        "query_excerpt": (user_msg.content or "")[:180] if user_msg else "",
        "context_counts": context_counts,
        "context_items": [_context_to_dict(c) for c in context_rows[:12]],
        **extra,
    }


def _artifact_to_dict(row: TaskArtifact, db: Session | None = None) -> dict:
    out = {
        "id": row.id,
        "session_id": row.session_id,
        "message_id": row.message_id,
        "kind": row.kind,
        "name": row.name,
        "mime_type": row.mime_type,
        "content": row.content,
        "source": row.source,
        "source_path": getattr(row, "source_path", "") or "",
        "run_id": getattr(row, "run_id", None),
        "employee_id": getattr(row, "employee_id", None),
        "version": int(getattr(row, "version", 1) or 1),
        "status": row.status,
        "archived": row.archived,
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }
    if db is not None:
        out["provenance"] = _artifact_provenance(db, row)
    return out


def _record_context_injections(
    db: Session,
    *,
    tenant_id: str,
    user_id: str,
    session_id: str,
    message_id: str | None,
    employee_id: str | None,
    turn_index: int | None,
    skills: list[dict],
    memories: list[dict],
    files: list[dict],
) -> list[ContextInjection]:
    rows: list[ContextInjection] = []
    for item in skills:
        row = ContextInjection(
            tenant_id=tenant_id,
            user_id=user_id,
            session_id=session_id,
            message_id=message_id,
            employee_id=employee_id,
            turn_index=turn_index,
            kind="skill",
            source_id=str(item.get("id") or ""),
            name=str(item.get("name") or item.get("slug") or "Skill"),
            scope=str(item.get("scope") or ""),
            status=str(item.get("status") or "injected"),
            summary=str(item.get("summary") or "")[:1200],
            payload=json.dumps(item, ensure_ascii=False),
        )
        db.add(row)
        rows.append(row)
    for item in memories:
        row = ContextInjection(
            tenant_id=tenant_id,
            user_id=user_id,
            session_id=session_id,
            message_id=message_id,
            employee_id=employee_id,
            turn_index=turn_index,
            kind="memory",
            source_id=str(item.get("id") or ""),
            name=str(item.get("title") or "Memory"),
            scope=str(item.get("scope") or ""),
            status="injected",
            summary=str(item.get("content") or "")[:1200],
            payload=json.dumps(item, ensure_ascii=False),
        )
        db.add(row)
        rows.append(row)
    for item in files:
        row = ContextInjection(
            tenant_id=tenant_id,
            user_id=user_id,
            session_id=session_id,
            message_id=message_id,
            employee_id=employee_id,
            turn_index=turn_index,
            kind="file",
            source_id=str(item.get("id") or ""),
            name=str(item.get("name") or item.get("original_name") or "File"),
            scope="session",
            status=str(item.get("status") or "injected"),
            summary=str(item.get("summary") or "")[:1200],
            payload=json.dumps(item, ensure_ascii=False),
        )
        db.add(row)
        rows.append(row)
    return rows


_FENCE_ARTIFACTS = {
    "html": ("html", "text/html;charset=utf-8", "html"),
    "htm": ("html", "text/html;charset=utf-8", "html"),
    "markdown": ("markdown", "text/markdown;charset=utf-8", "md"),
    "md": ("markdown", "text/markdown;charset=utf-8", "md"),
    "mermaid": ("mermaid", "text/plain;charset=utf-8", "mmd"),
    "json": ("json", "application/json;charset=utf-8", "json"),
    "csv": ("table", "text/csv;charset=utf-8", "csv"),
}


def _extract_task_artifacts(text: str, *, prefix: str = "交付物") -> list[dict]:
    artifacts: list[dict] = []
    seen_spans: list[tuple[int, int]] = []
    fence_re = re.compile(r"```([a-zA-Z0-9_+\-]*)[ \t]*\n?([\s\S]*?)```")
    for idx, m in enumerate(fence_re.finditer(text), start=1):
        lang = (m.group(1) or "").lower()
        if lang not in _FENCE_ARTIFACTS:
            continue
        kind, mime, ext = _FENCE_ARTIFACTS[lang]
        content = (m.group(2) or "").strip()
        if not content:
            continue
        seen_spans.append((m.start(), m.end()))
        if kind == "html":
            content = _wrap_html_artifact(content)
        artifacts.append({
            "kind": kind,
            "name": f"{prefix}-{idx}.{ext}",
            "mime_type": mime,
            "content": content,
        })
    stripped = text.strip()
    if not any(a["kind"] == "html" for a in artifacts):
        html = _extract_standalone_html(stripped)
        if html:
            artifacts.append({
                "kind": "html",
                "name": f"{prefix}-页面.html",
                "mime_type": "text/html;charset=utf-8",
                "content": _wrap_html_artifact(html),
            })
    return artifacts[:12]


def _extract_standalone_html(text: str) -> str:
    if not text:
        return ""
    if re.match(r"^\s*(<!doctype html|<html[\s>])", text, flags=re.I):
        return text
    m = re.search(r"(<(?:section|article|main|div|html|body|table|style|script)[\s\S]*?</(?:section|article|main|div|html|body|table|style|script)>)", text, flags=re.I)
    return m.group(1).strip() if m else ""


def _wrap_html_artifact(html: str) -> str:
    if re.match(r"^\s*(<!doctype|<html)", html, flags=re.I):
        return html
    return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>OpenAtlas Artifact</title>
</head>
<body>
{html}
</body>
</html>"""


def _tenant_dashboard_metrics(db: Session, tid: str, *, user_id: str | None = None) -> dict:
    msg_q = db.query(MessageRecord).join(SessionRecord, MessageRecord.session_id == SessionRecord.id).filter(
        SessionRecord.tenant_id == tid,
    )
    sess_q = db.query(SessionRecord).filter(SessionRecord.tenant_id == tid)
    file_q = db.query(FileAsset).filter(FileAsset.tenant_id == tid)
    audit_q = db.query(AuditLog).filter(AuditLog.tenant_id == tid)
    artifact_q = db.query(TaskArtifact).filter(TaskArtifact.tenant_id == tid)
    skill_run_q = db.query(SkillRun).filter(SkillRun.tenant_id == tid)
    if user_id:
        msg_q = msg_q.filter(SessionRecord.user_id == user_id)
        sess_q = sess_q.filter(SessionRecord.user_id == user_id)
        file_q = file_q.filter(FileAsset.user_id == user_id)
        audit_q = audit_q.filter(AuditLog.user_id == user_id)
        artifact_q = artifact_q.filter(TaskArtifact.user_id == user_id)
        skill_run_q = skill_run_q.filter(SkillRun.user_id == user_id)
    messages = msg_q.all()
    files = file_q.all()
    audits = audit_q.all()
    sessions = sess_q.all()
    skill_runs_all = skill_run_q.order_by(SkillRun.created_at.desc()).all()
    skill_window = int(os.environ.get("OPENATLAS_DASHBOARD_SKILL_RUN_WINDOW", "100"))
    skill_runs = skill_runs_all[:max(1, skill_window)]
    artifacts = artifact_q.all()
    input_tokens = sum(int(m.input_tokens or 0) for m in messages)
    output_tokens = sum(int(m.output_tokens or 0) for m in messages)
    failed = sum(1 for a in audits if any(k in (a.action or "").lower() for k in ("error", "failed", "fail")))
    task_status_counts: dict[str, int] = {}
    for s in sessions:
        status = getattr(s, "task_status", "draft") or "draft"
        task_status_counts[status] = task_status_counts.get(status, 0) + 1
    skill_fail = sum(1 for r in skill_runs if r.status in ("failed", "missing_in_hermes", "read_failed", "skipped_budget"))
    return {
        "active_users": db.query(User).filter(User.tenant_id == tid, User.is_active == True).count(),  # noqa: E712
        "messages": len(messages),
        "input_tokens": input_tokens,
        "output_tokens": output_tokens,
        "total_tokens": input_tokens + output_tokens,
        "input_chars": sum(len(m.content or "") for m in messages if m.role == "user"),
        "output_chars": sum(len(m.content or "") for m in messages if m.role == "assistant"),
        "files": len(files),
        "file_bytes": sum(int(f.size or 0) for f in files),
        "expired_files": sum(1 for f in files if _file_to_dict(f)["is_expired"]),
        "artifacts": len(artifacts),
        "artifact_kinds": {k: sum(1 for a in artifacts if a.kind == k) for k in sorted({a.kind for a in artifacts})},
        "skill_runs": len(skill_runs),
        "skill_runs_total": len(skill_runs_all),
        "skill_run_window": max(1, skill_window),
        "skill_failures": skill_fail,
        "skill_failure_rate": round(skill_fail / max(1, len(skill_runs)), 4),
        "audit_events": len(audits),
        "failure_events": failed,
        "failure_rate": round(failed / max(1, len(audits)), 4),
        "sessions": len(sessions),
        "task_status_counts": task_status_counts,
    }


def _template_to_dict(t: CollaborationTemplate) -> dict:
    try:
        participant_ids = json.loads(t.participant_ids or "[]")
    except Exception:
        participant_ids = []
    canvas_state = None
    if t.canvas_state:
        try:
            canvas_state = json.loads(t.canvas_state)
        except Exception:
            canvas_state = None
    plan_nodes: list[dict] = []
    plan_edges: list[dict] = []
    if canvas_state:
        raw_nodes = canvas_state.get("nodes") if isinstance(canvas_state, dict) else []
        raw_edges = canvas_state.get("edges") if isinstance(canvas_state, dict) else []
        for idx, node in enumerate(raw_nodes or []):
            if not isinstance(node, dict):
                continue
            data = node.get("data") if isinstance(node.get("data"), dict) else {}
            role = data.get("role") or ("user" if node.get("id") == "user" else "employee")
            plan_nodes.append({
                "node_id": node.get("id"),
                "order": idx,
                "role": role,
                "employee_id": data.get("employeeId") or data.get("employee_id"),
                "name": data.get("label") or data.get("name"),
                "default_prompt": data.get("prompt") or "",
                "output_type": data.get("outputType") or data.get("output_type") or "markdown",
                "skills": data.get("skills") if isinstance(data.get("skills"), list) else [],
            })
        for edge in raw_edges or []:
            if isinstance(edge, dict):
                plan_edges.append({
                    "source": edge.get("source"),
                    "target": edge.get("target"),
                    "mode": edge.get("mode") or "relay",
                })
    relay_count = max(0, len(participant_ids))
    meta = canvas_state.get("meta") if isinstance(canvas_state, dict) and isinstance(canvas_state.get("meta"), dict) else {}
    collaboration_plan = {
        "strategy": meta.get("strategy") or ("relay" if relay_count else "single"),
        "failure_strategy": meta.get("failure_strategy") or "continue_with_next_employee",
        "primary_employee_id": t.primary_employee_id,
        "participant_ids": participant_ids,
        "nodes": plan_nodes,
        "edges": plan_edges,
        "output_types": sorted({n.get("output_type") or "markdown" for n in plan_nodes if n.get("role") == "employee"}) or ["markdown"],
        "default_prompt": meta.get("default_prompt") or "",
        "relay_count": relay_count,
    }
    return {
        "id": t.id,
        "name": t.name,
        "description": t.description,
        "category": getattr(t, "category", "general") or "general",
        "visibility": getattr(t, "visibility", "private") or "private",
        "primary_employee_id": t.primary_employee_id,
        "participant_ids": participant_ids,
        "canvas_state": canvas_state,
        "collaboration_plan": collaboration_plan,
        "source_session_id": t.source_session_id,
        "created_at": t.created_at.isoformat() if t.created_at else None,
        "updated_at": t.updated_at.isoformat() if t.updated_at else None,
    }


def _canvas_event_to_dict(ev: CanvasEvent) -> dict:
    try:
        payload = json.loads(ev.payload or "{}")
    except Exception:
        payload = {}
    return {
        "id": ev.id,
        "tenant_id": ev.tenant_id,
        "user_id": ev.user_id,
        "session_id": ev.session_id,
        "employee_id": ev.employee_id,
        "event_type": ev.event_type,
        "node_id": ev.node_id,
        "edge_id": ev.edge_id,
        "payload": payload,
        "created_at": ev.created_at.isoformat() if ev.created_at else None,
    }


def _json_loads_obj(value: str | None, default: Any = None) -> Any:
    if default is None:
        default = {}
    try:
        return json.loads(value or "")
    except Exception:
        return default


def _session_run_to_dict(row: SessionRun) -> dict:
    return {
        "id": row.id,
        "session_id": row.session_id,
        "employee_id": row.employee_id,
        "hermes_run_id": row.hermes_run_id,
        "status": row.status,
        "stage": row.stage,
        "reason": row.reason,
        "last_event_type": row.last_event_type,
        "event_count": row.event_count,
        "payload": _json_loads_obj(row.payload, {}),
        "started_at": row.started_at.isoformat() if row.started_at else None,
        "completed_at": row.completed_at.isoformat() if row.completed_at else None,
        "created_at": row.created_at.isoformat() if row.created_at else None,
        "updated_at": row.updated_at.isoformat() if row.updated_at else None,
    }


def _workflow_node_to_dict(row: WorkflowNodeRun) -> dict:
    return {
        "id": row.id,
        "workflow_run_id": row.workflow_run_id,
        "session_id": row.session_id,
        "employee_id": row.employee_id,
        "node_id": row.node_id,
        "label": row.label,
        "status": row.status,
        "run_id": row.run_id,
        "hermes_run_id": row.hermes_run_id,
        "event_count": row.event_count,
        "input_summary": row.input_summary,
        "output_summary": row.output_summary,
        "artifact_ids": _json_loads_obj(row.artifact_ids, []),
        "error": row.error,
        "payload": _json_loads_obj(row.payload, {}),
        "started_at": row.started_at.isoformat() if row.started_at else None,
        "completed_at": row.completed_at.isoformat() if row.completed_at else None,
        "created_at": row.created_at.isoformat() if row.created_at else None,
        "updated_at": row.updated_at.isoformat() if row.updated_at else None,
    }


def _workflow_run_to_dict(row: WorkflowRun, nodes: list[WorkflowNodeRun] | None = None) -> dict:
    return {
        "id": row.id,
        "session_id": row.session_id,
        "template_id": row.template_id,
        "status": row.status,
        "strategy": row.strategy,
        "summary": row.summary,
        "payload": _json_loads_obj(row.payload, {}),
        "started_at": row.started_at.isoformat() if row.started_at else None,
        "completed_at": row.completed_at.isoformat() if row.completed_at else None,
        "created_at": row.created_at.isoformat() if row.created_at else None,
        "updated_at": row.updated_at.isoformat() if row.updated_at else None,
        "nodes": [_workflow_node_to_dict(n) for n in (nodes or [])],
    }


def _workflow_step_to_dict(row: WorkflowStepEvent) -> dict:
    return {
        "id": row.id,
        "session_id": row.session_id,
        "workflow_run_id": row.workflow_run_id,
        "workflow_node_run_id": row.workflow_node_run_id,
        "employee_id": row.employee_id,
        "event_type": row.event_type,
        "status": row.status,
        "title": row.title,
        "summary": row.summary,
        "input_summary": row.input_summary,
        "output_summary": row.output_summary,
        "raw_event_ref": row.raw_event_ref,
        "payload": _json_loads_obj(row.payload_json, {}),
        "risk_level": row.risk_level,
        "tool_name": row.tool_name,
        "artifact_ids": _json_loads_obj(row.artifact_ids, []),
        "file_ids": _json_loads_obj(row.file_ids, []),
        "is_checkpoint": bool(row.is_checkpoint),
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }


def _workflow_checkpoint_to_dict(row: WorkflowCheckpoint) -> dict:
    return {
        "id": row.id,
        "session_id": row.session_id,
        "workflow_run_id": row.workflow_run_id,
        "workflow_node_run_id": row.workflow_node_run_id,
        "step_event_id": row.step_event_id,
        "checkpoint_type": row.checkpoint_type,
        "status": row.status,
        "summary": row.summary,
        "context_snapshot": _json_loads_obj(row.context_snapshot_json, {}),
        "hermes_session_id": row.hermes_session_id,
        "hermes_run_id": row.hermes_run_id,
        "upstream_node_outputs": _json_loads_obj(row.upstream_node_outputs_json, {}),
        "artifact_policy": _json_loads_obj(row.artifact_policy_json, {}),
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }


def _workflow_fork_to_dict(row: WorkflowRunFork) -> dict:
    return {
        "id": row.id,
        "session_id": row.session_id,
        "parent_workflow_run_id": row.parent_workflow_run_id,
        "child_workflow_run_id": row.child_workflow_run_id,
        "forked_from_node_run_id": row.forked_from_node_run_id,
        "forked_from_checkpoint_id": row.forked_from_checkpoint_id,
        "reason": row.reason,
        "payload": _json_loads_obj(row.payload_json, {}),
        "created_by": row.created_by,
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }


def _step_status_for_event(event_type: str) -> str:
    if event_type.endswith(".failed") or event_type in {"node.failed", "run.failed", "error"}:
        return "failed"
    if event_type in {"tool.started", "run.started", "node.started"}:
        return "running"
    if event_type in {"approval.required", "openatlas.approval_required"}:
        return "waiting_approval"
    if event_type in {"openatlas.run_idle", "openatlas.run_detached", "node.stalled"}:
        return "stalled"
    return "completed"


def _step_title_for_event(event_type: str, payload: dict[str, Any]) -> str:
    if event_type.startswith("tool."):
        return f"工具: {payload.get('tool_name') or payload.get('name') or payload.get('tool') or 'tool'}"
    if event_type in {"openatlas.reasoning", "reasoning.summary"}:
        return "思考摘要"
    if event_type in {"openatlas.context", "context.injected"}:
        return "上下文注入"
    if event_type in {"openatlas.approval_required", "approval.required"}:
        return "等待人工确认"
    if event_type.startswith("artifact."):
        return "交付物"
    if event_type in {"openatlas.run_idle", "openatlas.run_detached", "node.stalled"}:
        return "长任务停滞"
    if event_type == "assistant.delta":
        return "模型回复"
    return event_type


def _should_checkpoint_event(event_type: str, status: str) -> bool:
    return (
        event_type in {
            "node.started",
            "context.injected",
            "tool.completed",
            "artifact.created",
            "node.completed",
            "node.failed",
            "openatlas.run_idle",
            "openatlas.run_detached",
            "approval.resolved",
        }
        or status in {"failed", "stalled", "waiting_approval"}
    )


def _record_workflow_step(
    db: Session,
    *,
    tenant_id: str,
    user_id: str,
    session_id: str,
    workflow_run_id: str,
    workflow_node_run_id: str | None,
    employee_id: str | None,
    event_type: str,
    payload: dict[str, Any] | None = None,
    title: str = "",
    summary: str = "",
    input_summary: str = "",
    output_summary: str = "",
    raw_event_ref: str = "",
    risk_level: str = "low",
    tool_name: str = "",
    artifact_ids: list[str] | None = None,
    file_ids: list[str] | None = None,
    checkpoint_type: str | None = None,
    hermes_session_id: str = "",
    hermes_run_id: str = "",
) -> WorkflowStepEvent:
    payload = payload or {}
    status = _step_status_for_event(event_type)
    artifact_ids = artifact_ids or []
    file_ids = file_ids or []
    row = WorkflowStepEvent(
        tenant_id=tenant_id,
        user_id=user_id,
        session_id=session_id,
        workflow_run_id=workflow_run_id,
        workflow_node_run_id=workflow_node_run_id,
        employee_id=employee_id,
        event_type=event_type,
        status=status,
        title=(title or _step_title_for_event(event_type, payload))[:255],
        summary=(summary or str(payload.get("summary") or payload.get("detail") or payload.get("message") or payload.get("label") or ""))[:4000],
        input_summary=input_summary[:4000],
        output_summary=output_summary[:4000],
        raw_event_ref=raw_event_ref[:128],
        payload_json=json.dumps(payload, ensure_ascii=False),
        risk_level=risk_level[:24],
        tool_name=(tool_name or str(payload.get("tool_name") or payload.get("name") or payload.get("tool") or ""))[:128],
        artifact_ids=json.dumps(artifact_ids, ensure_ascii=False),
        file_ids=json.dumps(file_ids, ensure_ascii=False),
        is_checkpoint=bool(checkpoint_type or _should_checkpoint_event(event_type, status)),
    )
    db.add(row)
    db.flush()
    if row.is_checkpoint:
        checkpoint = WorkflowCheckpoint(
            tenant_id=tenant_id,
            user_id=user_id,
            session_id=session_id,
            workflow_run_id=workflow_run_id,
            workflow_node_run_id=workflow_node_run_id,
            step_event_id=row.id,
            checkpoint_type=checkpoint_type or event_type,
            status="available",
            summary=row.summary or row.title,
            context_snapshot_json=json.dumps({
                "event_type": event_type,
                "input_summary": input_summary,
                "output_summary": output_summary,
                "payload": payload,
            }, ensure_ascii=False),
            hermes_session_id=hermes_session_id[:128],
            hermes_run_id=hermes_run_id[:128],
            upstream_node_outputs_json=json.dumps({}, ensure_ascii=False),
            artifact_policy_json=json.dumps({
                "artifact_ids": artifact_ids,
                "on_resume": "create_new_version",
            }, ensure_ascii=False),
        )
        db.add(checkpoint)
        db.flush()
    return row


def _record_workflow_step_safely(**kwargs: Any) -> None:
    db = SessionLocal()
    try:
        _record_workflow_step(db, **kwargs)
        db.commit()
    except Exception:
        db.rollback()
    finally:
        db.close()


async def _execute_checkpoint_resume(
    *,
    workflow_run_id: str,
    workflow_node_run_id: str,
    checkpoint_id: str,
    prompt: str,
    reasoning_effort: str | None = None,
) -> None:
    """Run a checkpoint recovery branch in the background.

    This intentionally keeps the first implementation linear: resume one node,
    persist its transcript/artifacts, and leave downstream replay/fan-out to the
    next executor iteration.
    """
    session_run_id = ""
    hermes_run_id = ""
    assistant_parts: list[str] = []
    reasoning_parts: list[str] = []
    tool_records: list[dict[str, Any]] = []
    artifact_ids: list[str] = []
    with SessionLocal() as db:
        wf = db.get(WorkflowRun, workflow_run_id)
        node = db.get(WorkflowNodeRun, workflow_node_run_id)
        checkpoint = db.get(WorkflowCheckpoint, checkpoint_id)
        if not wf or not node or not checkpoint:
            return
        rec = db.get(SessionRecord, wf.session_id)
        emp = db.get(DigitalEmployee, node.employee_id) if node.employee_id else None
        if not rec:
            return
        run_row = _create_session_run(
            db,
            tenant_id=wf.tenant_id,
            user_id=wf.user_id,
            session_id=wf.session_id,
            employee_id=node.employee_id,
            status="running",
            stage="recovery",
            reason="从检查点恢复执行",
            payload={
                "workflow_run_id": wf.id,
                "workflow_node_run_id": node.id,
                "checkpoint_id": checkpoint.id,
            },
        )
        session_run_id = run_row.id
        node.status = "running"
        node.run_id = session_run_id
        node.input_summary = prompt[:4000]
        node.started_at = node.started_at or datetime.now(timezone.utc)
        wf.status = "running"
        wf.summary = f"从检查点恢复执行: {checkpoint.summary or checkpoint.checkpoint_type}"
        rec.task_status = "running"
        rec.task_summary = "正在从协作检查点恢复执行。"
        _record_workflow_step(
            db,
            tenant_id=wf.tenant_id,
            user_id=wf.user_id,
            session_id=wf.session_id,
            workflow_run_id=wf.id,
            workflow_node_run_id=node.id,
            employee_id=node.employee_id,
            event_type="node.resume_started",
            title="从检查点恢复",
            summary=checkpoint.summary or "从检查点恢复执行",
            input_summary=prompt[:1000],
            payload={"checkpoint_id": checkpoint.id, "checkpoint_type": checkpoint.checkpoint_type},
            checkpoint_type="node.resume_started",
            hermes_session_id=rec.hermes_session_id,
        )
        db.commit()
        target = await hermes_client.resolve_target(db, wf.tenant_id)
        try:
            system_message = ""
            if emp:
                hermes_home = _tenant_hermes_home(db, wf.tenant_id)
                system_message, _eff, _skills = _build_employee_system_message(
                    db,
                    tenant_id=wf.tenant_id,
                    user_id=wf.user_id,
                    employee_id=emp.id,
                    hermes_home=hermes_home,
                    file_ctx_block="",
                    recent_context_block="",
                )
        except Exception:
            system_message = ""

    try:
        with SessionLocal() as db:
            wf = db.get(WorkflowRun, workflow_run_id)
            node = db.get(WorkflowNodeRun, workflow_node_run_id)
            rec = db.get(SessionRecord, wf.session_id) if wf else None
            if not wf or not node or not rec:
                return
            run = await hermes_client.create_run(
                target,
                message=prompt,
                session_id=rec.hermes_session_id,
                system_message=system_message or None,
                reasoning_effort=reasoning_effort,
            )
            hermes_run_id = str(run.get("run_id") or "")
            _remember_hermes_run(
                hermes_run_id,
                tenant_id=wf.tenant_id,
                user_id=wf.user_id,
                session_id=wf.session_id,
                employee_id=node.employee_id or "",
                hermes_session_id=rec.hermes_session_id,
            )
            _update_session_run(db, session_run_id, status="running", stage="runtime", reason="Hermes recovery run started", event_type="run.started", hermes_run_id=hermes_run_id)
            _update_workflow_node_run(
                db,
                workflow_run=wf,
                tenant_id=wf.tenant_id,
                user_id=wf.user_id,
                session_id=wf.session_id,
                employee_id=node.employee_id,
                node_id=node.node_id,
                label=node.label,
                status="running",
                event_type="run.started",
                session_run_id=session_run_id,
                hermes_run_id=hermes_run_id,
            )
            _record_workflow_step(
                db,
                tenant_id=wf.tenant_id,
                user_id=wf.user_id,
                session_id=wf.session_id,
                workflow_run_id=wf.id,
                workflow_node_run_id=node.id,
                employee_id=node.employee_id,
                event_type="run.started",
                title="Hermes 恢复 Run 启动",
                summary=f"Hermes run {hermes_run_id} 已启动",
                payload={**run, "hermes_run_id": hermes_run_id},
                checkpoint_type="run.started",
                hermes_session_id=rec.hermes_session_id,
                hermes_run_id=hermes_run_id,
            )
            db.commit()

        idle_timeout = max(15.0, float(os.environ.get("OPENATLAS_RUN_EVENT_IDLE_TIMEOUT_SECONDS", "60")))
        stream = hermes_client.stream_run_events(target, hermes_run_id).__aiter__()
        last_event_at = datetime.now(timezone.utc)
        while True:
            try:
                run_ev = await asyncio.wait_for(stream.__anext__(), timeout=idle_timeout)
                last_event_at = datetime.now(timezone.utc)
            except StopAsyncIteration:
                break
            except asyncio.TimeoutError:
                idle_for = int((datetime.now(timezone.utc) - last_event_at).total_seconds())
                with SessionLocal() as db:
                    wf = db.get(WorkflowRun, workflow_run_id)
                    node = db.get(WorkflowNodeRun, workflow_node_run_id)
                    rec = db.get(SessionRecord, wf.session_id) if wf else None
                    if wf and node and rec:
                        _update_session_run(db, session_run_id, status="stalled", stage="runtime", reason="Hermes 长时间未推送新事件", event_type="openatlas.run_idle", hermes_run_id=hermes_run_id)
                        node.status = "stalled"
                        wf.status = "stalled"
                        rec.task_status = "running"
                        rec.task_summary = "恢复执行暂时无新事件，后台补同步仍在继续。"
                        _record_workflow_step(
                            db,
                            tenant_id=wf.tenant_id,
                            user_id=wf.user_id,
                            session_id=wf.session_id,
                            workflow_run_id=wf.id,
                            workflow_node_run_id=node.id,
                            employee_id=node.employee_id,
                            event_type="openatlas.run_idle",
                            title="恢复执行停滞",
                            summary=f"{idle_for} 秒无新事件，已保留恢复检查点。",
                            payload={"idle_seconds": idle_for, "hermes_run_id": hermes_run_id},
                            checkpoint_type="stalled",
                            hermes_session_id=rec.hermes_session_id,
                            hermes_run_id=hermes_run_id,
                        )
                        db.commit()
                _schedule_detached_run_reconcile(hermes_run_id)
                return

            ev_name = run_ev.get("event") or "message"
            ev_data = run_ev.get("data") if isinstance(run_ev.get("data"), dict) else {"raw": run_ev.get("data")}
            with SessionLocal() as db:
                wf = db.get(WorkflowRun, workflow_run_id)
                node = db.get(WorkflowNodeRun, workflow_node_run_id)
                rec = db.get(SessionRecord, wf.session_id) if wf else None
                if not wf or not node or not rec:
                    return
                if ev_name == "message.delta":
                    delta = str(ev_data.get("delta") or "")
                    if delta:
                        assistant_parts.append(delta)
                elif ev_name == "reasoning.available":
                    text = str(ev_data.get("text") or "")
                    if text:
                        reasoning_parts.append(text)
                        _record_workflow_step(
                            db,
                            tenant_id=wf.tenant_id,
                            user_id=wf.user_id,
                            session_id=wf.session_id,
                            workflow_run_id=wf.id,
                            workflow_node_run_id=node.id,
                            employee_id=node.employee_id,
                            event_type="reasoning.summary",
                            title="思考摘要",
                            summary=text[:1000],
                            payload={**ev_data, "hermes_run_id": hermes_run_id},
                            hermes_session_id=rec.hermes_session_id,
                            hermes_run_id=hermes_run_id,
                        )
                elif ev_name == "tool.started":
                    tool_name = str(ev_data.get("tool") or ev_data.get("tool_name") or "tool")
                    tool_records.append({"name": tool_name, "status": "running", "payload": ev_data})
                    _record_workflow_step(
                        db,
                        tenant_id=wf.tenant_id,
                        user_id=wf.user_id,
                        session_id=wf.session_id,
                        workflow_run_id=wf.id,
                        workflow_node_run_id=node.id,
                        employee_id=node.employee_id,
                        event_type="tool.started",
                        title=f"工具: {tool_name}",
                        summary=str(ev_data.get("preview") or ev_data.get("command") or tool_name),
                        payload={**ev_data, "hermes_run_id": hermes_run_id},
                        risk_level="high" if _is_approval_sensitive_tool(tool_name) else "low",
                        tool_name=tool_name,
                        hermes_session_id=rec.hermes_session_id,
                        hermes_run_id=hermes_run_id,
                    )
                elif ev_name == "tool.completed":
                    payload = {
                        **ev_data,
                        "tool_name": ev_data.get("tool") or ev_data.get("tool_name") or "tool",
                        "result": ev_data.get("result") or ev_data.get("output"),
                        "error": ev_data.get("error"),
                        "hermes_run_id": hermes_run_id,
                    }
                    tool_name = str(payload.get("tool_name") or "tool")
                    live_artifacts = _persist_tool_event_artifacts(
                        tenant_id=wf.tenant_id,
                        user_id=wf.user_id,
                        session_id=wf.session_id,
                        payload=payload,
                        run_id=session_run_id,
                        employee_id=node.employee_id,
                        hermes_run_id=hermes_run_id,
                    )
                    new_artifact_ids = [str(a.get("id")) for a in live_artifacts if a.get("id")]
                    artifact_ids.extend(new_artifact_ids)
                    _record_workflow_step(
                        db,
                        tenant_id=wf.tenant_id,
                        user_id=wf.user_id,
                        session_id=wf.session_id,
                        workflow_run_id=wf.id,
                        workflow_node_run_id=node.id,
                        employee_id=node.employee_id,
                        event_type="tool.failed" if payload.get("error") else "tool.completed",
                        title=f"工具: {tool_name}",
                        summary=str(payload.get("error") or payload.get("label") or tool_name),
                        payload=payload,
                        risk_level="high" if _is_approval_sensitive_tool(tool_name) else "low",
                        tool_name=tool_name,
                        artifact_ids=new_artifact_ids,
                        checkpoint_type="tool.failed" if payload.get("error") else "tool.completed",
                        hermes_session_id=rec.hermes_session_id,
                        hermes_run_id=hermes_run_id,
                    )
                elif ev_name == "run.completed":
                    completed = str(ev_data.get("output") or ev_data.get("final_response") or "")
                    if completed:
                        assistant_parts.append(completed)
                    break
                elif ev_name == "run.failed":
                    raise RuntimeError(str(ev_data.get("error") or "Hermes recovery run failed"))
                db.commit()

        full = "".join(assistant_parts).strip()
        with SessionLocal() as db:
            wf = db.get(WorkflowRun, workflow_run_id)
            node = db.get(WorkflowNodeRun, workflow_node_run_id)
            rec = db.get(SessionRecord, wf.session_id) if wf else None
            if not wf or not node or not rec:
                return
            if full:
                msg_row = MessageRecord(
                    session_id=wf.session_id,
                    role="assistant",
                    content=full,
                    reasoning=json.dumps(reasoning_parts, ensure_ascii=False),
                    tool_calls=json.dumps(tool_records, ensure_ascii=False),
                    speaker_employee_id=node.employee_id,
                    speaker_name=node.label,
                )
                db.add(msg_row)
                db.flush()
                for art in _extract_task_artifacts(full, prefix=f"{node.label or '节点'}-恢复执行"):
                    artifact = TaskArtifact(
                        tenant_id=wf.tenant_id,
                        user_id=wf.user_id,
                        session_id=wf.session_id,
                        message_id=msg_row.id,
                        kind=art["kind"],
                        name=art["name"],
                        mime_type=art["mime_type"],
                        content=art["content"],
                        source="assistant",
                        run_id=session_run_id,
                        employee_id=node.employee_id,
                        version=_next_artifact_version(db, session_id=wf.session_id, name=art["name"]),
                        provenance_payload=json.dumps({
                            "origin": "workflow_checkpoint_resume",
                            "checkpoint_id": checkpoint_id,
                            "hermes_run_id": hermes_run_id,
                        }, ensure_ascii=False),
                    )
                    db.add(artifact)
                    db.flush()
                    artifact_ids.append(artifact.id)
            node.status = "done"
            node.output_summary = full[:4000]
            node.artifact_ids = json.dumps(list(dict.fromkeys(artifact_ids)), ensure_ascii=False)
            node.completed_at = datetime.now(timezone.utc)
            wf.status = "completed"
            wf.summary = "检查点恢复执行完成。"
            wf.completed_at = datetime.now(timezone.utc)
            rec.task_status = "completed" if full else "needs_input"
            rec.task_summary = "检查点恢复执行完成。" if full else "恢复执行没有返回可见内容。"
            _update_session_run(db, session_run_id, status="completed", stage="assistant", reason="checkpoint resume completed", event_type="run.completed", hermes_run_id=hermes_run_id)
            _record_workflow_step(
                db,
                tenant_id=wf.tenant_id,
                user_id=wf.user_id,
                session_id=wf.session_id,
                workflow_run_id=wf.id,
                workflow_node_run_id=node.id,
                employee_id=node.employee_id,
                event_type="node.completed",
                title="恢复节点完成",
                summary="检查点恢复执行完成。",
                output_summary=full[:1000],
                payload={"checkpoint_id": checkpoint_id, "hermes_run_id": hermes_run_id},
                artifact_ids=list(dict.fromkeys(artifact_ids)),
                checkpoint_type="node.completed",
                hermes_session_id=rec.hermes_session_id,
                hermes_run_id=hermes_run_id,
            )
            db.commit()
    except Exception as exc:
        with SessionLocal() as db:
            wf = db.get(WorkflowRun, workflow_run_id)
            node = db.get(WorkflowNodeRun, workflow_node_run_id)
            rec = db.get(SessionRecord, wf.session_id) if wf else None
            if wf and node and rec:
                node.status = "failed"
                node.error = str(exc)[:4000]
                wf.status = "failed"
                wf.summary = str(exc)[:1000]
                wf.completed_at = datetime.now(timezone.utc)
                rec.task_status = "failed"
                rec.task_summary = str(exc)[:1000]
                _update_session_run(db, session_run_id, status="failed", stage="recovery", reason=str(exc), event_type="node.failed", hermes_run_id=hermes_run_id)
                _record_workflow_step(
                    db,
                    tenant_id=wf.tenant_id,
                    user_id=wf.user_id,
                    session_id=wf.session_id,
                    workflow_run_id=wf.id,
                    workflow_node_run_id=node.id,
                    employee_id=node.employee_id,
                    event_type="node.failed",
                    title="恢复执行失败",
                    summary=str(exc)[:1000],
                    payload={"checkpoint_id": checkpoint_id, "hermes_run_id": hermes_run_id},
                    checkpoint_type="node.failed",
                    hermes_session_id=rec.hermes_session_id,
                    hermes_run_id=hermes_run_id,
                )
                db.commit()


def _create_session_run(
    db: Session,
    *,
    tenant_id: str,
    user_id: str,
    session_id: str,
    employee_id: str | None,
    status: str = "queued",
    stage: str = "",
    reason: str = "",
    payload: dict[str, Any] | None = None,
) -> SessionRun:
    now = datetime.now(timezone.utc)
    row = SessionRun(
        tenant_id=tenant_id,
        user_id=user_id,
        session_id=session_id,
        employee_id=employee_id,
        status=status,
        stage=stage,
        reason=reason,
        payload=json.dumps(payload or {}, ensure_ascii=False),
        started_at=now if status in {"running", "waiting_approval", "waiting_input"} else None,
        updated_at=now,
    )
    db.add(row)
    db.flush()
    return row


def _update_session_run(
    db: Session,
    run_id: str | None,
    *,
    status: str | None = None,
    stage: str | None = None,
    reason: str | None = None,
    event_type: str | None = None,
    hermes_run_id: str | None = None,
    payload_patch: dict[str, Any] | None = None,
) -> SessionRun | None:
    if not run_id:
        return None
    row = db.get(SessionRun, run_id)
    if not row:
        return None
    now = datetime.now(timezone.utc)
    if status:
        row.status = status
        if status in {"running", "waiting_approval", "waiting_input"} and not row.started_at:
            row.started_at = now
        if status in {"completed", "failed", "cancelled"}:
            row.completed_at = now
    if stage is not None:
        row.stage = stage[:64]
    if reason is not None:
        row.reason = reason[:2000]
    if event_type:
        row.last_event_type = event_type[:64]
        row.event_count = int(row.event_count or 0) + 1
    if hermes_run_id:
        row.hermes_run_id = hermes_run_id[:128]
    if payload_patch:
        payload = _json_loads_obj(row.payload, {})
        if not isinstance(payload, dict):
            payload = {}
        payload.update(payload_patch)
        row.payload = json.dumps(payload, ensure_ascii=False)
    row.updated_at = now
    return row


def _node_id_for_employee(employee_id: str | None) -> str:
    return f"employee-{employee_id}" if employee_id else "user"


def _ensure_workflow_run(
    db: Session,
    *,
    tenant_id: str,
    user_id: str,
    session_id: str,
    template_id: str | None = None,
    strategy: str = "relay",
    payload: dict[str, Any] | None = None,
) -> WorkflowRun:
    row = db.query(WorkflowRun).filter(
        WorkflowRun.tenant_id == tenant_id,
        WorkflowRun.session_id == session_id,
    ).order_by(WorkflowRun.created_at.desc()).first()
    if row and row.status in {"running", "waiting_approval", "waiting_input"}:
        return row
    now = datetime.now(timezone.utc)
    row = WorkflowRun(
        tenant_id=tenant_id,
        user_id=user_id,
        session_id=session_id,
        template_id=template_id,
        status="running",
        strategy=strategy,
        payload=json.dumps(payload or {}, ensure_ascii=False),
        started_at=now,
        updated_at=now,
    )
    db.add(row)
    db.flush()
    return row


def _ensure_workflow_node_run(
    db: Session,
    *,
    workflow_run: WorkflowRun,
    tenant_id: str,
    user_id: str,
    session_id: str,
    employee_id: str | None,
    node_id: str | None = None,
    label: str = "",
) -> WorkflowNodeRun:
    resolved_node_id = node_id or _node_id_for_employee(employee_id)
    row = db.query(WorkflowNodeRun).filter(
        WorkflowNodeRun.workflow_run_id == workflow_run.id,
        WorkflowNodeRun.node_id == resolved_node_id,
    ).first()
    if row:
        return row
    row = WorkflowNodeRun(
        tenant_id=tenant_id,
        user_id=user_id,
        workflow_run_id=workflow_run.id,
        session_id=session_id,
        employee_id=employee_id,
        node_id=resolved_node_id,
        label=label or resolved_node_id,
        status="idle",
    )
    db.add(row)
    db.flush()
    return row


def _update_workflow_node_run(
    db: Session,
    *,
    workflow_run: WorkflowRun | None,
    tenant_id: str,
    user_id: str,
    session_id: str,
    employee_id: str | None,
    node_id: str | None = None,
    label: str = "",
    status: str | None = None,
    event_type: str | None = None,
    session_run_id: str | None = None,
    hermes_run_id: str | None = None,
    input_summary: str | None = None,
    output_summary: str | None = None,
    artifact_ids: list[str] | None = None,
    error: str | None = None,
    payload_patch: dict[str, Any] | None = None,
) -> WorkflowNodeRun | None:
    wf = workflow_run or _ensure_workflow_run(
        db,
        tenant_id=tenant_id,
        user_id=user_id,
        session_id=session_id,
    )
    row = _ensure_workflow_node_run(
        db,
        workflow_run=wf,
        tenant_id=tenant_id,
        user_id=user_id,
        session_id=session_id,
        employee_id=employee_id,
        node_id=node_id,
        label=label,
    )
    now = datetime.now(timezone.utc)
    if status:
        row.status = status
        if status in {"running", "waiting_approval", "waiting_input"} and not row.started_at:
            row.started_at = now
        if status in {"done", "completed", "failed", "cancelled"}:
            row.completed_at = now
    if event_type:
        row.event_count = int(row.event_count or 0) + 1
    if session_run_id:
        row.run_id = session_run_id
    if hermes_run_id:
        row.hermes_run_id = hermes_run_id[:128]
    if input_summary is not None:
        row.input_summary = input_summary[:4000]
    if output_summary is not None:
        row.output_summary = output_summary[:4000]
    if artifact_ids:
        existing = _json_loads_obj(row.artifact_ids, [])
        if not isinstance(existing, list):
            existing = []
        row.artifact_ids = json.dumps(list(dict.fromkeys([*existing, *artifact_ids])), ensure_ascii=False)
    if error is not None:
        row.error = error[:4000]
    if payload_patch:
        payload = _json_loads_obj(row.payload, {})
        if not isinstance(payload, dict):
            payload = {}
        payload.update(payload_patch)
        row.payload = json.dumps(payload, ensure_ascii=False)
    row.updated_at = now
    wf.updated_at = now
    return row


def _record_canvas_runtime_event(
    *,
    tenant_id: str,
    user_id: str,
    session_id: str,
    employee_id: str | None,
    event_type: str,
    payload: dict[str, Any] | None = None,
    node_id: str | None = None,
    edge_id: str | None = None,
) -> None:
    db = SessionLocal()
    try:
        db.add(CanvasEvent(
            tenant_id=tenant_id,
            user_id=user_id,
            session_id=session_id,
            employee_id=employee_id or None,
            event_type=event_type,
            node_id=node_id or (f"employee-{employee_id}" if employee_id else ""),
            edge_id=edge_id or "",
            payload=json.dumps(payload or {}, ensure_ascii=False),
        ))
        db.commit()
    except Exception:
        db.rollback()
    finally:
        db.close()


def _sanitize_hermes_text(s: str) -> str:
    """P3.12 (2026-06-07) Bug 5b: 历史消息清洗 — 剥 <system_prompt> <context> <file_context> 三种注入块.

    只剥完整匹配的 <tag>...</tag> 块, 不动用户原文. 用来兼容老 session (P3.11 之前没 MessageRecord).
    """
    if not s:
        return ""
    import re as _re
    # 用非贪婪匹配, 多行支持, 不区分大小写
    for tag in ("system_prompt", "context", "file_context", "skills", "hermes_skills"):
        s = _re.sub(rf"<{tag}>.*?</{tag}>\s*", "", s, flags=_re.DOTALL | _re.IGNORECASE)
    return s.strip()


def _hermes_message_role_content(item: dict[str, Any]) -> tuple[str, str]:
    role = str(item.get("role") or item.get("type") or "").strip().lower()
    content = item.get("content") or item.get("text") or item.get("message") or ""
    if isinstance(content, list):
        parts: list[str] = []
        for part in content:
            if isinstance(part, dict):
                parts.append(str(part.get("text") or part.get("content") or ""))
            else:
                parts.append(str(part))
        content = "\n".join(p for p in parts if p)
    elif not isinstance(content, str):
        content = json.dumps(content, ensure_ascii=False) if content else ""
    return role, _sanitize_hermes_text(content)


def _artifact_from_path(path: str | Path, *, source: str = "hermes_tool") -> dict | None:
    if isinstance(path, str) and (not path.strip() or len(path) > 4096 or "\n" in path or "\r" in path):
        return None
    p = Path(path).expanduser()
    try:
        if not p.exists() or not p.is_file():
            return None
    except OSError:
        return None
    try:
        resolved = p.resolve()
        home = Path.home().resolve()
        tmp_roots = {Path("/tmp").resolve(), Path("/private/tmp").resolve()}
        is_tmp_helper = any(resolved == root or root in resolved.parents for root in tmp_roots)
        is_user_visible = resolved == home or home in resolved.parents
    except Exception:
        is_tmp_helper = False
        is_user_visible = False
    try:
        file_text = p.read_text(encoding="utf-8", errors="replace")
    except Exception:
        return None
    ext = p.suffix.lower()
    if ext in {".html", ".htm"}:
        if is_tmp_helper:
            return None
        return {
            "kind": "html",
            "name": p.name,
            "mime_type": "text/html;charset=utf-8",
            "content": _wrap_html_artifact(file_text),
            "source": source,
            "source_path": str(resolved),
        }
    if ext in {".md", ".markdown"}:
        return {
            "kind": "markdown",
            "name": p.name,
            "mime_type": "text/markdown;charset=utf-8",
            "content": file_text,
            "source": source,
            "source_path": str(resolved),
        }
    if ext == ".csv":
        return {
            "kind": "table",
            "name": p.name,
            "mime_type": "text/csv;charset=utf-8",
            "content": file_text[:200000],
            "source": source,
            "source_path": str(resolved),
        }
    if ext == ".json":
        return {
            "kind": "json",
            "name": p.name,
            "mime_type": "application/json;charset=utf-8",
            "content": file_text[:200000],
            "source": source,
            "source_path": str(resolved),
        }
    if is_tmp_helper or ext in {".py", ".js", ".ts", ".tsx", ".jsx", ".sh", ".pyc"}:
        return None
    if not is_user_visible:
        return None
    return {
        "kind": "report",
        "name": p.name,
        "mime_type": mimetypes.guess_type(str(p))[0] or "text/plain;charset=utf-8",
        "content": file_text[:200000],
        "source": source,
        "source_path": str(resolved),
    }


def _candidate_paths_from_payload(payload: Any) -> list[str]:
    paths: list[str] = []

    def add_from_string(value: str) -> None:
        if not value:
            return
        if "\n" in value or "\r" in value:
            for line in re.split(r"[\r\n]+", value):
                add_from_string(line.strip())
            return
        # JSON strings often carry path fields as plain values.
        if value.startswith("~") or value.startswith("/"):
            paths.append(value.strip().strip("'\"`，,;"))
        for m in re.finditer(r"(?P<path>(?:~|/Users/|/tmp/|/private/tmp/)[^\s'\"`<>|]+)", value):
            paths.append(m.group("path").strip().strip("'\"`，,;"))

    def walk(value: Any) -> None:
        if isinstance(value, dict):
            for k, v in value.items():
                lk = str(k).lower()
                if lk in {"path", "file", "filename", "filepath", "file_path", "resolved_path", "output_path"}:
                    if isinstance(v, str):
                        add_from_string(v)
                    else:
                        walk(v)
                else:
                    walk(v)
        elif isinstance(value, list):
            for item in value:
                walk(item)
        elif isinstance(value, str):
            add_from_string(value)

    if isinstance(payload, str):
        try:
            walk(json.loads(payload))
        except Exception:
            walk(payload)
    else:
        walk(payload)
    return list(dict.fromkeys(paths))[:20]


def _tool_artifacts_from_payload(payload: Any, *, source: str = "hermes_tool") -> list[dict]:
    artifacts: list[dict] = []
    for path in _candidate_paths_from_payload(payload):
        art = _artifact_from_path(path, source=source)
        if art:
            artifacts.append(art)
    return artifacts[:12]


def _next_artifact_version(db: Session, *, session_id: str, name: str, exclude_id: str | None = None) -> int:
    q = db.query(TaskArtifact).filter(
        TaskArtifact.session_id == session_id,
        TaskArtifact.name == name,
    )
    if exclude_id:
        q = q.filter(TaskArtifact.id != exclude_id)
    versions = [
        int(getattr(row, "version", 1) or 1)
        for row in q.all()
    ]
    return (max(versions) + 1) if versions else 1


def _tool_artifacts_from_hermes_item(item: dict[str, Any], *, prefix: str = "Hermes 产物") -> list[dict]:
    role, content = _hermes_message_role_content(item)
    tool_name = str(item.get("tool_name") or item.get("name") or "").strip()
    if role != "tool" or tool_name not in {"write_file", "write", "terminal", "execute_code"} or not content:
        return []
    try:
        payload = json.loads(content)
    except Exception:
        payload = content
    return _tool_artifacts_from_payload(payload, source="hermes_tool")


def _persist_tool_event_artifacts(
    *,
    tenant_id: str,
    user_id: str,
    session_id: str,
    payload: Any,
    run_id: str | None = None,
    employee_id: str | None = None,
    hermes_run_id: str | None = None,
) -> list[dict]:
    artifacts = _tool_artifacts_from_payload(payload, source="hermes_tool_live")
    if not artifacts:
        return []
    db = SessionLocal()
    persisted: list[dict] = []
    try:
        seen = {(a.name, a.source, a.source_path) for a in db.query(TaskArtifact).filter_by(session_id=session_id).all()}
        for art in artifacts:
            key = (art["name"], art.get("source") or "hermes_tool_live", art.get("source_path") or "")
            if key in seen:
                continue
            row = TaskArtifact(
                tenant_id=tenant_id,
                user_id=user_id,
                session_id=session_id,
                message_id=None,
                kind=art["kind"],
                name=art["name"],
                mime_type=art["mime_type"],
                content=art["content"],
                source=art.get("source") or "hermes_tool_live",
                source_path=art.get("source_path") or "",
                run_id=run_id,
                employee_id=employee_id,
                version=_next_artifact_version(db, session_id=session_id, name=art["name"]),
                provenance_payload=json.dumps({
                    "origin": "tool_event",
                    "hermes_run_id": hermes_run_id or "",
                }, ensure_ascii=False),
            )
            db.add(row)
            db.flush()
            persisted.append(_artifact_to_dict(row, db))
            seen.add(key)
        db.commit()
        return persisted
    finally:
        db.close()


async def _reconcile_hermes_session_transcript(
    db: Session,
    *,
    target: Any,
    rec: SessionRecord,
    tenant_id: str,
    user_id: str,
    employee_id: str | None = None,
    speaker_name: str = "",
) -> int:
    """Import late Hermes messages/artifacts that arrived after Atlas stopped listening."""
    if not rec.hermes_session_id:
        return 0
    try:
        raw = await hermes_client.get_session_messages(target, rec.hermes_session_id)
    except Exception:
        return 0
    items = (raw.get("items") or raw.get("data") or []) if isinstance(raw, dict) else raw
    if not isinstance(items, list):
        return 0

    existing = {
        (r.role, (r.content or "").strip())
        for r in db.query(MessageRecord).filter_by(session_id=rec.id).all()
    }
    existing_assistant_texts = [
        content for role, content in existing
        if role == "assistant" and content
    ]
    imported = 0
    artifact_seen = {
        (a.name, a.source, getattr(a, "source_path", "") or "")
        for a in db.query(TaskArtifact).filter_by(session_id=rec.id).all()
    }
    for item in items:
        role, content = _hermes_message_role_content(item)
        for art in _tool_artifacts_from_hermes_item(item, prefix="Hermes 产物"):
            key = (art["name"], art.get("source") or "hermes_tool", art.get("source_path") or "")
            if key in artifact_seen:
                continue
            db.add(TaskArtifact(
                tenant_id=tenant_id,
                user_id=user_id,
                session_id=rec.id,
                message_id=None,
                kind=art["kind"],
                name=art["name"],
                mime_type=art["mime_type"],
                content=art["content"],
                source=art.get("source") or "hermes_tool",
                source_path=art.get("source_path") or "",
                employee_id=employee_id or rec.employee_id,
                version=_next_artifact_version(db, session_id=rec.id, name=art["name"]),
                provenance_payload=json.dumps({"origin": "reconcile_tool"}, ensure_ascii=False),
            ))
            artifact_seen.add(key)
            imported += 1
        if role != "assistant" or not content or ("assistant", content) in existing:
            continue
        if any(content in old or old in content for old in existing_assistant_texts):
            continue
        msg_row = MessageRecord(
            session_id=rec.id,
            role="assistant",
            content=content,
            tool_calls=(
                item.get("tool_calls")
                if isinstance(item.get("tool_calls"), str)
                else json.dumps(item.get("tool_calls") or [], ensure_ascii=False)
            ),
            reasoning=(
                item.get("reasoning")
                if isinstance(item.get("reasoning"), str)
                else json.dumps(item.get("reasoning") or item.get("reasoning_content") or [], ensure_ascii=False)
            ),
            output_tokens=_rough_token_count(content),
            total_tokens=_rough_token_count(content),
            speaker_employee_id=employee_id or rec.employee_id,
            speaker_name=speaker_name,
        )
        db.add(msg_row)
        db.flush()
        for art in _extract_task_artifacts(content, prefix=f"{speaker_name or 'Hermes'}-补同步"):
            key = (art["name"], "assistant_reconcile", "")
            if key in artifact_seen:
                continue
            db.add(TaskArtifact(
                tenant_id=tenant_id,
                user_id=user_id,
                session_id=rec.id,
                message_id=msg_row.id,
                kind=art["kind"],
                name=art["name"],
                mime_type=art["mime_type"],
                content=art["content"],
                source="assistant_reconcile",
                employee_id=employee_id or rec.employee_id,
                version=_next_artifact_version(db, session_id=rec.id, name=art["name"]),
                provenance_payload=json.dumps({"origin": "reconcile_message"}, ensure_ascii=False),
            ))
            artifact_seen.add(key)
        existing.add(("assistant", content))
        existing_assistant_texts.append(content)
        imported += 1
    if imported:
        rec.message_count = _session_message_count(db, rec.id, rec.message_count)
        rec.updated_at = datetime.now(timezone.utc)
        if rec.task_status == "running":
            rec.task_status = "completed"
        db.commit()
    return imported


# ════════════════════════════════════════════════════════════════════════════
# Routes
# ════════════════════════════════════════════════════════════════════════════


@app.get("/api/health")
def health() -> dict:
    return {"status": "ok", "service": "openatlas-backend"}


# ── Auth ────────────────────────────────────────────────────────────────────
@app.post("/api/auth/login", response_model=LoginOut)
def login(body: LoginIn, request: Request, db: Session = Depends(get_db)) -> LoginOut:
    user = db.execute(select(User).where(User.email == body.email)).scalar_one_or_none()
    if not user or not verify_password(body.password, user.password_hash):
        raise HTTPException(401, "invalid credentials")
    if not user.is_active:
        raise HTTPException(403, "user disabled")
    tenant = db.get(Tenant, user.tenant_id)
    user.last_login_at = datetime.utcnow()
    token = issue_token(user_id=user.id, tenant_id=user.tenant_id, role=user.role.value)
    db.commit()
    audit(db, principal=Principal(user, tenant), action="auth.login", resource_type="user",
          resource_id=user.id, request=request)
    db.commit()
    return LoginOut(
        access_token=token,
        user={"id": user.id, "email": user.email, "username": user.username, "role": user.role.value},
        tenant={"id": tenant.id, "slug": tenant.slug, "name": tenant.name},
    )


@app.get("/api/auth/me")
def me(p: Principal = Depends(get_principal)) -> dict:
    return {
        "user": {"id": p.user.id, "email": p.user.email, "username": p.user.username,
                 "role": p.user.role.value, "tenant_id": p.tenant.id},
        "tenant": {"id": p.tenant.id, "slug": p.tenant.slug, "name": p.tenant.name},
    }


# ── Files (P3.12 / 3.4.1 Bug 7) ─────────────────────────────────────────────
_MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB 上限, 防 OOM / 滥用


@app.post("/api/files/upload")
async def upload_file(
    request: Request,
    file: UploadFile = File(...),
    session_id: str | None = Query(default=None, alias="session_id"),
    employee_id: str | None = Query(default=None, alias="employee_id"),
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    """上传附件. multipart/form-data 接收. 存到 $OPENATLAS_HOME/uploads/{tenant}/{user}/{session}/{file_id}/.

    权限: 同租户 / 同 user. session_id 必须属于本 user. 跨租户 / 跨 user 访问会被 403.
    """
    if not file.filename:
        raise HTTPException(400, "missing filename")
    # 读 bytes
    content = await file.read()
    if len(content) > _MAX_FILE_BYTES:
        raise HTTPException(413, f"file too large (max {_MAX_FILE_BYTES} bytes)")
    # 校验 session 归属
    if session_id:
        rec = db.get(SessionRecord, session_id)
        if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
            raise HTTPException(403, "session not in this tenant/user")
    # 校验 employee 归属
    if employee_id:
        emp = db.get(DigitalEmployee, employee_id)
        if not emp or emp.tenant_id != p.tenant.id:
            raise HTTPException(403, "employee not in this tenant")
    # 落盘
    file_id = str(_uuid.uuid4())
    safe_name = _safe_filename(file.filename)
    rel_dir = f"uploads/{p.tenant.id}/{p.user.id}/{session_id or 'no-session'}/{file_id}"
    full_dir = Path(OPENATLAS_HOME) / rel_dir
    full_dir.mkdir(parents=True, exist_ok=True)
    full_path = full_dir / safe_name
    full_path.write_bytes(content)
    # 提取文本 (sync, 文件已 < 10MB, 应 < 5s)
    mime = _guess_mime(str(full_path), file.content_type)
    text = _extract_text_from_file(str(full_path), mime)
    status = "extracted" if text else ("uploaded" if mime.startswith("image/") else "uploaded")
    # 落库
    fa = FileAsset(
        id=file_id,
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=session_id,
        employee_id=employee_id,
        original_name=file.filename,
        mime_type=mime,
        size=len(content),
        storage_path=str(full_path),
        status=status,
        extracted_text=text,
    )
    db.add(fa)
    db.commit()
    db.refresh(fa)
    # 审计
    audit(db, principal=p, action="file.upload", resource_type="file",
          resource_id=fa.id, request=request,
          extra={"original_name": fa.original_name, "size": fa.size,
                 "session_id": session_id, "extracted_chars": len(text)})
    return _file_to_dict(fa)


@app.get("/api/files/{fid}")
def get_file_meta(
    fid: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    """附件元数据. 跨租户 / 跨 user → 403. 不返回 extracted_text 全文 (避免 leak)."""
    fa = db.get(FileAsset, fid)
    if not fa:
        raise HTTPException(404, "file not found")
    if fa.tenant_id != p.tenant.id or fa.user_id != p.user.id:
        raise HTTPException(403, "forbidden")
    return _file_to_dict(fa)


@app.get("/api/files")
def list_files(
    session_id: str | None = Query(default=None),
    employee_id: str | None = Query(default=None),
    include_expired: bool = Query(default=True),
    limit: int = Query(default=50, ge=1, le=200),
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    q = db.query(FileAsset).filter(
        FileAsset.tenant_id == p.tenant.id,
        FileAsset.user_id == p.user.id,
    )
    if session_id:
        rec = db.get(SessionRecord, session_id)
        if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
            raise HTTPException(404, "session not found")
        q = q.filter(FileAsset.session_id == session_id)
    if employee_id:
        emp = db.get(DigitalEmployee, employee_id)
        if not emp or emp.tenant_id != p.tenant.id:
            raise HTTPException(404, "employee not found")
        q = q.filter(FileAsset.employee_id == employee_id)
    rows = q.order_by(FileAsset.created_at.desc()).limit(limit).all()
    if not include_expired:
        rows = [f for f in rows if not _file_to_dict(f)["is_expired"]]
    return {"items": [_file_to_dict(f) for f in rows]}


@app.post("/api/files/prune-expired")
def prune_expired_files(
    request: Request,
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    if _FILE_TTL_DAYS <= 0:
        return {"ok": True, "deleted": 0, "ttl_days": _FILE_TTL_DAYS}
    cutoff = datetime.now(timezone.utc) - timedelta(days=_FILE_TTL_DAYS)
    rows = db.query(FileAsset).filter(
        FileAsset.tenant_id == p.tenant.id,
        FileAsset.user_id == p.user.id,
        FileAsset.created_at < cutoff,
    ).all()
    deleted = 0
    for fa in rows:
        try:
            Path(fa.storage_path).unlink(missing_ok=True)
        except Exception:
            pass
        db.delete(fa)
        deleted += 1
    audit(db, principal=p, action="file.prune_expired", resource_type="file",
          resource_id=None, request=request, extra={"deleted": deleted, "ttl_days": _FILE_TTL_DAYS})
    db.commit()
    return {"ok": True, "deleted": deleted, "ttl_days": _FILE_TTL_DAYS}


@app.delete("/api/files/{fid}", status_code=204)
def delete_file(
    fid: str, request: Request, p: Principal = Depends(get_principal), db: Session = Depends(get_db),
):
    fa = db.get(FileAsset, fid)
    if not fa:
        raise HTTPException(404, "file not found")
    if fa.tenant_id != p.tenant.id or fa.user_id != p.user.id:
        raise HTTPException(403, "forbidden")
    # 删磁盘
    try:
        Path(fa.storage_path).unlink(missing_ok=True)
    except Exception:
        pass
    db.delete(fa)
    db.commit()
    audit(db, principal=p, action="file.delete", resource_type="file",
          resource_id=fid, request=request)


# ── Digital Employees (CRUD) ───────────────────────────────────────────────
@app.get("/api/employees")
def list_employees(p: Principal = Depends(get_principal), db: Session = Depends(get_db)) -> dict:
    rows = db.execute(
        select(DigitalEmployee).where(
            DigitalEmployee.tenant_id == p.tenant.id,
            DigitalEmployee.status != EmployeeStatus.archived,
        )
        .order_by(DigitalEmployee.created_at.desc())
    ).scalars().all()
    return {"items": [_employee_to_dict(e) for e in rows]}


@app.post("/api/employees")
def create_employee(
    body: EmployeeIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    if p.user.role == UserRole.user:
        raise HTTPException(403, "only tenant_admin or system_admin can create employees")
    initial_skills = _resolve_initial_employee_skills(db, body.initial_skill_ids, p)
    # Phase 3.4 — resource limit
    if p.tenant.max_employees is not None:
        current = db.query(DigitalEmployee).filter(
            DigitalEmployee.tenant_id == p.tenant.id,
            DigitalEmployee.status != EmployeeStatus.archived,
        ).count()
        if current >= p.tenant.max_employees:
            raise HTTPException(
                429, f"tenant '{p.tenant.slug}' has reached max_employees={p.tenant.max_employees}"
            )
    slug = _slugify(body.display_name)
    profile_name = f"tenant_{p.tenant.slug}__employee_{slug}"
    existing = db.execute(
        select(DigitalEmployee).where(
            DigitalEmployee.tenant_id == p.tenant.id,
            DigitalEmployee.profile_name == profile_name,
        )
    ).scalar_one_or_none()
    if existing:
        raise HTTPException(409, f"employee with slug '{slug}' already exists")
    emp = DigitalEmployee(
        tenant_id=p.tenant.id,
        display_name=body.display_name,
        profile_name=profile_name,
        description=body.description,
        avatar=(body.avatar or body.display_name[:1] or "?").upper()[:1],
        status=EmployeeStatus.active,
        model=body.model,
        provider=body.provider,
        temperature=body.temperature,
        max_tokens=body.max_tokens,
        system_prompt=body.system_prompt,
        toolsets=json.dumps(body.toolsets),
        created_by=p.user.id,
    )
    db.add(emp)
    db.flush()
    # Bind any initial skills (already validated as enabled + visible).
    for skill in initial_skills:
        exists_binding = db.execute(
            select(SkillBinding).where(
                SkillBinding.tenant_id == p.tenant.id,
                SkillBinding.skill_id == skill.id,
                SkillBinding.target_type == "employee",
                SkillBinding.target_id == emp.id,
            )
        ).scalar_one_or_none()
        if exists_binding:
            continue
        db.add(SkillBinding(
            tenant_id=p.tenant.id, skill_id=skill.id, target_type="employee",
            target_id=emp.id, binding_mode="inherited",
            enabled=True, locked=(skill.scope == Scope.global_),
            created_by=p.user.id,
        ))
    audit(db, principal=p, action="employee.create", resource_type="digital_employee",
          resource_id=emp.id, request=request, extra={"profile_name": profile_name})
    db.commit()
    db.refresh(emp)
    return _employee_to_dict(emp)


@app.get("/api/employees/{emp_id}")
def get_employee(
    emp_id: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    emp = db.get(DigitalEmployee, emp_id)
    if not emp or emp.tenant_id != p.tenant.id:
        raise HTTPException(404, "employee not found")
    # Skills bound to this employee
    bindings = db.execute(
        select(SkillBinding, SkillPackage)
        .join(SkillPackage, SkillPackage.id == SkillBinding.skill_id)
        .where(SkillBinding.target_type == "employee", SkillBinding.target_id == emp_id)
    ).all()
    skills = []
    for b, s in bindings:
        d = _skill_to_dict(s)
        d["binding"] = {
            "id": b.id, "binding_mode": b.binding_mode, "enabled": b.enabled, "locked": b.locked,
        }
        skills.append(d)
    out = _employee_to_dict(emp)
    out["skills"] = skills
    return out


@app.patch("/api/employees/{emp_id}")
def patch_employee(
    emp_id: str, body: EmployeePatch, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    emp = db.get(DigitalEmployee, emp_id)
    if not emp or emp.tenant_id != p.tenant.id:
        raise HTTPException(404, "employee not found")
    for fld in ("display_name", "description", "avatar", "model", "provider",
                "temperature", "max_tokens", "system_prompt", "toolsets"):
        val = getattr(body, fld, None)
        if val is not None:
            if fld == "toolsets":
                emp.toolsets = json.dumps(val)
            else:
                setattr(emp, fld, val)
    audit(db, principal=p, action="employee.update", resource_type="digital_employee",
          resource_id=emp.id, request=request)
    db.commit()
    db.refresh(emp)
    return _employee_to_dict(emp)


@app.patch("/api/employees/{emp_id}/toolsets")
async def patch_employee_toolsets(
    emp_id: str, body: EmployeeToolsetsPatchIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "employee toolset governance")
    emp = db.get(DigitalEmployee, emp_id)
    if not emp or emp.tenant_id != p.tenant.id:
        raise HTTPException(404, "employee not found")
    target = await hermes_client.resolve_target(db, p.tenant.id)
    payload = await hermes_client.list_toolsets(target)
    valid = {item["name"] for item in _normalize_toolset_items(payload)}
    requested = []
    for name in body.toolsets:
        clean = str(name).strip()
        if not clean:
            continue
        if clean not in valid:
            raise HTTPException(400, f"unknown toolset: {clean}")
        if clean not in requested:
            requested.append(clean)
    previous = json.loads(emp.toolsets or "[]")
    emp.toolsets = json.dumps(requested)
    emp.updated_at = datetime.now(timezone.utc)
    audit(db, principal=p, action="employee.toolsets.update", resource_type="digital_employee",
          resource_id=emp.id, request=request, extra={"previous": previous, "toolsets": requested})
    db.commit()
    db.refresh(emp)
    return _employee_to_dict(emp)


@app.delete("/api/employees/{emp_id}")
def delete_employee(
    emp_id: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    emp = db.get(DigitalEmployee, emp_id)
    if not emp or emp.tenant_id != p.tenant.id:
        raise HTTPException(404, "employee not found")
    emp.status = EmployeeStatus.archived
    audit(db, principal=p, action="employee.archive", resource_type="digital_employee",
          resource_id=emp.id, request=request)
    db.commit()
    return {"ok": True}


# ── Sessions + SSE chat stream ─────────────────────────────────────────────
@app.get("/api/sessions")
def list_sessions(
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    # Bug 10 (2026-06-06): 0 消息会话不展现.  之前 list_sessions 不读 message_count
    # 也不过滤, 5+ 个 "0 条消息" 会话都展示在 LeftAside 列表里.  修法: 加 .where(
    # message_count > 0).  注: message_count 在 session_chat_stream 入口 +1, 只
    # 要发过消息就 >0, 还没说话的不显示.
    rows = db.execute(
        select(SessionRecord).where(SessionRecord.tenant_id == p.tenant.id,
                                    SessionRecord.user_id == p.user.id,
                                    SessionRecord.archived == False,  # noqa: E712
                                    SessionRecord.message_count > 0)
        .order_by(SessionRecord.pinned.desc(), SessionRecord.updated_at.desc())
    ).scalars().all()
    return {"items": [
        {
            "id": r.id,
            "employee_id": r.employee_id,
            "hermes_session_id": r.hermes_session_id,
            # Bug 3/4/11: 接力员工列表透传给前端, 群聊 relay 显示
            "participant_ids": json.loads(r.participant_ids or "[]"),
            "is_group": bool(json.loads(r.participant_ids or "[]")),
            "title": r.title,
            "task_status": getattr(r, "task_status", "draft") or "draft",
            "task_summary": getattr(r, "task_summary", "") or "",
            "pinned": bool(getattr(r, "pinned", False)),
            "workspace": getattr(r, "workspace", "") or "",
            "model_override": getattr(r, "model_override", "") or "",
            "status": "archived" if getattr(r, "archived", False) else "active",
            "archived": bool(getattr(r, "archived", False)),
            "last_message": r.last_message,
            "message_count": _session_message_count(db, r.id, r.message_count),
            # Bug 9 (2026-06-06): models.py _now 改 timezone-aware UTC,
            # isoformat() 输出带 "+00:00" 后缀, 前端 new Date() 按 UTC 解析不再丢 8h.
            "created_at": r.created_at.isoformat(),
            "updated_at": r.updated_at.isoformat(),
        } for r in rows
    ]}


@app.get("/api/run-queue")
def run_queue(
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    statuses = {"running", "needs_input", "failed"}
    rows = (
        db.query(SessionRecord, DigitalEmployee)
        .outerjoin(DigitalEmployee, SessionRecord.employee_id == DigitalEmployee.id)
        .filter(
            SessionRecord.tenant_id == p.tenant.id,
            SessionRecord.user_id == p.user.id,
            SessionRecord.archived == False,  # noqa: E712
            SessionRecord.message_count > 0,
            SessionRecord.task_status.in_(statuses),
        )
        .order_by(SessionRecord.updated_at.desc())
        .limit(30)
        .all()
    )
    return {
        "items": [
            {
                "id": r.id,
                "employee_id": r.employee_id,
                "employee_name": emp.display_name if emp else "",
                "title": r.title,
                "task_status": r.task_status or "draft",
                "task_summary": r.task_summary or "",
                "is_stale": _is_stale_running_session(r),
                "last_message": r.last_message,
                "message_count": _session_message_count(db, r.id, r.message_count),
                "updated_at": r.updated_at.isoformat() if r.updated_at else None,
                "participant_ids": json.loads(r.participant_ids or "[]"),
                "is_group": bool(json.loads(r.participant_ids or "[]")),
            }
            for r, emp in rows
        ]
    }


@app.post("/api/sessions")
async def create_session(
    body: SessionCreateIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    # Phase 3.4 — resource limit
    if p.tenant.max_sessions is not None:
        current = db.query(SessionRecord).filter(
            SessionRecord.tenant_id == p.tenant.id,
        ).count()
        if current >= p.tenant.max_sessions:
            raise HTTPException(
                429, f"tenant '{p.tenant.slug}' has reached max_sessions={p.tenant.max_sessions}"
            )
    if body.employee_id:
        emp = db.get(DigitalEmployee, body.employee_id)
        if not emp or emp.tenant_id != p.tenant.id:
            raise HTTPException(404, "employee not found in this tenant")
    # Hermes enforces unique titles, so always uniquify what we send.
    import uuid as _uuid
    unique_title = (body.title or "openatlas session") + "-" + _uuid.uuid4().hex[:6]
    target = await hermes_client.resolve_target(db, p.tenant.id)
    try:
        hermes = await hermes_client.create_session(target, title=unique_title)
    except Exception as ex:
        raise HTTPException(502, f"hermes create_session failed: {ex}")
    # Hermes can wrap the session in either {"id": ...} or {"session": {"id": ...}} or {"data": {"id": ...}}
    cand = hermes.get("id") or hermes.get("session_id")
    if not cand:
        inner = hermes.get("session") or hermes.get("data") or {}
        cand = inner.get("id") or inner.get("session_id")
    hermes_sid = cand
    if not hermes_sid:
        raise HTTPException(502, f"hermes did not return session id: {hermes}")
    rec = SessionRecord(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        employee_id=body.employee_id,
        hermes_session_id=hermes_sid,
        title=body.title or "新会话",
        # Bug 3/4/11 (2026-06-06): 群聊接力员工列表落库 (json 字符串).
        # 之前完全没存, 群聊变单聊.  现在 list_sessions 透传给前端 CommandCenter.
        participant_ids=json.dumps(body.participant_ids or []),
    )
    db.add(rec)
    audit(db, principal=p, action="session.create", resource_type="session",
          resource_id=rec.id, request=request,
          extra={"hermes_sid": hermes_sid, "participant_ids": body.participant_ids or []})
    db.commit()
    db.refresh(rec)
    return {
        "id": rec.id,
        "employee_id": rec.employee_id,
        "hermes_session_id": rec.hermes_session_id,
        "participant_ids": json.loads(rec.participant_ids or "[]"),
        "is_group": bool(json.loads(rec.participant_ids or "[]")),
        "title": rec.title,
        "task_status": getattr(rec, "task_status", "draft") or "draft",
        "task_summary": getattr(rec, "task_summary", "") or "",
        "pinned": bool(getattr(rec, "pinned", False)),
        "workspace": getattr(rec, "workspace", "") or "",
        "model_override": getattr(rec, "model_override", "") or "",
        "last_message": rec.last_message,
        "message_count": _session_message_count(db, rec.id, rec.message_count),
        "created_at": rec.created_at.isoformat(),
        "updated_at": rec.updated_at.isoformat(),
    }


@app.post("/api/sessions/maintenance/stale")
async def maintain_stale_sessions(
    request: Request,
    limit: int = Query(default=20, ge=1, le=100),
    dry_run: bool = Query(default=False),
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    """Reconcile and close stale current-user running sessions.

    The operation is deliberately non-destructive: old running sessions are
    either completed by transcript reconciliation or moved to needs_input with a
    clear summary so they no longer pollute the active run queue indefinitely.
    """
    rows = (
        db.query(SessionRecord)
        .filter(
            SessionRecord.tenant_id == p.tenant.id,
            SessionRecord.user_id == p.user.id,
            SessionRecord.archived == False,  # noqa: E712
            SessionRecord.task_status == "running",
        )
        .order_by(SessionRecord.updated_at.asc())
        .limit(limit)
        .all()
    )
    stale_rows = [r for r in rows if _is_stale_running_session(r)]
    target = None
    changed: list[dict[str, Any]] = []
    for rec in stale_rows:
        imported = 0
        previous = rec.task_status or "draft"
        if not dry_run:
            if target is None:
                target = await hermes_client.resolve_target(db, p.tenant.id)
            emp = db.get(DigitalEmployee, rec.employee_id) if rec.employee_id else None
            imported = await _reconcile_hermes_session_transcript(
                db,
                target=target,
                rec=rec,
                tenant_id=p.tenant.id,
                user_id=p.user.id,
                employee_id=rec.employee_id,
                speaker_name=emp.display_name if emp else "",
            )
            db.refresh(rec)
            if rec.task_status == "running" and _is_stale_running_session(rec):
                rec.task_status = "needs_input"
                rec.task_summary = (
                    "任务超过 30 分钟没有新事件。OpenAtlas 已停止运行等待，"
                    "请打开会话补充信息、重新进行或检查 Hermes Runtime。"
                )
                latest_run = db.query(SessionRun).filter(
                    SessionRun.session_id == rec.id,
                    SessionRun.tenant_id == p.tenant.id,
                ).order_by(SessionRun.created_at.desc()).first()
                if latest_run:
                    _update_session_run(
                        db,
                        latest_run.id,
                        status="waiting_input",
                        stage="stale",
                        reason=rec.task_summary,
                        event_type="maintenance.stale",
                    )
                latest_wf = db.query(WorkflowRun).filter(
                    WorkflowRun.session_id == rec.id,
                    WorkflowRun.tenant_id == p.tenant.id,
                ).order_by(WorkflowRun.created_at.desc()).first()
                if latest_wf:
                    latest_wf.status = "waiting_input"
                    latest_wf.summary = rec.task_summary
                    latest_wf.updated_at = datetime.now(timezone.utc)
                if (rec.title or "").strip() in {"", "新会话"} and (rec.last_message or "").strip():
                    rec.title = (rec.last_message or "").strip().replace("\n", " ")[:64]
                rec.updated_at = datetime.now(timezone.utc)
            audit(db, principal=p, action="session.maintain_stale", resource_type="session",
                  resource_id=rec.id, request=request,
                  extra={"previous_status": previous, "task_status": rec.task_status, "imported": imported})
        changed.append({
            "id": rec.id,
            "title": rec.title,
            "previous_status": previous,
            "task_status": rec.task_status,
            "imported": imported,
            "is_stale": _is_stale_running_session(rec),
        })
    if not dry_run:
        db.commit()
    return {
        "ok": True,
        "dry_run": dry_run,
        "scanned": len(rows),
        "stale": len(stale_rows),
        "updated": 0 if dry_run else len(changed),
        "items": changed,
    }


@app.post("/api/sessions/{sid}/resume")
def resume_session(
    sid: str, body: SessionResumeIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id:
        raise HTTPException(404, "session not found")
    previous = rec.task_status or "draft"
    rec.task_status = "running" if body.message.strip() else "draft"
    rec.updated_at = datetime.now(timezone.utc)
    audit(db, principal=p, action="session.resume", resource_type="session",
          resource_id=rec.id, request=request,
          extra={"previous_status": previous, "has_message": bool(body.message.strip())})
    db.commit()
    return {
        "id": rec.id,
        "task_status": rec.task_status,
        "previous_status": previous,
        "next_action": "send_message" if not body.message.strip() else "stream_chat",
    }


@app.patch("/api/sessions/{sid}")
def patch_session(
    sid: str, body: SessionPatchIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    changes: dict[str, Any] = {}
    if body.title is not None:
        title = body.title.strip()
        if not title:
            raise HTTPException(400, "title cannot be empty")
        rec.title = title[:255]
        changes["title"] = rec.title
    if body.pinned is not None:
        rec.pinned = bool(body.pinned)
        changes["pinned"] = rec.pinned
    if body.workspace is not None:
        rec.workspace = body.workspace.strip()[:128]
        changes["workspace"] = rec.workspace
    if body.model_override is not None:
        rec.model_override = body.model_override.strip()[:128]
        changes["model_override"] = rec.model_override
    if not changes:
        raise HTTPException(400, "no fields to update")
    rec.updated_at = datetime.now(timezone.utc)
    audit(db, principal=p, action="session.update", resource_type="session",
          resource_id=rec.id, request=request, extra=changes)
    db.commit()
    db.refresh(rec)
    return {
        "id": rec.id,
        "employee_id": rec.employee_id,
        "hermes_session_id": rec.hermes_session_id,
        "participant_ids": json.loads(rec.participant_ids or "[]"),
        "is_group": bool(json.loads(rec.participant_ids or "[]")),
        "title": rec.title,
        "task_status": getattr(rec, "task_status", "draft") or "draft",
        "task_summary": getattr(rec, "task_summary", "") or "",
        "pinned": bool(getattr(rec, "pinned", False)),
        "workspace": getattr(rec, "workspace", "") or "",
        "model_override": getattr(rec, "model_override", "") or "",
        "status": "active",
        "last_message": rec.last_message,
        "message_count": _session_message_count(db, rec.id, rec.message_count),
        "created_at": rec.created_at.isoformat(),
        "updated_at": rec.updated_at.isoformat(),
    }


@app.delete("/api/sessions/{sid}")
def delete_session(
    sid: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id:
        raise HTTPException(404, "session not found")
    previous = rec.task_status or "draft"
    if not rec.archived:
        rec.archived = True
        if rec.task_status == "running":
            rec.task_status = "completed"
        rec.updated_at = datetime.now(timezone.utc)
        audit(db, principal=p, action="session.archive", resource_type="session",
              resource_id=rec.id, request=request,
              extra={"hermes_session_id": rec.hermes_session_id, "previous_status": previous})
        db.commit()
    return {"ok": True, "id": rec.id, "previous_status": previous, "task_status": rec.task_status}


@app.post("/api/sessions/{sid}/fork")
async def fork_session(
    sid: str, body: SessionForkIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    if p.tenant.max_sessions is not None:
        current = db.query(SessionRecord).filter(SessionRecord.tenant_id == p.tenant.id).count()
        if current >= p.tenant.max_sessions:
            raise HTTPException(429, f"tenant '{p.tenant.slug}' has reached max_sessions={p.tenant.max_sessions}")

    fork_title = (body.title or f"{rec.title or '新会话'} · 分支").strip()
    hermes_fork_id = f"openatlas_fork_{_uuid.uuid4().hex[:16]}"
    target = await hermes_client.resolve_target(db, p.tenant.id)
    try:
        hermes = await hermes_client.fork_session(
            target,
            rec.hermes_session_id,
            title=fork_title,
            fork_id=hermes_fork_id,
        )
    except Exception as ex:
        raise HTTPException(502, f"hermes fork_session failed: {ex}")
    hermes_sid = _extract_hermes_session_id(hermes) or hermes_fork_id

    fork = SessionRecord(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        employee_id=rec.employee_id,
        hermes_session_id=hermes_sid,
        title=fork_title,
        last_message=rec.last_message,
        message_count=_session_message_count(db, rec.id, rec.message_count),
        task_status="draft",
        task_summary=rec.task_summary,
        pinned=False,
        workspace=getattr(rec, "workspace", "") or "",
        model_override=getattr(rec, "model_override", "") or "",
        participant_ids=rec.participant_ids or "[]",
        canvas_state=rec.canvas_state or "",
        reusable_template_id=rec.reusable_template_id,
    )
    db.add(fork)
    db.flush()

    source_messages = db.query(MessageRecord).filter(
        MessageRecord.session_id == rec.id,
    ).order_by(MessageRecord.created_at.asc()).all()
    for msg in source_messages:
        db.add(MessageRecord(
            session_id=fork.id,
            role=msg.role,
            content=msg.content,
            model_message=msg.model_message,
            tool_calls=msg.tool_calls,
            reasoning=msg.reasoning,
            attachments=msg.attachments,
            input_tokens=msg.input_tokens,
            output_tokens=msg.output_tokens,
            total_tokens=msg.total_tokens,
            speaker_employee_id=msg.speaker_employee_id,
            speaker_name=msg.speaker_name,
            turn_index=msg.turn_index,
        ))
    audit(db, principal=p, action="session.fork", resource_type="session",
          resource_id=fork.id, request=request, extra={
              "source_session_id": rec.id,
              "source_hermes_session_id": rec.hermes_session_id,
              "hermes_session_id": hermes_sid,
              "copied_messages": len(source_messages),
          })
    db.commit()
    db.refresh(fork)
    return {
        "id": fork.id,
        "employee_id": fork.employee_id,
        "hermes_session_id": fork.hermes_session_id,
        "participant_ids": json.loads(fork.participant_ids or "[]"),
        "is_group": bool(json.loads(fork.participant_ids or "[]")),
        "title": fork.title,
        "task_status": fork.task_status or "draft",
        "task_summary": fork.task_summary or "",
        "pinned": bool(getattr(fork, "pinned", False)),
        "workspace": getattr(fork, "workspace", "") or "",
        "model_override": getattr(fork, "model_override", "") or "",
        "last_message": fork.last_message,
        "message_count": _session_message_count(db, fork.id, fork.message_count),
        "created_at": fork.created_at.isoformat() if fork.created_at else None,
        "updated_at": fork.updated_at.isoformat() if fork.updated_at else None,
    }


@app.get("/api/sessions/{sid}")
async def session_detail(
    sid: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    last_msg = db.query(MessageRecord).filter_by(session_id=sid).order_by(MessageRecord.created_at.desc()).first()
    if (rec.task_status == "running" or _is_stale_running_session(rec) or (last_msg and last_msg.role == "user")):
        target = await hermes_client.resolve_target(db, p.tenant.id)
        emp = db.get(DigitalEmployee, rec.employee_id) if rec.employee_id else None
        await _reconcile_hermes_session_transcript(
            db,
            target=target,
            rec=rec,
            tenant_id=p.tenant.id,
            user_id=p.user.id,
            employee_id=rec.employee_id,
            speaker_name=emp.display_name if emp else "",
        )
    # P3.11 (2026-06-06): 透传 participant_ids / is_group, 跟 list_sessions / create_session 一致
    import json as _json
    pid_list = _json.loads(rec.participant_ids or "[]")
    canvas_state = None
    if rec.canvas_state:
        try:
            canvas_state = _json.loads(rec.canvas_state)
        except Exception:
            canvas_state = None
    artifacts = db.query(TaskArtifact).filter(
        TaskArtifact.session_id == rec.id,
        TaskArtifact.tenant_id == p.tenant.id,
    ).order_by(TaskArtifact.created_at.desc()).all()
    context_rows = db.query(ContextInjection).filter(
        ContextInjection.session_id == rec.id,
        ContextInjection.tenant_id == p.tenant.id,
    ).order_by(ContextInjection.created_at.desc()).limit(120).all()
    return {
        "id": rec.id,
        "employee_id": rec.employee_id,
        "hermes_session_id": rec.hermes_session_id,
        "title": rec.title or "新会话",
        "task_status": getattr(rec, "task_status", "draft") or "draft",
        "task_summary": getattr(rec, "task_summary", "") or "",
        "pinned": bool(getattr(rec, "pinned", False)),
        "workspace": getattr(rec, "workspace", "") or "",
        "model_override": getattr(rec, "model_override", "") or "",
        "summary_updated_at": rec.summary_updated_at.isoformat() if getattr(rec, "summary_updated_at", None) else None,
        "created_at": rec.created_at.isoformat() if rec.created_at else None,
        "updated_at": rec.updated_at.isoformat() if rec.updated_at else None,
        "last_message": rec.last_message,
        "message_count": _session_message_count(db, rec.id, rec.message_count),
        "participant_ids": pid_list,
        "is_group": bool(pid_list),
        "canvas_state": canvas_state,
        "reusable_template_id": rec.reusable_template_id,
        "summary": _session_summary(db, rec),
        "artifacts": [_artifact_to_dict(a, db) for a in artifacts],
        "context_injections": [_context_to_dict(c) for c in context_rows],
        "health": _session_health_snapshot(db, rec),
    }


@app.get("/api/sessions/{sid}/health")
async def session_health(
    sid: str,
    reconcile: bool = Query(default=False),
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    imported = 0
    if reconcile:
        target = await hermes_client.resolve_target(db, p.tenant.id)
        emp = db.get(DigitalEmployee, rec.employee_id) if rec.employee_id else None
        imported = await _reconcile_hermes_session_transcript(
            db,
            target=target,
            rec=rec,
            tenant_id=p.tenant.id,
            user_id=p.user.id,
            employee_id=rec.employee_id,
            speaker_name=emp.display_name if emp else "",
        )
    return {"id": rec.id, "imported": imported, "health": _session_health_snapshot(db, rec)}


@app.get("/api/sessions/{sid}/runs")
def list_session_runs(
    sid: str,
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    rows = db.query(SessionRun).filter(
        SessionRun.session_id == sid,
        SessionRun.tenant_id == p.tenant.id,
    ).order_by(SessionRun.created_at.desc()).limit(50).all()
    return {"items": [_session_run_to_dict(r) for r in rows], "total": len(rows)}


@app.get("/api/sessions/{sid}/replay")
def get_session_replay(
    sid: str,
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    wf = db.query(WorkflowRun).filter(
        WorkflowRun.session_id == sid,
        WorkflowRun.tenant_id == p.tenant.id,
    ).order_by(WorkflowRun.created_at.desc()).first()
    nodes: list[WorkflowNodeRun] = []
    if wf:
        nodes = db.query(WorkflowNodeRun).filter(
            WorkflowNodeRun.workflow_run_id == wf.id,
        ).order_by(WorkflowNodeRun.created_at.asc()).all()
    events = db.query(CanvasEvent).filter(
        CanvasEvent.session_id == sid,
        CanvasEvent.tenant_id == p.tenant.id,
    ).order_by(CanvasEvent.created_at.asc()).limit(300).all()
    runs = db.query(SessionRun).filter(
        SessionRun.session_id == sid,
        SessionRun.tenant_id == p.tenant.id,
    ).order_by(SessionRun.created_at.asc()).limit(80).all()
    steps: list[WorkflowStepEvent] = []
    checkpoints: list[WorkflowCheckpoint] = []
    forks: list[WorkflowRunFork] = []
    if wf:
        steps = db.query(WorkflowStepEvent).filter(
            WorkflowStepEvent.workflow_run_id == wf.id,
            WorkflowStepEvent.tenant_id == p.tenant.id,
        ).order_by(WorkflowStepEvent.created_at.asc()).limit(500).all()
        checkpoints = db.query(WorkflowCheckpoint).filter(
            WorkflowCheckpoint.workflow_run_id == wf.id,
            WorkflowCheckpoint.tenant_id == p.tenant.id,
        ).order_by(WorkflowCheckpoint.created_at.asc()).limit(200).all()
        forks = db.query(WorkflowRunFork).filter(
            WorkflowRunFork.session_id == sid,
            WorkflowRunFork.tenant_id == p.tenant.id,
        ).order_by(WorkflowRunFork.created_at.asc()).limit(50).all()
    artifacts = db.query(TaskArtifact).filter(
        TaskArtifact.session_id == sid,
        TaskArtifact.tenant_id == p.tenant.id,
    ).order_by(TaskArtifact.created_at.asc()).limit(80).all()
    node_steps: dict[str, list[dict[str, Any]]] = {}
    node_checkpoints: dict[str, list[dict[str, Any]]] = {}
    for step in steps:
        key = step.workflow_node_run_id or ""
        node_steps.setdefault(key, []).append(_workflow_step_to_dict(step))
    for checkpoint in checkpoints:
        key = checkpoint.workflow_node_run_id or ""
        node_checkpoints.setdefault(key, []).append(_workflow_checkpoint_to_dict(checkpoint))
    workflow_payload = _workflow_run_to_dict(wf, nodes) if wf else None
    if workflow_payload:
        for node in workflow_payload.get("nodes") or []:
            node_id = node.get("id") or ""
            node["steps"] = node_steps.get(node_id, [])
            node["checkpoints"] = node_checkpoints.get(node_id, [])
    return {
        "session": {
            "id": rec.id,
            "title": rec.title or "新会话",
            "task_status": rec.task_status,
            "task_summary": rec.task_summary,
        },
        "workflow_run": workflow_payload,
        "steps": [_workflow_step_to_dict(s) for s in steps],
        "checkpoints": [_workflow_checkpoint_to_dict(c) for c in checkpoints],
        "forks": [_workflow_fork_to_dict(f) for f in forks],
        "runs": [_session_run_to_dict(r) for r in runs],
        "events": [_canvas_event_to_dict(e) for e in events],
        "artifacts": [_artifact_to_dict(a, db) for a in artifacts],
        "total": len(events) + len(runs) + len(nodes) + len(steps),
    }


@app.post("/api/sessions/{sid}/workflow-nodes/{node_run_id}/action")
def workflow_node_action(
    sid: str,
    node_run_id: str,
    body: WorkflowNodeActionIn,
    request: Request,
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    node = db.get(WorkflowNodeRun, node_run_id)
    if not node or node.tenant_id != p.tenant.id or node.session_id != sid:
        raise HTTPException(404, "workflow node run not found")
    wf = db.get(WorkflowRun, node.workflow_run_id)
    action = body.action or "continue"
    node.status = "retry_requested" if action == "retry" else "continue_requested"
    node.updated_at = datetime.now(timezone.utc)
    if wf:
        wf.status = "waiting_input"
        wf.summary = f"节点 {node.label or node.node_id} 已请求{'重试' if action == 'retry' else '继续'}，等待用户确认发送。"
        wf.updated_at = datetime.now(timezone.utc)
    rec.task_status = "needs_input"
    rec.task_summary = f"协作节点「{node.label or node.node_id}」等待用户确认{'重试' if action == 'retry' else '继续'}。"
    rec.updated_at = datetime.now(timezone.utc)
    prompt = body.message.strip() or (
        f"请从协作节点「{node.label or node.node_id}」{'重新执行' if action == 'retry' else '继续执行'}。"
        f"\n节点输入摘要: {(node.input_summary or rec.last_message or '')[:800]}"
        f"\n上次输出/错误: {(node.error or node.output_summary or '')[:800]}"
        "\n请只处理该节点负责的部分，并在完成后说明生成了哪些交付物。"
    )
    ev = CanvasEvent(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=rec.id,
        employee_id=node.employee_id,
        event_type=f"node.{action}_requested",
        node_id=node.node_id,
        payload=json.dumps({
            "node_run_id": node.id,
            "workflow_run_id": node.workflow_run_id,
            "continuation_message": prompt,
        }, ensure_ascii=False),
    )
    db.add(ev)
    audit(db, principal=p, action=f"workflow_node.{action}", resource_type="workflow_node_run",
          resource_id=node.id, request=request, extra={"session_id": sid, "node_id": node.node_id})
    db.commit()
    db.refresh(node)
    return {
        "ok": True,
        "action": action,
        "session_id": sid,
        "task_status": rec.task_status,
        "node": _workflow_node_to_dict(node),
        "workflow_run": _workflow_run_to_dict(wf) if wf else None,
        "continuation_message": prompt,
    }


@app.post("/api/sessions/{sid}/workflow-checkpoints/{checkpoint_id}/resume")
async def resume_workflow_checkpoint(
    sid: str,
    checkpoint_id: str,
    body: WorkflowCheckpointResumeIn,
    request: Request,
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    checkpoint = db.get(WorkflowCheckpoint, checkpoint_id)
    if not checkpoint or checkpoint.tenant_id != p.tenant.id or checkpoint.session_id != sid:
        raise HTTPException(404, "workflow checkpoint not found")
    parent_wf = db.get(WorkflowRun, checkpoint.workflow_run_id)
    parent_node = db.get(WorkflowNodeRun, checkpoint.workflow_node_run_id) if checkpoint.workflow_node_run_id else None
    if not parent_wf or parent_wf.tenant_id != p.tenant.id:
        raise HTTPException(404, "workflow run not found")
    if not parent_node or not parent_node.employee_id:
        raise HTTPException(400, "checkpoint is not attached to an executable employee node")
    snapshot = _json_loads_obj(checkpoint.context_snapshot_json, {})
    prompt = body.message.strip() or (
        "请从 OpenAtlas 协作工作流检查点继续执行。\n"
        f"检查点类型: {checkpoint.checkpoint_type}\n"
        f"检查点摘要: {checkpoint.summary or parent_node.output_summary or parent_node.input_summary}\n"
        f"原节点: {parent_node.label or parent_node.node_id}\n"
        f"原始用户任务: {rec.last_message or parent_node.input_summary}\n"
        f"检查点上下文: {json.dumps(snapshot, ensure_ascii=False)[:1800]}\n"
        "要求: 只继续该节点负责的部分；复用检查点前已完成的信息；如生成文件或报告，请说明交付物名称。"
    )
    if body.mode == "prompt_only":
        return {
            "ok": True,
            "mode": "prompt_only",
            "checkpoint": _workflow_checkpoint_to_dict(checkpoint),
            "continuation_message": prompt,
        }

    now = datetime.now(timezone.utc)
    child_wf = WorkflowRun(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=sid,
        template_id=parent_wf.template_id,
        status="queued",
        strategy=parent_wf.strategy,
        summary=f"从检查点恢复: {checkpoint.summary or checkpoint.checkpoint_type}",
        payload=json.dumps({
            "parent_workflow_run_id": parent_wf.id,
            "forked_from_checkpoint_id": checkpoint.id,
            "forked_from_node_run_id": parent_node.id,
            "resume_prompt": prompt[:1000],
        }, ensure_ascii=False),
        started_at=now,
        updated_at=now,
    )
    db.add(child_wf)
    db.flush()
    child_node = WorkflowNodeRun(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        workflow_run_id=child_wf.id,
        session_id=sid,
        employee_id=parent_node.employee_id,
        node_id=parent_node.node_id,
        label=parent_node.label,
        status="queued",
        input_summary=prompt[:4000],
        payload=json.dumps({
            "parent_node_run_id": parent_node.id,
            "checkpoint_id": checkpoint.id,
            "recovery": True,
        }, ensure_ascii=False),
    )
    db.add(child_node)
    db.flush()
    fork = WorkflowRunFork(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=sid,
        parent_workflow_run_id=parent_wf.id,
        child_workflow_run_id=child_wf.id,
        forked_from_node_run_id=parent_node.id,
        forked_from_checkpoint_id=checkpoint.id,
        reason="checkpoint_resume",
        payload_json=json.dumps({"prompt": prompt[:1000], "mode": body.mode}, ensure_ascii=False),
        created_by=p.user.id,
    )
    db.add(fork)
    _record_workflow_step(
        db,
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=sid,
        workflow_run_id=child_wf.id,
        workflow_node_run_id=child_node.id,
        employee_id=child_node.employee_id,
        event_type="replay.fork_created",
        title="创建恢复分支",
        summary=f"从检查点创建 Replay Fork: {checkpoint.summary or checkpoint.checkpoint_type}",
        input_summary=prompt[:1000],
        payload={
            "parent_workflow_run_id": parent_wf.id,
            "forked_from_node_run_id": parent_node.id,
            "forked_from_checkpoint_id": checkpoint.id,
        },
        checkpoint_type="replay.fork_created",
        hermes_session_id=rec.hermes_session_id,
        hermes_run_id=checkpoint.hermes_run_id,
    )
    rec.task_status = "running"
    rec.task_summary = "已创建恢复分支，正在从检查点继续执行。"
    db.add(CanvasEvent(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=sid,
        employee_id=child_node.employee_id,
        event_type="replay.fork_created",
        node_id=child_node.node_id,
        payload=json.dumps({
            "parent_workflow_run_id": parent_wf.id,
            "child_workflow_run_id": child_wf.id,
            "checkpoint_id": checkpoint.id,
        }, ensure_ascii=False),
    ))
    audit(db, principal=p, action="workflow_checkpoint.resume", resource_type="workflow_checkpoint",
          resource_id=checkpoint.id, request=request, extra={
              "session_id": sid,
              "parent_workflow_run_id": parent_wf.id,
              "child_workflow_run_id": child_wf.id,
          })
    db.commit()
    asyncio.create_task(_execute_checkpoint_resume(
        workflow_run_id=child_wf.id,
        workflow_node_run_id=child_node.id,
        checkpoint_id=checkpoint.id,
        prompt=prompt,
    ))
    return {
        "ok": True,
        "mode": "fork_resume",
        "task_status": rec.task_status,
        "checkpoint": _workflow_checkpoint_to_dict(checkpoint),
        "fork": _workflow_fork_to_dict(fork),
        "workflow_run": _workflow_run_to_dict(child_wf, [child_node]),
        "continuation_message": prompt,
    }


@app.post("/api/sessions/{sid}/workflow-steps/{step_event_id}/action")
async def workflow_step_action(
    sid: str,
    step_event_id: str,
    body: WorkflowStepActionIn,
    request: Request,
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    step = db.get(WorkflowStepEvent, step_event_id)
    if not step or step.tenant_id != p.tenant.id or step.session_id != sid:
        raise HTTPException(404, "workflow step not found")
    parent_wf = db.get(WorkflowRun, step.workflow_run_id)
    parent_node = db.get(WorkflowNodeRun, step.workflow_node_run_id) if step.workflow_node_run_id else None
    if not parent_wf or parent_wf.tenant_id != p.tenant.id:
        raise HTTPException(404, "workflow run not found")
    if not parent_node or not parent_node.employee_id:
        raise HTTPException(400, "workflow step is not attached to an executable employee node")

    checkpoint = db.query(WorkflowCheckpoint).filter(
        WorkflowCheckpoint.tenant_id == p.tenant.id,
        WorkflowCheckpoint.session_id == sid,
        WorkflowCheckpoint.step_event_id == step.id,
    ).order_by(WorkflowCheckpoint.created_at.desc()).first()
    payload = _json_loads_obj(step.payload_json, {})
    if not checkpoint:
        checkpoint = WorkflowCheckpoint(
            tenant_id=p.tenant.id,
            user_id=p.user.id,
            session_id=sid,
            workflow_run_id=parent_wf.id,
            workflow_node_run_id=parent_node.id,
            step_event_id=step.id,
            checkpoint_type=f"tool.{body.action}",
            status="available",
            summary=step.summary or step.title or step.event_type,
            context_snapshot_json=json.dumps({
                "event_type": step.event_type,
                "title": step.title,
                "summary": step.summary,
                "input_summary": step.input_summary,
                "output_summary": step.output_summary,
                "tool_name": step.tool_name,
                "payload": payload,
            }, ensure_ascii=False),
            hermes_session_id=rec.hermes_session_id,
            hermes_run_id=str(payload.get("hermes_run_id") or ""),
            upstream_node_outputs_json=json.dumps({}, ensure_ascii=False),
            artifact_policy_json=json.dumps({
                "artifact_ids": _json_loads_obj(step.artifact_ids, []),
                "on_resume": "create_new_version",
            }, ensure_ascii=False),
        )
        db.add(checkpoint)
        db.flush()

    action = body.action or "retry"
    verb = "重新执行该工具步骤" if action == "retry" else "跳过该工具步骤并继续后续任务"
    prompt = body.message.strip() or (
        "请从 OpenAtlas 协作工作流的具体工具步骤恢复执行。\n"
        f"恢复动作: {verb}\n"
        f"节点: {parent_node.label or parent_node.node_id}\n"
        f"工具: {step.tool_name or payload.get('tool_name') or payload.get('name') or '未知工具'}\n"
        f"步骤状态: {step.status}\n"
        f"步骤摘要: {step.summary or step.output_summary or step.input_summary or step.event_type}\n"
        f"原始用户任务: {rec.last_message or parent_node.input_summary}\n"
        f"工具上下文: {json.dumps(payload, ensure_ascii=False)[:1600]}\n"
        "要求: 复用该步骤之前已经完成的信息；不要重复已成功完成的上游工作；"
        + ("优先重新调用或替代该工具完成目标。" if action == "retry" else "明确说明该工具被跳过后的影响，并继续产出可用结果。")
    )

    now = datetime.now(timezone.utc)
    child_wf = WorkflowRun(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=sid,
        template_id=parent_wf.template_id,
        status="queued",
        strategy=parent_wf.strategy,
        summary=f"{'重试' if action == 'retry' else '跳过'}工具步骤: {step.tool_name or step.title or step.event_type}",
        payload=json.dumps({
            "parent_workflow_run_id": parent_wf.id,
            "forked_from_checkpoint_id": checkpoint.id,
            "forked_from_node_run_id": parent_node.id,
            "forked_from_step_event_id": step.id,
            "step_action": action,
            "resume_prompt": prompt[:1000],
        }, ensure_ascii=False),
        started_at=now,
        updated_at=now,
    )
    db.add(child_wf)
    db.flush()
    child_node = WorkflowNodeRun(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        workflow_run_id=child_wf.id,
        session_id=sid,
        employee_id=parent_node.employee_id,
        node_id=parent_node.node_id,
        label=parent_node.label,
        status="queued",
        input_summary=prompt[:4000],
        payload=json.dumps({
            "parent_node_run_id": parent_node.id,
            "checkpoint_id": checkpoint.id,
            "step_event_id": step.id,
            "step_action": action,
            "recovery": True,
        }, ensure_ascii=False),
    )
    db.add(child_node)
    db.flush()
    fork = WorkflowRunFork(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=sid,
        parent_workflow_run_id=parent_wf.id,
        child_workflow_run_id=child_wf.id,
        forked_from_node_run_id=parent_node.id,
        forked_from_checkpoint_id=checkpoint.id,
        reason=f"step_{action}",
        payload_json=json.dumps({"prompt": prompt[:1000], "step_event_id": step.id, "action": action}, ensure_ascii=False),
        created_by=p.user.id,
    )
    db.add(fork)
    _record_workflow_step(
        db,
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=sid,
        workflow_run_id=child_wf.id,
        workflow_node_run_id=child_node.id,
        employee_id=child_node.employee_id,
        event_type=f"tool.{action}_requested",
        title="工具步骤恢复",
        summary=f"{'重试' if action == 'retry' else '跳过并继续'}工具步骤: {step.tool_name or step.title or step.event_type}",
        input_summary=prompt[:1000],
        payload={
            "parent_workflow_run_id": parent_wf.id,
            "forked_from_node_run_id": parent_node.id,
            "forked_from_checkpoint_id": checkpoint.id,
            "forked_from_step_event_id": step.id,
            "step_action": action,
        },
        checkpoint_type=f"tool.{action}_requested",
        hermes_session_id=rec.hermes_session_id,
        hermes_run_id=checkpoint.hermes_run_id,
    )
    rec.task_status = "running"
    rec.task_summary = f"已创建工具步骤恢复分支，正在{'重试' if action == 'retry' else '跳过并继续'}。"
    if action == "skip":
        step.status = "skipped"
    db.add(CanvasEvent(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=sid,
        employee_id=child_node.employee_id,
        event_type=f"tool.{action}_requested",
        node_id=child_node.node_id,
        payload=json.dumps({
            "parent_workflow_run_id": parent_wf.id,
            "child_workflow_run_id": child_wf.id,
            "checkpoint_id": checkpoint.id,
            "step_event_id": step.id,
            "action": action,
        }, ensure_ascii=False),
    ))
    audit(db, principal=p, action=f"workflow_step.{action}", resource_type="workflow_step_event",
          resource_id=step.id, request=request, extra={
              "session_id": sid,
              "parent_workflow_run_id": parent_wf.id,
              "child_workflow_run_id": child_wf.id,
          })
    db.commit()
    asyncio.create_task(_execute_checkpoint_resume(
        workflow_run_id=child_wf.id,
        workflow_node_run_id=child_node.id,
        checkpoint_id=checkpoint.id,
        prompt=prompt,
    ))
    return {
        "ok": True,
        "action": action,
        "task_status": rec.task_status,
        "step": _workflow_step_to_dict(step),
        "checkpoint": _workflow_checkpoint_to_dict(checkpoint),
        "fork": _workflow_fork_to_dict(fork),
        "workflow_run": _workflow_run_to_dict(child_wf, [child_node]),
        "continuation_message": prompt,
    }


@app.post("/api/sessions/{sid}/recover")
async def recover_session(
    sid: str,
    request: Request,
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    target = await hermes_client.resolve_target(db, p.tenant.id)
    emp = db.get(DigitalEmployee, rec.employee_id) if rec.employee_id else None
    imported = await _reconcile_hermes_session_transcript(
        db,
        target=target,
        rec=rec,
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        employee_id=rec.employee_id,
        speaker_name=emp.display_name if emp else "",
    )
    snapshot = _session_summary(db, rec)
    artifacts = db.query(TaskArtifact).filter(
        TaskArtifact.session_id == rec.id,
        TaskArtifact.tenant_id == p.tenant.id,
        TaskArtifact.archived == False,  # noqa: E712
    ).order_by(TaskArtifact.created_at.desc()).limit(20).all()
    parts = [snapshot.get("summary") or "本会话还没有可总结的内容。"]
    if artifacts:
        parts.append("交付物: " + " / ".join(a.name for a in artifacts[:6]))
    rec.task_summary = "\n".join(parts)
    rec.summary_updated_at = datetime.now(timezone.utc)
    if rec.task_status == "running" and not _is_stale_running_session(rec):
        pass
    elif rec.task_status == "running" and not imported:
        rec.task_status = "needs_input"
    rec.updated_at = datetime.now(timezone.utc)
    audit(db, principal=p, action="session.recover", resource_type="session",
          resource_id=rec.id, request=request, extra={"imported": imported, "artifacts": len(artifacts)})
    db.commit()
    db.refresh(rec)
    return {
        "id": rec.id,
        "imported": imported,
        "task_status": rec.task_status,
        "task_summary": rec.task_summary,
        "summary_updated_at": rec.summary_updated_at.isoformat() if rec.summary_updated_at else None,
        "artifacts": [_artifact_to_dict(a, db) for a in artifacts],
        "health": _session_health_snapshot(db, rec),
    }


@app.get("/api/sessions/{sid}/canvas-state")
def get_canvas_state(
    sid: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    if not rec.canvas_state:
        return {"nodes": [], "edges": [], "version": 1}
    try:
        return json.loads(rec.canvas_state)
    except Exception:
        return {"nodes": [], "edges": [], "version": 1}


@app.api_route("/api/sessions/{sid}/canvas-state", methods=["PATCH", "POST"])
def patch_canvas_state(
    sid: str, body: CanvasStateIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    state = {"nodes": body.nodes, "edges": body.edges, "version": body.version}
    rec.canvas_state = json.dumps(state, ensure_ascii=False)
    audit(db, principal=p, action="session.canvas.update", resource_type="session",
          resource_id=rec.id, request=request,
          extra={"nodes": len(body.nodes), "edges": len(body.edges), "version": body.version})
    db.commit()
    return {"ok": True, "canvas_state": state}


@app.post("/api/sessions/{sid}/canvas-events")
def create_canvas_event(
    sid: str, body: CanvasEventIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    if body.employee_id:
        emp = db.get(DigitalEmployee, body.employee_id)
        if not emp or emp.tenant_id != p.tenant.id:
            raise HTTPException(400, "employee not in this tenant")
    ev = CanvasEvent(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=rec.id,
        employee_id=body.employee_id,
        event_type=body.event_type.strip() or "canvas.event",
        node_id=body.node_id.strip(),
        edge_id=body.edge_id.strip(),
        payload=json.dumps(body.payload or {}, ensure_ascii=False),
    )
    db.add(ev)
    audit(db, principal=p, action="session.canvas.event", resource_type="session",
          resource_id=rec.id, request=request,
          extra={"event_type": ev.event_type, "employee_id": ev.employee_id, "node_id": ev.node_id})
    db.commit()
    return _canvas_event_to_dict(ev)


@app.get("/api/canvas-events")
def list_canvas_events(
    employee_id: str | None = Query(default=None),
    session_id: str | None = Query(default=None),
    limit: int = Query(default=100, ge=1, le=500),
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    q = db.query(CanvasEvent).filter(CanvasEvent.tenant_id == p.tenant.id, CanvasEvent.user_id == p.user.id)
    if employee_id:
        q = q.filter(CanvasEvent.employee_id == employee_id)
    if session_id:
        q = q.filter(CanvasEvent.session_id == session_id)
    rows = q.order_by(CanvasEvent.created_at.desc()).limit(limit).all()
    return {"items": [_canvas_event_to_dict(r) for r in rows], "total": len(rows)}


@app.post("/api/sessions/{sid}/save-template")
def save_session_template(
    sid: str, body: TemplateSaveIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    canvas_state = {}
    if rec.canvas_state:
        try:
            canvas_state = json.loads(rec.canvas_state)
        except Exception:
            canvas_state = {}
    if not isinstance(canvas_state, dict):
        canvas_state = {}
    canvas_state.setdefault("nodes", [])
    canvas_state.setdefault("edges", [])
    canvas_state["meta"] = {
        **(canvas_state.get("meta") if isinstance(canvas_state.get("meta"), dict) else {}),
        "strategy": (body.strategy or "relay").strip() or "relay",
        "failure_strategy": (body.failure_strategy or "continue_with_next_employee").strip() or "continue_with_next_employee",
        "default_prompt": body.default_prompt.strip(),
        "output_type": (body.output_type or "markdown").strip() or "markdown",
    }
    tpl = CollaborationTemplate(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        name=body.name.strip(),
        description=body.description.strip(),
        category=(body.category or "general").strip() or "general",
        visibility=(body.visibility or "private").strip() or "private",
        primary_employee_id=rec.employee_id,
        participant_ids=rec.participant_ids or "[]",
        canvas_state=json.dumps(canvas_state, ensure_ascii=False),
        source_session_id=rec.id,
    )
    db.add(tpl)
    db.flush()
    rec.reusable_template_id = tpl.id
    audit(db, principal=p, action="collaboration_template.create", resource_type="collaboration_template",
          resource_id=tpl.id, request=request,
          extra={"source_session_id": rec.id, "participant_ids": json.loads(rec.participant_ids or "[]")})
    db.commit()
    return _template_to_dict(tpl)


@app.get("/api/collaboration-templates")
def list_collaboration_templates(
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rows = db.query(CollaborationTemplate).filter(
        CollaborationTemplate.tenant_id == p.tenant.id,
        or_(CollaborationTemplate.user_id == p.user.id, CollaborationTemplate.visibility == "tenant"),
    ).order_by(CollaborationTemplate.updated_at.desc()).all()
    return {"items": [_template_to_dict(t) for t in rows]}


@app.delete("/api/collaboration-templates/{template_id}")
def delete_collaboration_template(
    template_id: str,
    request: Request,
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    tpl = db.get(CollaborationTemplate, template_id)
    if not tpl or tpl.tenant_id != p.tenant.id:
        raise HTTPException(404, "template not found")
    is_admin = p.user.role in (UserRole.tenant_admin, UserRole.system_admin)
    if tpl.user_id != p.user.id and not is_admin:
        raise HTTPException(403, "only the owner or tenant admin can delete this template")
    db.delete(tpl)
    audit(
        db,
        principal=p,
        action="collaboration_template.delete",
        resource_type="collaboration_template",
        resource_id=template_id,
        request=request,
        extra={"name": tpl.name, "category": tpl.category},
    )
    db.commit()
    return {"ok": True, "deleted": template_id}


@app.post("/api/collaboration-templates/{template_id}/sessions")
async def create_session_from_template(
    template_id: str, body: TemplateUseIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    tpl = db.get(CollaborationTemplate, template_id)
    if not tpl or tpl.tenant_id != p.tenant.id or (tpl.user_id != p.user.id and tpl.visibility != "tenant"):
        raise HTTPException(404, "template not found")
    participant_ids = json.loads(tpl.participant_ids or "[]")
    target = await hermes_client.resolve_target(db, p.tenant.id)
    try:
        unique_title = (body.title or tpl.name or "template session") + "-" + _uuid.uuid4().hex[:6]
        hermes = await hermes_client.create_session(target, title=unique_title)
    except Exception as ex:
        raise HTTPException(502, f"hermes create_session failed: {ex}")
    inner = hermes.get("session") or hermes.get("data") or {}
    hermes_sid = hermes.get("id") or hermes.get("session_id") or inner.get("id") or inner.get("session_id")
    if not hermes_sid:
        raise HTTPException(502, f"hermes did not return session id: {hermes}")
    rec = SessionRecord(
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        employee_id=tpl.primary_employee_id,
        hermes_session_id=hermes_sid,
        title=body.title or tpl.name,
        participant_ids=json.dumps(participant_ids, ensure_ascii=False),
        canvas_state=tpl.canvas_state or "",
        reusable_template_id=tpl.id,
    )
    db.add(rec)
    audit(db, principal=p, action="collaboration_template.use", resource_type="session",
          resource_id=rec.id, request=request,
          extra={"template_id": tpl.id, "participant_ids": participant_ids})
    db.commit()
    db.refresh(rec)
    return {
        "id": rec.id,
        "employee_id": rec.employee_id,
        "hermes_session_id": rec.hermes_session_id,
        "participant_ids": participant_ids,
        "is_group": bool(participant_ids),
        "title": rec.title,
        "last_message": rec.last_message,
        "message_count": rec.message_count,
        "created_at": rec.created_at.isoformat(),
        "updated_at": rec.updated_at.isoformat(),
        "reusable_template_id": rec.reusable_template_id,
    }


@app.get("/api/sessions/{sid}/messages")
async def session_messages(
    sid: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    last_msg = db.query(MessageRecord).filter_by(session_id=sid).order_by(MessageRecord.created_at.desc()).first()
    if (rec.task_status == "running" or _is_stale_running_session(rec) or (last_msg and last_msg.role == "user")):
        target = await hermes_client.resolve_target(db, p.tenant.id)
        emp = db.get(DigitalEmployee, rec.employee_id) if rec.employee_id else None
        await _reconcile_hermes_session_transcript(
            db,
            target=target,
            rec=rec,
            tenant_id=p.tenant.id,
            user_id=p.user.id,
            employee_id=rec.employee_id,
            speaker_name=emp.display_name if emp else "",
        )
    # P3.12 (2026-06-07) Bug 5b / 3.4.5: 默认返 OpenAtlas 侧 sanitized messages
    # (没注入 system_prompt / context / file_context 的真实 display_message).
    # 老 session (P3.11 之前) 没 MessageRecord, fallback 拿 Hermes raw + 清洗.
    rows = db.query(MessageRecord).filter_by(session_id=sid).order_by(MessageRecord.created_at.asc()).all()
    if rows:
        def _message_reasoning(raw: str | None) -> list[str]:
            if not raw:
                return []
            try:
                parsed = json.loads(raw)
                if isinstance(parsed, list):
                    return [str(x) for x in parsed if str(x).strip()]
                if isinstance(parsed, str) and parsed.strip():
                    return [parsed]
            except Exception:
                pass
            return [raw] if raw.strip() else []

        return {
            "items": [
                {
                    "id": r.id,
                    "role": r.role,
                    "content": r.content,
                    "tool_calls": json.loads(r.tool_calls or "[]"),
                    "reasoning": _message_reasoning(getattr(r, "reasoning", "") or ""),
                    "attachments": json.loads(getattr(r, "attachments", "") or "[]"),
                    "input_tokens": int(getattr(r, "input_tokens", 0) or 0),
                    "output_tokens": int(getattr(r, "output_tokens", 0) or 0),
                    "total_tokens": int(getattr(r, "total_tokens", 0) or 0),
                    "token_count": int(
                        getattr(r, "total_tokens", 0)
                        or getattr(r, "output_tokens", 0)
                        or getattr(r, "input_tokens", 0)
                        or _rough_token_count(r.content or "")
                    ),
                    "speaker_employee_id": r.speaker_employee_id,
                    "speaker_name": r.speaker_name,
                    "turn_index": r.turn_index,
                    "created_at": r.created_at.isoformat() if r.created_at else None,
                }
                for r in rows
            ],
            "source": "openatlas",
        }
    # fallback: Hermes raw + 清洗 (剥 <system_prompt> <context> <file_context>)
    # P3.12 late: hermes session 可能已被回收 (404) → 返空 items, 不抛 500
    target = await hermes_client.resolve_target(db, p.tenant.id)
    try:
        raw = await hermes_client.get_session_messages(target, rec.hermes_session_id)
    except Exception as e:  # noqa: BLE001
        # hermes session 失效 (404/410/500) → 老 session 当空返, 不阻塞 history 列表
        return {"items": [], "source": "hermes_missing", "_error": str(e)[:200]}
    items = (raw.get("items") or raw.get("data") or []) if isinstance(raw, dict) else raw
    sanitized = []
    for it in items:
        content = it.get("content") or it.get("text") or ""
        clean = _sanitize_hermes_text(content)
        if not clean:
            continue
        sanitized.append({
            **it,
            "content": clean,
            "_sanitized": True,  # 标记已清洗, 前端可显示
        })
    return {"items": sanitized, "source": "hermes_sanitized"}


@app.post("/api/sessions/{sid}/summary")
def refresh_session_summary(
    sid: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    snapshot = _session_summary(db, rec)
    artifact_rows = db.query(TaskArtifact).filter(
        TaskArtifact.session_id == sid,
        TaskArtifact.tenant_id == p.tenant.id,
        TaskArtifact.archived == False,  # noqa: E712
    ).order_by(TaskArtifact.created_at.desc()).limit(20).all()
    context_rows = db.query(ContextInjection).filter(
        ContextInjection.session_id == sid,
        ContextInjection.tenant_id == p.tenant.id,
    ).order_by(ContextInjection.created_at.desc()).limit(40).all()
    context_counts: dict[str, int] = {}
    for c in context_rows:
        context_counts[c.kind] = context_counts.get(c.kind, 0) + 1
    parts = [
        snapshot.get("summary") or "本会话还没有可总结的内容。",
        f"任务状态: {getattr(rec, 'task_status', 'draft') or 'draft'}",
    ]
    if artifact_rows:
        parts.append("交付物: " + " / ".join(a.name for a in artifact_rows[:6]))
    if context_counts:
        parts.append("本次上下文: " + " / ".join(f"{k} {v}" for k, v in sorted(context_counts.items())))
    rec.task_summary = "\n".join(parts)
    rec.summary_updated_at = datetime.now(timezone.utc)
    audit(db, principal=p, action="session.summary.refresh", resource_type="session",
          resource_id=rec.id, request=request, extra={"artifacts": len(artifact_rows), "context": context_counts})
    db.commit()
    return {
        "task_summary": rec.task_summary,
        "task_status": rec.task_status,
        "summary_updated_at": rec.summary_updated_at.isoformat() if rec.summary_updated_at else None,
        "summary": snapshot,
    }


@app.patch("/api/sessions/{sid}/task-status")
def patch_session_task_status(
    sid: str, body: TaskStatusPatchIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    allowed = {"draft", "running", "needs_input", "completed", "failed"}
    if body.task_status not in allowed:
        raise HTTPException(400, f"task_status must be one of {sorted(allowed)}")
    rec.task_status = body.task_status
    rec.updated_at = datetime.now(timezone.utc)
    audit(db, principal=p, action="session.status.update", resource_type="session",
          resource_id=rec.id, request=request, extra={"task_status": body.task_status})
    db.commit()
    return {
        "id": rec.id,
        "task_status": rec.task_status,
        "updated_at": rec.updated_at.isoformat() if rec.updated_at else None,
    }


@app.get("/api/sessions/{sid}/context")
def list_session_context(
    sid: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    rows = db.query(ContextInjection).filter(
        ContextInjection.session_id == sid,
        ContextInjection.tenant_id == p.tenant.id,
    ).order_by(ContextInjection.created_at.desc()).limit(200).all()
    return {"items": [_context_to_dict(r) for r in rows]}


@app.get("/api/sessions/{sid}/artifacts")
def list_session_artifacts(
    sid: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    rows = db.query(TaskArtifact).filter(
        TaskArtifact.session_id == sid,
        TaskArtifact.tenant_id == p.tenant.id,
    ).order_by(TaskArtifact.created_at.desc()).all()
    return {"items": [_artifact_to_dict(r, db) for r in rows]}


@app.patch("/api/artifacts/{artifact_id}")
def patch_artifact(
    artifact_id: str,
    body: ArtifactPatchIn,
    request: Request,
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    art = db.get(TaskArtifact, artifact_id)
    if not art or art.tenant_id != p.tenant.id or art.user_id != p.user.id:
        raise HTTPException(404, "artifact not found")
    changes: dict[str, Any] = {}
    if body.name is not None:
        name = body.name.strip()
        if not name:
            raise HTTPException(400, "name cannot be empty")
        art.name = name[:255]
        art.version = _next_artifact_version(db, session_id=art.session_id, name=art.name, exclude_id=art.id)
        changes["name"] = art.name
        changes["version"] = art.version
    if body.status is not None:
        status = body.status.strip()
        allowed = {"active", "draft", "final", "archived"}
        if status not in allowed:
            raise HTTPException(400, f"status must be one of {sorted(allowed)}")
        art.status = status
        art.archived = status == "archived"
        changes["status"] = art.status
        changes["archived"] = art.archived
    if not changes:
        raise HTTPException(400, "no fields to update")
    audit(db, principal=p, action="artifact.update", resource_type="artifact",
          resource_id=art.id, request=request, extra={"session_id": art.session_id, **changes})
    db.commit()
    db.refresh(art)
    return _artifact_to_dict(art, db)


@app.post("/api/artifacts/{artifact_id}/archive")
def archive_artifact(
    artifact_id: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    art = db.get(TaskArtifact, artifact_id)
    if not art or art.tenant_id != p.tenant.id or art.user_id != p.user.id:
        raise HTTPException(404, "artifact not found")
    art.archived = True
    art.status = "archived"
    audit(db, principal=p, action="artifact.archive", resource_type="artifact",
          resource_id=art.id, request=request, extra={"session_id": art.session_id, "kind": art.kind})
    db.commit()
    return _artifact_to_dict(art, db)


@app.post("/api/hermes-runs/{run_id}/approval")
async def approve_hermes_run(
    run_id: str, body: RunApprovalIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    """Resolve a pending approval on the tenant-scoped Hermes runtime."""
    run_meta = _get_authorized_hermes_run(run_id, p)
    target = await hermes_client.resolve_target(db, p.tenant.id)
    choice = {"approve": "once", "allow": "once", "approved": "once"}.get(body.choice, body.choice)
    out = await hermes_client.respond_run_approval(
        target,
        run_id,
        choice=choice,
        resolve_all=body.resolve_all,
        approval_id=body.approval_id,
    )
    audit(db, principal=p, action="hermes_run.approval", resource_type="hermes_run",
          resource_id=run_id, request=request, extra={
              "choice": choice,
              "resolve_all": body.resolve_all,
              "approval_id": body.approval_id,
              "session_id": run_meta.get("session_id"),
              "employee_id": run_meta.get("employee_id"),
              "hermes_session_id": run_meta.get("hermes_session_id"),
          })
    return out


@app.post("/api/hermes-runs/{run_id}/stop")
async def stop_hermes_run(
    run_id: str, body: RunStopIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    """Stop a tenant-scoped Hermes run instead of only aborting the browser stream."""
    run_meta = _get_authorized_hermes_run(run_id, p)
    target = await hermes_client.resolve_target(db, p.tenant.id)
    out = await hermes_client.stop_run(target, run_id)
    audit(db, principal=p, action="hermes_run.stop", resource_type="hermes_run",
          resource_id=run_id, request=request, extra={
              "reason": body.reason,
              "session_id": run_meta.get("session_id"),
              "employee_id": run_meta.get("employee_id"),
              "hermes_session_id": run_meta.get("hermes_session_id"),
          })
    db.commit()
    return out


@app.post("/api/sessions/{sid}/chat/stream")
async def session_chat_stream(
    sid: str, body: ChatIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> StreamingResponse:
    rec = db.get(SessionRecord, sid)
    if not rec or rec.tenant_id != p.tenant.id or rec.user_id != p.user.id or rec.archived:
        raise HTTPException(404, "session not found")
    stored_participants = json.loads(rec.participant_ids or "[]")
    requested_relays = body.relay_employee_ids or []
    primary_employee_id = body.primary_employee_id or rec.employee_id
    turn_employee_ids: list[str] = []
    for eid in [primary_employee_id, *stored_participants, *requested_relays]:
        if eid and eid not in turn_employee_ids:
            turn_employee_ids.append(eid)
    if not turn_employee_ids and rec.employee_id:
        turn_employee_ids.append(rec.employee_id)

    turn_employees: list[DigitalEmployee] = []
    for eid in turn_employee_ids:
        emp = db.get(DigitalEmployee, eid)
        if not emp or emp.tenant_id != p.tenant.id or emp.status != EmployeeStatus.active:
            raise HTTPException(404, f"employee {eid} not found in this tenant")
        turn_employees.append(emp)
    if not turn_employees:
        raise HTTPException(400, "session has no employee to chat with")

    # P3.12 (2026-06-07) Bug 7: 附件 file_context 注入
    # 校验 attachment_ids 全部属于本 user / 本 session
    file_assets: list[FileAsset] = []
    if body.attachment_ids:
        for fid in body.attachment_ids:
            fa = db.get(FileAsset, fid)
            if not fa:
                raise HTTPException(400, f"attachment {fid} not found")
            if fa.tenant_id != p.tenant.id or fa.user_id != p.user.id:
                raise HTTPException(403, f"attachment {fid} not in this tenant/user")
            # 若 attachment 上传时没绑 session, 帮它绑上
            if not fa.session_id and rec.id:
                fa.session_id = rec.id
            if not fa.employee_id and primary_employee_id:
                fa.employee_id = primary_employee_id
            file_assets.append(fa)
        db.commit()
    file_ctx_block = _build_file_context_block(file_assets)
    file_provenance = []
    for f in file_assets:
        meta = _file_to_dict(f)
        file_provenance.append({
            "id": meta.get("id"),
            "name": meta.get("original_name"),
            "mime": meta.get("mime_type"),
            "size": meta.get("size"),
            "status": meta.get("status"),
            "extracted_chars": meta.get("extracted_chars", 0),
            "summary": meta.get("summary"),
            "snippets": (meta.get("snippets") or [])[:3],
            "expires_at": meta.get("expires_at"),
            "is_expired": meta.get("is_expired"),
        })
    attachment_ids_injected = [f["id"] for f in file_provenance if f.get("id")]
    attachment_chars = sum(int(f.get("extracted_chars") or 0) for f in file_provenance)
    hermes_home = _tenant_hermes_home(db, p.tenant.id)
    recent_context_block = _build_recent_conversation_block(
        db,
        session_id=rec.id,
        current_user_message=body.message,
    )

    # Update last_message + count + audit, all in one commit
    rec.last_message = body.message[:200]
    rec.message_count = (rec.message_count or 0) + 1
    rec.task_status = "running"
    primary_system_message, primary_eff, primary_skills = _build_employee_system_message(
        db,
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        employee_id=turn_employees[0].id,
        hermes_home=hermes_home,
        file_ctx_block=file_ctx_block,
        recent_context_block=recent_context_block,
    )
    audit(db, principal=p, action="session.chat", resource_type="session",
          resource_id=rec.id, request=request,
          extra={"memories_injected": [m["id"] for m in primary_eff],
                 "speaker_employee_ids": [e.id for e in turn_employees],
                 "attachments_injected": attachment_ids_injected,
                 "attachment_chars": attachment_chars})
    # P3.12 (2026-06-07) Bug 5b: 落 MessageRecord, 只存 display_message (用户真实输入)
    user_msg = MessageRecord(
        session_id=rec.id,
        role="user",
        content=body.message,
        model_message="",  # model_message 在 event_gen 后填, 异步 commit
        attachments=json.dumps(file_provenance, ensure_ascii=False),
    )
    db.add(user_msg)
    db.commit()
    db.refresh(user_msg)
    db.refresh(rec)
    # Resolve tenant-scoped runtime target (Phase 2: per-tenant Hermes gateway)
    target = await hermes_client.resolve_target(db, p.tenant.id)
    use_run_events = await hermes_client.supports_run_events(target)
    msg = body.message
    workflow_row = _ensure_workflow_run(
        db,
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=rec.id,
        template_id=rec.reusable_template_id,
        strategy="relay" if len(turn_employees) > 1 else "single",
        payload={
            "message_id": user_msg.id,
            "employee_ids": [e.id for e in turn_employees],
            "query_excerpt": body.message[:240],
        },
    )
    _update_workflow_node_run(
        db,
        workflow_run=workflow_row,
        tenant_id=p.tenant.id,
        user_id=p.user.id,
        session_id=rec.id,
        employee_id=None,
        node_id="user",
        label="用户",
        status="done",
        event_type="user.message",
        input_summary=body.message[:1000],
    )
    workflow_run_id = workflow_row.id
    db.commit()

    async def event_gen():
        # P3.12 (2026-06-07) Bug 3.4.3: 每个 SSE event 附 openatlas_session_id
        # 让前端做归属校验 (切会话后旧流事件被忽略)
        sess_marker = rec.id

        def trace_event(stage: str, title: str, detail: str = "", speaker: dict | None = None, **extra: object) -> str:
            payload = {
                "openatlas_session_id": sess_marker,
                "stage": stage,
                "title": title,
                "detail": detail,
                **extra,
            }
            if speaker:
                payload.update(speaker)
            return f"event: openatlas.trace\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"

        def task_state_event(status: str, reason: str = "", speaker: dict | None = None) -> str:
            payload = {
                "openatlas_session_id": sess_marker,
                "session_id": sess_marker,
                "task_status": status,
                "reason": reason,
            }
            if speaker:
                payload.update(speaker)
            return f"event: openatlas.task_state\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"

        # 注入文件 provenance (不返 content, 只返 metadata)
        yield task_state_event("running", "chat stream started")
        yield trace_event(
            "files",
            "附件上下文准备",
            f"已准备 {len(file_provenance)} 个附件，合计可注入 {attachment_chars} 字",
            file_count=len(file_provenance),
            extracted_chars=attachment_chars,
        )
        yield f"event: openatlas.files\ndata: {json.dumps({'items': file_provenance, 'openatlas_session_id': sess_marker}, ensure_ascii=False)}\n\n"

        try:
            relay_outputs: list[str] = []
            for turn_index, emp in enumerate(turn_employees):
                session_run_id = ""
                current_hermes_run_id = ""
                current_node_run_id = ""
                db_run = SessionLocal()
                try:
                    wf = db_run.get(WorkflowRun, workflow_run_id)
                    run_row = _create_session_run(
                        db_run,
                        tenant_id=p.tenant.id,
                        user_id=p.user.id,
                        session_id=rec.id,
                        employee_id=emp.id,
                        status="running",
                        stage="agent",
                        reason=f"{emp.display_name} 开始处理",
                        payload={"turn_index": turn_index, "query_excerpt": msg[:240]},
                    )
                    session_run_id = run_row.id
                    node_row = _update_workflow_node_run(
                        db_run,
                        workflow_run=wf,
                        tenant_id=p.tenant.id,
                        user_id=p.user.id,
                        session_id=rec.id,
                        employee_id=emp.id,
                        label=emp.display_name,
                        status="running",
                        event_type="agent.join",
                        session_run_id=session_run_id,
                        input_summary=msg[:1000],
                        payload_patch={"turn_index": turn_index},
                    )
                    current_node_run_id = node_row.id if node_row else ""
                    _record_workflow_step(
                        db_run,
                        tenant_id=p.tenant.id,
                        user_id=p.user.id,
                        session_id=rec.id,
                        workflow_run_id=workflow_run_id,
                        workflow_node_run_id=current_node_run_id or None,
                        employee_id=emp.id,
                        event_type="node.started",
                        title=f"{emp.display_name} 开始执行",
                        summary=f"{emp.display_name} 开始处理本轮任务",
                        input_summary=msg[:1000],
                        payload={"turn_index": turn_index, "employee_name": emp.display_name},
                        checkpoint_type="node.started",
                        hermes_session_id=rec.hermes_session_id,
                    )
                    session_row = db_run.get(SessionRecord, rec.id)
                    if session_row:
                        session_row.task_status = "running"
                        session_row.updated_at = datetime.now(timezone.utc)
                    db_run.commit()
                finally:
                    db_run.close()

                def mark_run_progress(
                    *,
                    status: str | None = None,
                    node_status: str | None = None,
                    stage: str | None = None,
                    reason: str | None = None,
                    event_type: str | None = None,
                    hermes_run_id: str | None = None,
                    artifact_ids: list[str] | None = None,
                    output_summary: str | None = None,
                    error: str | None = None,
                    payload_patch: dict[str, Any] | None = None,
                ) -> None:
                    db_progress = SessionLocal()
                    try:
                        wf = db_progress.get(WorkflowRun, workflow_run_id)
                        _update_session_run(
                            db_progress,
                            session_run_id,
                            status=status,
                            stage=stage,
                            reason=reason,
                            event_type=event_type,
                            hermes_run_id=hermes_run_id,
                            payload_patch=payload_patch,
                        )
                        _update_workflow_node_run(
                            db_progress,
                            workflow_run=wf,
                            tenant_id=p.tenant.id,
                            user_id=p.user.id,
                            session_id=rec.id,
                            employee_id=emp.id,
                            label=emp.display_name,
                            status=node_status or status,
                            event_type=event_type,
                            session_run_id=session_run_id,
                            hermes_run_id=hermes_run_id,
                            output_summary=output_summary,
                            artifact_ids=artifact_ids,
                            error=error,
                            payload_patch=payload_patch,
                        )
                        if wf and status in {"completed", "failed", "cancelled"}:
                            wf.status = "failed" if status == "failed" else "running"
                            wf.updated_at = datetime.now(timezone.utc)
                        db_progress.commit()
                    except Exception:
                        db_progress.rollback()
                    finally:
                        db_progress.close()

                def record_step_event(
                    event_type: str,
                    payload: dict[str, Any] | None = None,
                    *,
                    title: str = "",
                    summary: str = "",
                    input_summary: str = "",
                    output_summary: str = "",
                    risk_level: str = "low",
                    tool_name: str = "",
                    artifact_ids: list[str] | None = None,
                    file_ids: list[str] | None = None,
                    checkpoint_type: str | None = None,
                ) -> None:
                    _record_workflow_step_safely(
                        tenant_id=p.tenant.id,
                        user_id=p.user.id,
                        session_id=rec.id,
                        workflow_run_id=workflow_run_id,
                        workflow_node_run_id=current_node_run_id or None,
                        employee_id=emp.id,
                        event_type=event_type,
                        payload=payload or {},
                        title=title,
                        summary=summary,
                        input_summary=input_summary,
                        output_summary=output_summary,
                        risk_level=risk_level,
                        tool_name=tool_name,
                        artifact_ids=artifact_ids,
                        file_ids=file_ids,
                        checkpoint_type=checkpoint_type,
                        hermes_session_id=rec.hermes_session_id,
                        hermes_run_id=current_hermes_run_id,
                    )

                speaker = {
                    "openatlas_session_id": sess_marker,
                    "speaker_employee_id": emp.id,
                    "speaker_name": emp.display_name,
                    "turn_index": turn_index,
                    "session_run_id": session_run_id,
                    "workflow_run_id": workflow_run_id,
                }
                yield trace_event(
                    "agent",
                    "调度数字员工",
                    f"{emp.display_name} 开始处理本轮任务",
                    speaker,
                    employee_id=emp.id,
                )
                _record_canvas_runtime_event(
                    tenant_id=p.tenant.id,
                    user_id=p.user.id,
                    session_id=rec.id,
                    employee_id=emp.id,
                    event_type="agent.join",
                    payload={"speaker_name": emp.display_name, "turn_index": turn_index},
                )
                yield f"event: agent_join\ndata: {json.dumps(speaker, ensure_ascii=False)}\n\n"

                relay_context = "\n\n".join(relay_outputs)
                if turn_index == 0:
                    system_message, eff, skill_evidence = primary_system_message, primary_eff, primary_skills
                else:
                    db_ctx = SessionLocal()
                    try:
                        system_message, eff, skill_evidence = _build_employee_system_message(
                            db_ctx,
                            tenant_id=p.tenant.id,
                            user_id=p.user.id,
                            employee_id=emp.id,
                            hermes_home=hermes_home,
                            file_ctx_block=file_ctx_block,
                            recent_context_block=recent_context_block,
                            relay_context=relay_context,
                        )
                    finally:
                        db_ctx.close()
                speaker_role_block = (
                    "<current_speaker>\n"
                    f"You are now speaking as OpenAtlas digital employee: {emp.display_name}.\n"
                    f"speaker_employee_id: {emp.id}\n"
                    f"turn_index: {turn_index}\n"
                    "In this group chat turn, answer only for this current employee. "
                    "If the user assigned different instructions to the primary employee, relay employee, or other participants, "
                    "follow only the instruction that matches this current speaker and do not repeat previous employees' answers.\n"
                    "</current_speaker>"
                )
                system_message = "\n\n".join([p for p in (system_message, speaker_role_block) if p and p.strip()])
                context_payload = {
                    **speaker,
                    "skills": skill_evidence,
                    "memories": eff,
                    "files": file_provenance,
                }
                db_ctx = SessionLocal()
                try:
                    _record_context_injections(
                        db_ctx,
                        tenant_id=p.tenant.id,
                        user_id=p.user.id,
                        session_id=rec.id,
                        message_id=user_msg.id,
                        employee_id=emp.id,
                        turn_index=turn_index,
                        skills=skill_evidence,
                        memories=eff,
                        files=file_provenance,
                    )
                    for skill_item in skill_evidence:
                        db_ctx.add(SkillRun(
                            tenant_id=p.tenant.id,
                            user_id=p.user.id,
                            session_id=rec.id,
                            employee_id=emp.id,
                            skill_id=skill_item.get("id") or None,
                            skill_name=str(skill_item.get("name") or ""),
                            skill_slug=str(skill_item.get("slug") or ""),
                            status=("used" if skill_item.get("status") == "injected" else str(skill_item.get("status") or "used")),
                            error="" if skill_item.get("status") == "injected" else str(skill_item.get("summary") or ""),
                        ))
                    _record_workflow_step(
                        db_ctx,
                        tenant_id=p.tenant.id,
                        user_id=p.user.id,
                        session_id=rec.id,
                        workflow_run_id=workflow_run_id,
                        workflow_node_run_id=current_node_run_id or None,
                        employee_id=emp.id,
                        event_type="context.injected",
                        title="注入本轮上下文",
                        summary=f"Skill {len(skill_evidence)} 个，记忆 {len(eff)} 条，附件 {len(file_provenance)} 个",
                        payload={
                            "skills": skill_evidence,
                            "memories": eff,
                            "files": file_provenance,
                            "turn_index": turn_index,
                        },
                        file_ids=[str(f.get("id")) for f in file_provenance if f.get("id")],
                        checkpoint_type="context.injected",
                        hermes_session_id=rec.hermes_session_id,
                        hermes_run_id=current_hermes_run_id,
                    )
                    db_ctx.commit()
                finally:
                    db_ctx.close()
                yield trace_event(
                    "context",
                    "注入本轮上下文",
                    f"Skill {len(skill_evidence)} 个，记忆 {len(eff)} 条，附件 {len(file_provenance)} 个",
                    speaker,
                    skill_count=len(skill_evidence),
                    memory_count=len(eff),
                    file_count=len(file_provenance),
                )
                yield f"event: openatlas.context\ndata: {json.dumps(context_payload, ensure_ascii=False)}\n\n"
                yield f"event: openatlas.memories\ndata: {json.dumps({**speaker, 'items': eff}, ensure_ascii=False)}\n\n"

                assistant_text_parts: list[str] = []
                reasoning_parts: list[str] = []
                tool_call_records: list[dict[str, Any]] = []
                persisted_artifact_ids: list[str] = []
                final_task_status = "running"
                final_task_reason = ""
                assistant_persisted = False
                assistant_usage = {"input_tokens": 0, "output_tokens": 0, "total_tokens": 0}
                runtime_detached = False
                scoped_msg = msg
                if len(turn_employees) > 1:
                    other_names = [x.display_name for x in turn_employees if x.id != emp.id]
                    scoped_msg = (
                        f"{msg}\n\n"
                        "<current_turn_instruction>\n"
                        f"本轮只允许「{emp.display_name}」发言。请只完成用户任务中分配给「{emp.display_name}」的部分，"
                        "不要替其他员工回答、不要输出其他员工的小节、不要生成最终汇总。\n"
                        f"其他员工将单独发言: {', '.join(other_names) or '无'}。\n"
                        "如果用户的原始任务包含多个角色，请忽略不属于当前员工的角色要求。\n"
                        "</current_turn_instruction>"
                    )

                def persist_assistant_once() -> None:
                    nonlocal assistant_persisted, final_task_status, final_task_reason, persisted_artifact_ids
                    if assistant_persisted:
                        return
                    assistant_persisted = True
                    full_assistant = "".join(assistant_text_parts)
                    if not full_assistant.strip():
                        final_task_status = "needs_input"
                        final_task_reason = "assistant returned no visible content"
                        db2 = SessionLocal()
                        try:
                            session_row = db2.get(SessionRecord, sid)
                            if session_row:
                                session_row.task_status = final_task_status
                                session_row.updated_at = datetime.now(timezone.utc)
                            db2.commit()
                        finally:
                            db2.close()
                        return
                    final_task_status, final_task_reason = _infer_task_status_from_assistant(full_assistant)
                    db2 = SessionLocal()
                    try:
                        msg_row = MessageRecord(
                            session_id=sid,
                            role="assistant",
                            content=full_assistant,
                            reasoning=json.dumps(reasoning_parts, ensure_ascii=False),
                            tool_calls=json.dumps(tool_call_records, ensure_ascii=False),
                            input_tokens=int(assistant_usage.get("input_tokens") or 0),
                            output_tokens=int(assistant_usage.get("output_tokens") or 0),
                            total_tokens=int(assistant_usage.get("total_tokens") or 0),
                            speaker_employee_id=emp.id,
                            speaker_name=emp.display_name,
                            turn_index=turn_index,
                        )
                        db2.add(msg_row)
                        db2.flush()
                        artifacts = _extract_task_artifacts(
                            full_assistant,
                            prefix=f"{emp.display_name}-回复-{turn_index + 1}",
                        )
                        for art in artifacts:
                            artifact_row = TaskArtifact(
                                tenant_id=p.tenant.id,
                                user_id=p.user.id,
                                session_id=sid,
                                message_id=msg_row.id,
                                kind=art["kind"],
                                name=art["name"],
                                mime_type=art["mime_type"],
                                content=art["content"],
                                source="assistant",
                                run_id=session_run_id,
                                employee_id=emp.id,
                                version=_next_artifact_version(db2, session_id=sid, name=art["name"]),
                                provenance_payload=json.dumps({
                                    "origin": "assistant_message",
                                    "hermes_run_id": current_hermes_run_id,
                                    "turn_index": turn_index,
                                }, ensure_ascii=False),
                            )
                            db2.add(artifact_row)
                            db2.flush()
                            persisted_artifact_ids.append(artifact_row.id)
                        session_row = db2.get(SessionRecord, sid)
                        if session_row:
                            session_row.task_status = final_task_status
                            session_row.updated_at = datetime.now(timezone.utc)
                        db2.commit()
                    finally:
                        db2.close()

                yield trace_event(
                    "runtime",
                    "调用 Hermes Gateway",
                    f"使用{'Run Events' if use_run_events else 'Session Stream'} 链路 {target.base_url}",
                    speaker,
                    hermes_session_id=rec.hermes_session_id,
                )

                async def runtime_stream():
                    nonlocal current_hermes_run_id
                    if use_run_events:
                        run = await hermes_client.create_run(
                            target,
                            message=scoped_msg,
                            session_id=rec.hermes_session_id,
                            system_message=system_message or None,
                            reasoning_effort=body.reasoning_effort,
                        )
                        run_id = str(run.get("run_id") or "")
                        current_hermes_run_id = run_id
                        mark_run_progress(
                            status="running",
                            node_status="running",
                            stage="runtime",
                            reason="Hermes run started",
                            event_type="run.started",
                            hermes_run_id=run_id,
                        )
                        record_step_event(
                            "run.started",
                            {**run, "hermes_run_id": run_id},
                            title="Hermes Run 启动",
                            summary=f"Hermes run {run_id} 已启动",
                            checkpoint_type="run.started",
                        )
                        _remember_hermes_run(
                            run_id,
                            tenant_id=p.tenant.id,
                            user_id=p.user.id,
                            session_id=rec.id,
                            employee_id=emp.id,
                            hermes_session_id=rec.hermes_session_id,
                        )
                        yield {"event": "run.started", "data": {**run, "hermes_run_id": run_id}}
                        run_events = hermes_client.stream_run_events(target, run_id).__aiter__()
                        try:
                            idle_timeout_seconds = max(
                                15.0,
                                float(os.environ.get("OPENATLAS_RUN_EVENT_IDLE_TIMEOUT_SECONDS", "60")),
                            )
                        except ValueError:
                            idle_timeout_seconds = 60.0
                        try:
                            detach_after_seconds = max(
                                idle_timeout_seconds,
                                float(os.environ.get("OPENATLAS_RUN_EVENT_DETACH_AFTER_SECONDS", "900")),
                            )
                        except ValueError:
                            detach_after_seconds = 900.0
                        pending_dangerous_tool: dict[str, Any] | None = None
                        synthetic_approval_sent = False
                        synthetic_approval_keys: set[str] = set()
                        last_event_at = datetime.now(timezone.utc)
                        while True:
                            try:
                                wait_seconds = 8 if pending_dangerous_tool and not synthetic_approval_sent else idle_timeout_seconds
                                run_ev = await asyncio.wait_for(run_events.__anext__(), timeout=wait_seconds)
                                last_event_at = datetime.now(timezone.utc)
                            except StopAsyncIteration:
                                break
                            except asyncio.TimeoutError:
                                if pending_dangerous_tool and not synthetic_approval_sent:
                                    synthetic_approval_sent = True
                                    approval_payload = {
                                        "command": pending_dangerous_tool.get("command") or "",
                                        "description": "Hermes 工具已进入高风险等待态，需要人工确认",
                                        "pattern_key": "openatlas synthetic dangerous terminal confirmation",
                                        "choices": ["once", "session", "always", "deny"],
                                        "hermes_run_id": run_id,
                                        "source": "openatlas_timeout_guard",
                                    }
                                    record_step_event(
                                        "approval.required",
                                        approval_payload,
                                        title="等待人工确认",
                                        summary=approval_payload["description"],
                                        risk_level="high",
                                        checkpoint_type="approval.required",
                                    )
                                    yield {
                                        "event": "openatlas.approval_required",
                                        "data": approval_payload,
                                    }
                                    continue
                                idle_for = (datetime.now(timezone.utc) - last_event_at).total_seconds()
                                payload = {
                                    "message": (
                                        f"Hermes Run Events 已超过 {int(idle_timeout_seconds)} 秒没有新事件，"
                                        "OpenAtlas 会继续等待后台结果。"
                                    ),
                                    "idle_seconds": int(idle_for),
                                    "idle_timeout_seconds": int(idle_timeout_seconds),
                                    "detach_after_seconds": int(detach_after_seconds),
                                    "hermes_run_id": run_id,
                                }
                                yield {"event": "openatlas.run_idle", "data": payload}
                                if idle_for >= detach_after_seconds:
                                    record_step_event(
                                        "openatlas.run_detached",
                                        payload,
                                        title="长任务转后台",
                                        summary="Hermes 长时间未推送新事件，OpenAtlas 转入后台补同步",
                                        risk_level="medium",
                                        checkpoint_type="stalled",
                                    )
                                    yield {
                                        "event": "openatlas.run_detached",
                                        "data": {
                                            **payload,
                                            "message": "Hermes 后台可能仍在运行，OpenAtlas 已停止前端长连接等待，稍后会从 Hermes 会话补同步结果。",
                                        },
                                    }
                                    break
                                continue
                            ev_name = run_ev.get("event") or "message"
                            ev_data = run_ev.get("data") or {}
                            if _is_approval_event(ev_name, ev_data):
                                pending_dangerous_tool = None
                                synthetic_approval_sent = False
                                approval_payload = _normalize_approval_payload(ev_name, ev_data, run_id, source="hermes")
                                mark_run_progress(
                                    status="waiting_approval",
                                    node_status="waiting_approval",
                                    stage="approval",
                                    reason="等待用户确认授权",
                                    event_type=ev_name,
                                    hermes_run_id=run_id,
                                )
                                record_step_event(
                                    "approval.required",
                                    approval_payload,
                                    title="等待人工确认",
                                    summary=str(approval_payload.get("description") or approval_payload.get("command") or "等待用户确认授权"),
                                    risk_level="high",
                                    checkpoint_type="approval.required",
                                )
                                yield {
                                    "event": "openatlas.approval_required",
                                    "data": approval_payload,
                                }
                                continue
                            if ev_name == "message.delta":
                                yield {"event": "assistant.delta", "data": {"delta": ev_data.get("delta") or "", "hermes_run_id": run_id}}
                            elif ev_name == "reasoning.available":
                                yield {"event": "openatlas.reasoning", "data": {"text": ev_data.get("text") or "", "hermes_run_id": run_id}}
                            elif ev_name == "approval.responded":
                                yield {"event": "openatlas.approval_responded", "data": {**ev_data, "hermes_run_id": run_id}}
                            elif ev_name == "tool.started":
                                preview = str(ev_data.get("preview") or ev_data.get("command") or "")
                                tool_name = str(ev_data.get("tool") or ev_data.get("tool_name") or "tool")
                                tool_call_id = ev_data.get("tool_call_id") or ev_data.get("id")
                                tool_started_payload = {
                                    **ev_data,
                                    "tool_name": tool_name,
                                    "label": preview or ev_data.get("label") or ev_data.get("tool") or "",
                                    "tool_call_id": tool_call_id,
                                    "args": ev_data.get("args") or ev_data.get("input") or ev_data.get("parameters"),
                                    "preview": preview,
                                    "command": ev_data.get("command") or preview,
                                    "hermes_run_id": run_id,
                                }
                                if _is_approval_sensitive_tool(tool_name):
                                    pending_dangerous_tool = {"command": tool_started_payload.get("command") or preview, "tool_name": tool_name}
                                yield {"event": "tool.started", "data": tool_started_payload}
                                if _is_approval_sensitive_tool(tool_name) and os.environ.get("OPENATLAS_APPROVE_EXEC_TOOLS", "1") != "0":
                                    approval_key = f"{run_id}:{tool_call_id or tool_name}:{tool_started_payload.get('command') or preview}"
                                    if approval_key not in synthetic_approval_keys:
                                        synthetic_approval_keys.add(approval_key)
                                        synthetic_approval_sent = True
                                        mark_run_progress(
                                            status="waiting_approval",
                                            node_status="waiting_approval",
                                            stage="approval",
                                            reason="等待用户确认工具调用",
                                            event_type="openatlas.approval_required",
                                            hermes_run_id=run_id,
                                        )
                                        approval_payload = _normalize_approval_payload(
                                            "openatlas.synthetic_tool_approval",
                                            tool_started_payload,
                                            run_id,
                                            source="openatlas_exec_tool_guard",
                                        )
                                        record_step_event(
                                            "approval.required",
                                            approval_payload,
                                            title="等待工具授权",
                                            summary=str(approval_payload.get("description") or approval_payload.get("command") or "等待用户确认工具调用"),
                                            risk_level="high",
                                            checkpoint_type="approval.required",
                                        )
                                        yield {
                                            "event": "openatlas.approval_required",
                                            "data": approval_payload,
                                        }
                            elif ev_name == "tool.completed":
                                pending_dangerous_tool = None
                                yield {"event": "tool.failed" if ev_data.get("error") else "tool.completed", "data": {
                                    **ev_data,
                                    "tool_name": ev_data.get("tool") or ev_data.get("tool_name") or "tool",
                                    "label": ev_data.get("label") or ev_data.get("tool") or "",
                                    "tool_call_id": ev_data.get("tool_call_id") or ev_data.get("id"),
                                    "args": ev_data.get("args") or ev_data.get("input") or ev_data.get("parameters"),
                                    "result": ev_data.get("result") or ev_data.get("output"),
                                    "error": ev_data.get("error"),
                                    "duration": ev_data.get("duration"),
                                    "hermes_run_id": run_id,
                                }}
                            elif ev_name == "run.failed":
                                yield {"event": "error", "data": {"message": ev_data.get("error") or "Hermes run failed", "hermes_run_id": run_id}}
                            else:
                                yield {"event": ev_name, "data": ev_data}
                    else:
                        async for old_ev in hermes_client.stream_chat(
                            target,
                            rec.hermes_session_id,
                            message=scoped_msg,
                            system_message=system_message or None,
                        ):
                            yield old_ev

                async for ev in runtime_stream():
                    name = ev.get("event") or "message"
                    data = ev.get("data")
                    completed_fallback = ""
                    if name in {"openatlas.run_detached", "openatlas.run_idle"}:
                        runtime_detached = True
                        detached_run_id = ""
                        if isinstance(data, dict):
                            detached_run_id = str(data.get("hermes_run_id") or data.get("run_id") or "")
                        if name == "openatlas.run_detached":
                            _schedule_detached_run_reconcile(detached_run_id or current_hermes_run_id)
                        mark_run_progress(
                            status="stalled",
                            node_status="stalled",
                            stage="runtime",
                            reason="Hermes 长时间未推送新事件，等待后台补同步",
                            event_type=name,
                            hermes_run_id=detached_run_id or current_hermes_run_id,
                        )
                        record_step_event(
                            name,
                            data if isinstance(data, dict) else {"message": str(data)},
                            title="长任务停滞",
                            summary="Hermes 长时间未推送新事件，OpenAtlas 已保留检查点并等待补同步",
                            risk_level="medium",
                            checkpoint_type="stalled",
                        )
                    if name == "run.completed" and isinstance(data, dict):
                        usage = data.get("usage") if isinstance(data.get("usage"), dict) else {}
                        assistant_usage["input_tokens"] = int(usage.get("input_tokens") or 0)
                        assistant_usage["output_tokens"] = int(usage.get("output_tokens") or 0)
                        assistant_usage["total_tokens"] = int(
                            usage.get("total_tokens")
                            or assistant_usage["input_tokens"] + assistant_usage["output_tokens"]
                        )
                    if name == "run.completed" and isinstance(data, dict) and not assistant_text_parts:
                        output = data.get("output") or data.get("final_response") or ""
                        if isinstance(output, str):
                            completed_fallback += output
                        for msg_item in data.get("messages") or []:
                            if isinstance(msg_item, dict) and msg_item.get("role") == "assistant":
                                content = msg_item.get("content") or ""
                                if isinstance(content, str):
                                    completed_fallback += content
                        if completed_fallback:
                            assistant_text_parts.append(completed_fallback)
                            fallback_data = {**speaker, "delta": completed_fallback, "content": completed_fallback}
                            yield f"event: assistant.delta\ndata: {json.dumps(fallback_data, ensure_ascii=False)}\n\n"
                    if name == "assistant.completed" and isinstance(data, dict):
                        data = {k: v for k, v in data.items() if k != "content"}
                    if name == "assistant.delta" and isinstance(data, dict):
                        msg_obj = data.get("message") if isinstance(data.get("message"), dict) else {}
                        d = (
                            data.get("delta")
                            or data.get("content")
                            or data.get("text")
                            or msg_obj.get("content")
                            or msg_obj.get("delta")
                            or ""
                        )
                        if isinstance(d, str):
                            assistant_text_parts.append(d)
                    if name == "openatlas.reasoning" and isinstance(data, dict):
                        reasoning_text = str(data.get("text") or "").strip()
                        if reasoning_text:
                            reasoning_parts.append(reasoning_text)
                            record_step_event(
                                "reasoning.summary",
                                data,
                                title="思考摘要",
                                summary=reasoning_text[:1000],
                            )
                    if name == "openatlas.approval_responded" and isinstance(data, dict):
                        record_step_event(
                            "approval.resolved",
                            data,
                            title="审批已处理",
                            summary=str(data.get("choice") or data.get("status") or "用户已处理审批"),
                            checkpoint_type="approval.resolved",
                        )
                    if name.startswith("tool."):
                        tool_payload = data if isinstance(data, dict) else {}
                        raw_call_id = (
                            tool_payload.get("tool_call_id")
                            or tool_payload.get("id")
                            or tool_payload.get("call_id")
                            or ""
                        )
                        call_id = str(raw_call_id)
                        tool_name = str(
                            tool_payload.get("tool_name")
                            or tool_payload.get("name")
                            or tool_payload.get("tool")
                            or "tool"
                        )
                        label = str(tool_payload.get("label") or tool_payload.get("preview") or tool_payload.get("command") or "")
                        status = (
                            "running" if name == "tool.started"
                            else "completed" if name == "tool.completed"
                            else "failed" if name == "tool.failed"
                            else "progress"
                        )
                        if not call_id:
                            call_id = f"{tool_name}:{label or 'default'}"
                        next_record = {
                            "tool_call_id": call_id,
                            "name": tool_name,
                            "label": label,
                            "status": status,
                            "args": tool_payload.get("args"),
                            "result": tool_payload.get("result"),
                            "error": tool_payload.get("error"),
                            "duration": tool_payload.get("duration"),
                            "preview": tool_payload.get("preview"),
                            "command": tool_payload.get("command"),
                            "hermes_run_id": tool_payload.get("hermes_run_id"),
                            "runtime_event": name,
                        }
                        existing_idx = next(
                            (
                                i for i, item in enumerate(tool_call_records)
                                if item.get("tool_call_id") == call_id
                                or (
                                    not item.get("tool_call_id")
                                    and item.get("name") == tool_name
                                    and item.get("label") == label
                                )
                                or (
                                    status in ("completed", "failed")
                                    and item.get("name") == tool_name
                                    and item.get("status") in ("running", "progress")
                                )
                            ),
                            None,
                        )
                        if existing_idx is None:
                            tool_call_records.append({k: v for k, v in next_record.items() if v not in (None, "")})
                        else:
                            if not raw_call_id:
                                next_record.pop("tool_call_id", None)
                            if label in ("", tool_name, "tool"):
                                next_record.pop("label", None)
                            merged = {**tool_call_records[existing_idx], **{k: v for k, v in next_record.items() if v not in (None, "")}}
                            tool_call_records[existing_idx] = merged
                        if name == "tool.completed":
                            live_artifacts = _persist_tool_event_artifacts(
                                tenant_id=p.tenant.id,
                                user_id=p.user.id,
                                session_id=rec.id,
                                payload=tool_payload,
                                run_id=session_run_id,
                                employee_id=emp.id,
                                hermes_run_id=current_hermes_run_id or tool_payload.get("hermes_run_id"),
                            )
                            if live_artifacts:
                                live_artifact_ids = [str(a.get("id")) for a in live_artifacts if a.get("id")]
                                mark_run_progress(
                                    status="running",
                                    node_status="running",
                                    stage="artifact",
                                    reason=f"登记 {len(live_artifacts)} 个工具交付物",
                                    event_type="artifact.created",
                                    hermes_run_id=current_hermes_run_id or tool_payload.get("hermes_run_id"),
                                    artifact_ids=live_artifact_ids,
                                )
                                record_step_event(
                                    "artifact.created",
                                    {"count": len(live_artifacts), "items": live_artifacts, "turn_index": turn_index},
                                    title="工具产物入库",
                                    summary=f"从工具事件登记 {len(live_artifacts)} 个交付物",
                                    artifact_ids=live_artifact_ids,
                                    checkpoint_type="artifact.created",
                                )
                                yield trace_event(
                                    "artifact",
                                    "工具产物入库",
                                    f"从工具事件登记 {len(live_artifacts)} 个交付物",
                                    speaker,
                                    artifact_count=len(live_artifacts),
                                )
                                yield (
                                    "event: openatlas.artifacts\n"
                                    f"data: {json.dumps({**speaker, 'items': live_artifacts}, ensure_ascii=False)}\n\n"
                                )
                                _record_canvas_runtime_event(
                                    tenant_id=p.tenant.id,
                                    user_id=p.user.id,
                                    session_id=rec.id,
                                    employee_id=emp.id,
                                    event_type="artifact.created",
                                    payload={"count": len(live_artifacts), "items": live_artifacts, "turn_index": turn_index},
                                )
                        _record_canvas_runtime_event(
                            tenant_id=p.tenant.id,
                            user_id=p.user.id,
                            session_id=rec.id,
                            employee_id=emp.id,
                            event_type=name,
                            payload={
                                "tool_name": tool_payload.get("tool_name") or tool_payload.get("name") or tool_payload.get("tool"),
                                "label": tool_payload.get("label") or tool_payload.get("preview") or tool_payload.get("command"),
                                "status": status,
                                "turn_index": turn_index,
                            },
                        )
                        yield trace_event(
                            "tool",
                            "工具调用事件",
                            f"{name}: {tool_payload.get('tool_name') or tool_payload.get('name') or tool_payload.get('label') or 'tool'}",
                            speaker,
                            runtime_event=name,
                        )
                        mark_run_progress(
                            status="failed" if name == "tool.failed" else "running",
                            node_status="failed" if name == "tool.failed" else "running",
                            stage="tool",
                            reason=str(tool_payload.get("error") or tool_payload.get("label") or tool_payload.get("tool_name") or tool_payload.get("name") or ""),
                            event_type=name,
                            hermes_run_id=current_hermes_run_id or tool_payload.get("hermes_run_id"),
                            error=str(tool_payload.get("error") or "") if name == "tool.failed" else None,
                        )
                        record_step_event(
                            name,
                            tool_payload,
                            title=f"工具: {tool_name}",
                            summary=str(tool_payload.get("error") or tool_payload.get("label") or tool_payload.get("preview") or tool_payload.get("command") or tool_name),
                            risk_level="high" if _is_approval_sensitive_tool(tool_name) else "low",
                            tool_name=tool_name,
                            checkpoint_type=("tool.completed" if name == "tool.completed" else "tool.failed" if name == "tool.failed" else None),
                        )
                    if name in ("run.completed", "done"):
                        persist_assistant_once()
                        done_status = "completed" if final_task_status == "completed" else (
                            "waiting_input" if final_task_status == "needs_input" else final_task_status
                        )
                        mark_run_progress(
                            status=done_status,
                            node_status="done" if done_status == "completed" else done_status,
                            stage="assistant",
                            reason=final_task_reason,
                            event_type=name,
                            hermes_run_id=current_hermes_run_id,
                            artifact_ids=persisted_artifact_ids,
                            output_summary="".join(assistant_text_parts)[:1000],
                        )
                        record_step_event(
                            "node.completed" if done_status == "completed" else "node.waiting_input",
                            {"runtime_event": name, "task_status": final_task_status, "reason": final_task_reason},
                            title="节点完成" if done_status == "completed" else "节点等待补充",
                            summary=final_task_reason or ("节点已完成" if done_status == "completed" else "节点需要补充信息"),
                            output_summary="".join(assistant_text_parts)[:1000],
                            artifact_ids=persisted_artifact_ids,
                            checkpoint_type="node.completed" if done_status == "completed" else "node.waiting_input",
                        )
                    if isinstance(data, dict):
                        data = {**data, **speaker}
                    if isinstance(data, (dict, list)):
                        data_str = json.dumps(data, ensure_ascii=False)
                    else:
                        data_str = str(data)
                    yield f"event: {name}\ndata: {data_str}\n\n"
                    if name in ("run.completed", "done"):
                        break

                if runtime_detached and not "".join(assistant_text_parts).strip():
                    db_detached = SessionLocal()
                    try:
                        session_row = db_detached.get(SessionRecord, sid)
                        if session_row:
                            session_row.task_status = "running"
                            session_row.updated_at = datetime.now(timezone.utc)
                        db_detached.commit()
                    finally:
                        db_detached.close()
                    yield trace_event(
                        "runtime",
                        "后台继续运行",
                        "Hermes 长时间未推送新事件，前端已停止等待；稍后打开历史会话会自动补同步最终结果。",
                        speaker,
                        detached=True,
                    )
                    yield task_state_event("running", "Hermes 后台仍在运行，等待补同步", speaker)
                    yield f"event: done\ndata: {json.dumps({'openatlas_session_id': sess_marker, 'detached': True}, ensure_ascii=False)}\n\n"
                    return

                persist_assistant_once()
                full = "".join(assistant_text_parts).strip()
                if full:
                    relay_outputs.append(f"{emp.display_name}:\n{full}")
                yield trace_event(
                    "artifact",
                    "回复归档完成",
                    f"生成 {len(full)} 字回复，并抽取可下载交付物；状态 {final_task_status}",
                    speaker,
                    content_length=len(full),
                    task_status=final_task_status,
                )
                yield task_state_event(final_task_status, final_task_reason, speaker)
                _record_canvas_runtime_event(
                    tenant_id=p.tenant.id,
                    user_id=p.user.id,
                    session_id=rec.id,
                    employee_id=emp.id,
                    event_type="agent.leave",
                    payload={
                        "speaker_name": emp.display_name,
                        "turn_index": turn_index,
                        "content_length": len(full),
                        "task_status": final_task_status,
                    },
                )
                yield f"event: agent_leave\ndata: {json.dumps({**speaker, 'content_length': len(full)}, ensure_ascii=False)}\n\n"
            db_wf_done = SessionLocal()
            try:
                wf = db_wf_done.get(WorkflowRun, workflow_run_id)
                if wf:
                    node_rows = db_wf_done.query(WorkflowNodeRun).filter(
                        WorkflowNodeRun.workflow_run_id == wf.id,
                    ).all()
                    has_failed = any(n.status == "failed" for n in node_rows)
                    has_waiting = any(n.status in {"waiting_approval", "waiting_input"} for n in node_rows)
                    wf.status = "failed" if has_failed else "waiting_input" if has_waiting else "completed"
                    wf.completed_at = datetime.now(timezone.utc) if wf.status in {"completed", "failed"} else None
                    wf.summary = f"协作运行结束: {len(node_rows)} 个节点，状态 {wf.status}"
                    wf.updated_at = datetime.now(timezone.utc)
                    db_wf_done.commit()
            finally:
                db_wf_done.close()
            yield trace_event("done", "任务流结束", "本轮会话执行完成")
            yield f"event: done\ndata: {json.dumps({'openatlas_session_id': sess_marker}, ensure_ascii=False)}\n\n"
        except Exception as e:
            db_err = SessionLocal()
            try:
                session_row = db_err.get(SessionRecord, rec.id)
                if session_row:
                    session_row.task_status = "failed"
                    session_row.updated_at = datetime.now(timezone.utc)
                wf = db_err.get(WorkflowRun, workflow_run_id)
                if wf:
                    wf.status = "failed"
                    wf.summary = str(e)[:1000]
                    wf.completed_at = datetime.now(timezone.utc)
                    wf.updated_at = datetime.now(timezone.utc)
                db_err.commit()
            finally:
                db_err.close()
            yield task_state_event("failed", str(e)[:500])
            yield f"event: error\ndata: {json.dumps({'message': str(e), 'openatlas_session_id': sess_marker}, ensure_ascii=False)}\n\n"

    return StreamingResponse(event_gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


# ── Capabilities proxy (models / skills / toolsets / health) ───────────────
@app.get("/api/capabilities")
async def capabilities(p: Principal = Depends(get_principal), db: Session = Depends(get_db)) -> dict:
    target = await hermes_client.resolve_target(db, p.tenant.id)
    return await hermes_client.get_capabilities(target)


@app.get("/api/models")
async def models(p: Principal = Depends(get_principal), db: Session = Depends(get_db)) -> dict:
    target = await hermes_client.resolve_target(db, p.tenant.id)
    return await hermes_client.get_models(target)


@app.get("/api/toolsets")
async def toolsets(p: Principal = Depends(get_principal), db: Session = Depends(get_db)) -> dict:
    target = await hermes_client.resolve_target(db, p.tenant.id)
    return await hermes_client.list_toolsets(target)


@app.get("/api/toolsets/governance")
async def toolsets_governance(
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "toolset governance")
    target = await hermes_client.resolve_target(db, p.tenant.id)
    payload = await hermes_client.list_toolsets(target)
    toolset_items = _normalize_toolset_items(payload)
    employees = db.execute(
        select(DigitalEmployee).where(
            DigitalEmployee.tenant_id == p.tenant.id,
            DigitalEmployee.status != EmployeeStatus.archived,
        ).order_by(DigitalEmployee.display_name.asc())
    ).scalars().all()
    employee_items = []
    for emp in employees:
        assigned = json.loads(emp.toolsets or "[]")
        employee_items.append({
            "id": emp.id,
            "display_name": emp.display_name,
            "avatar": emp.avatar,
            "status": emp.status.value if hasattr(emp.status, "value") else str(emp.status),
            "toolsets": assigned,
            "high_risk_toolsets": [name for name in assigned if _toolset_risk(name) == "high"],
        })
    assigned_counts = {
        item["name"]: sum(1 for emp in employee_items if item["name"] in emp["toolsets"])
        for item in toolset_items
    }
    return {
        "toolsets": [{**item, "assigned_employee_count": assigned_counts.get(item["name"], 0)} for item in toolset_items],
        "employees": employee_items,
        "runtime": {"base_url": target.base_url, "tenant_id": target.tenant_id},
    }


@app.get("/api/runtime/health")
async def runtime_health(p: Principal = Depends(get_principal), db: Session = Depends(get_db)) -> dict:
    target = await hermes_client.resolve_target(db, p.tenant.id)
    return await hermes_client.get_health(target)


@app.get("/api/runtime/status")
async def runtime_status(p: Principal = Depends(get_principal), db: Session = Depends(get_db)) -> dict:
    if p.user.role not in (UserRole.tenant_admin, UserRole.system_admin):
        raise HTTPException(403, "runtime status requires admin role")
    runtime = (
        db.query(HermesRuntime)
        .filter(HermesRuntime.tenant_id == p.tenant.id)
        .order_by(HermesRuntime.created_at.desc())
        .first()
    )
    target = await hermes_client.resolve_target(db, p.tenant.id)
    async def probe(label: str, fn):
        try:
            return {"ok": True, "data": await fn(target)}
        except Exception as exc:
            return {"ok": False, "error": str(exc)[:1000]}
    health = await probe("health", hermes_client.get_health)
    capabilities = await probe("capabilities", hermes_client.get_capabilities)
    models_payload = await probe("models", hermes_client.get_models)
    skills_payload = await probe("skills", hermes_client.list_skills)
    recent_errors = db.query(AuditLog).filter(
        AuditLog.tenant_id == p.tenant.id,
        AuditLog.action.ilike("%error%"),
    ).order_by(AuditLog.created_at.desc()).limit(10).all()
    return {
        "tenant": {"id": p.tenant.id, "slug": p.tenant.slug, "name": p.tenant.name},
        "runtime": {
            "id": runtime.id if runtime else None,
            "status": runtime.status.value if runtime and hasattr(runtime.status, "value") else (runtime.status if runtime else "fallback"),
            "gateway_base_url": runtime.gateway_base_url if runtime else target.base_url,
            "port": runtime.port if runtime else None,
            "pid": runtime.pid if runtime else None,
            "hermes_home": runtime.hermes_home_path if runtime else str(HERMES_HOME),
            "health_checked_at": runtime.health_checked_at.isoformat() if runtime and runtime.health_checked_at else None,
        },
        "checks": {
            "health": health,
            "capabilities": capabilities,
            "models": models_payload,
            "skills": skills_payload,
        },
        "recent_errors": [
            {
                "id": r.id,
                "action": r.action,
                "resource_type": r.resource_type,
                "resource_id": r.resource_id,
                "metadata": json.loads(r.extra or "{}"),
                "created_at": r.created_at.isoformat() if r.created_at else None,
            }
            for r in recent_errors
        ],
    }


@app.get("/api/runtime/diagnostics")
async def runtime_diagnostics(
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "runtime diagnostics")
    target = await hermes_client.resolve_target(db, p.tenant.id)

    async def probe(name: str, fn):
        started = datetime.now(timezone.utc)
        try:
            data = await fn(target)
            return {
                "name": name,
                "ok": True,
                "latency_ms": int((datetime.now(timezone.utc) - started).total_seconds() * 1000),
                "data": data,
            }
        except Exception as exc:
            return {
                "name": name,
                "ok": False,
                "latency_ms": int((datetime.now(timezone.utc) - started).total_seconds() * 1000),
                "error": str(exc)[:1000],
            }

    checks = [
        await probe("health", hermes_client.get_health),
        await probe("capabilities", hermes_client.get_capabilities),
        await probe("models", hermes_client.get_models),
        await probe("skills", hermes_client.list_skills),
        await probe("toolsets", hermes_client.list_toolsets),
    ]
    hermes_home = _tenant_hermes_home(db, p.tenant.id)
    return {
        "tenant": {"id": p.tenant.id, "slug": p.tenant.slug, "name": p.tenant.name},
        "runtime": {"base_url": target.base_url, "hermes_home": str(hermes_home)},
        "checks": checks,
        "all_ok": all(c["ok"] for c in checks),
    }


@app.get("/api/runtime/logs")
def runtime_logs(
    kind: str = Query(default="gateway", max_length=64),
    tail: int = Query(default=120, ge=20, le=1000),
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "runtime logs")
    hermes_home = _tenant_hermes_home(db, p.tenant.id)
    allowed = {"agent", "errors", "gateway", "gui", "openatlas"}
    if kind not in allowed:
        raise HTTPException(400, f"invalid log kind: {kind}")
    if kind == "openatlas":
        runtime = (
            db.query(HermesRuntime)
            .filter(HermesRuntime.tenant_id == p.tenant.id)
            .order_by(HermesRuntime.created_at.desc())
            .first()
        )
        path = _tenant_log_path(runtime.hermes_home_path if runtime else str(hermes_home))
        return {"ok": bool(path.exists()), "kind": kind, "path": str(path), "output": _tail_file(path, lines=tail)}
    try:
        result = _run_hermes_cli(hermes_home, ["logs", kind, "-n", str(tail)], timeout=30, allow_error=True)
    except Exception as exc:
        log_path = hermes_home / "logs" / f"{kind}.log"
        return {"ok": False, "kind": kind, "path": str(log_path), "output": _tail_file(log_path, lines=tail), "error": str(exc)[:500]}
    return {"kind": kind, **result}


@app.post("/api/runtime/ops/{action}")
def runtime_operation(
    action: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "runtime operation")
    hermes_home = _tenant_hermes_home(db, p.tenant.id)
    clean_action = action.strip().lower()
    if clean_action == "doctor":
        result = _run_hermes_cli(hermes_home, ["doctor"], timeout=120, allow_error=True)
    elif clean_action == "security-audit":
        result = _run_hermes_cli(hermes_home, ["security", "audit", "--json"], timeout=180, allow_error=True)
    elif clean_action == "prompt-size":
        result = _run_hermes_cli(hermes_home, ["prompt-size", "--platform", "api_server", "--json"], timeout=60, allow_error=True)
    elif clean_action == "logs-list":
        result = _run_hermes_cli(hermes_home, ["logs", "list"], timeout=30, allow_error=True)
    elif clean_action == "backup":
        stamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
        out = Path(OPENATLAS_HOME) / "backups" / f"hermes-{p.tenant.slug}-{stamp}.zip"
        try:
            result = _run_hermes_cli(hermes_home, ["backup", "--output", str(out)], timeout=240, allow_error=True)
        except Exception as exc:
            size = _zip_directory(hermes_home, out)
            result = {"ok": True, "exit_code": 0, "argv": ["openatlas", "zip-backup"], "output": f"fallback backup created: {out}", "fallback_error": str(exc)[:500], "source_bytes": size}
        result["backup_path"] = str(out)
        result["backup_size"] = out.stat().st_size if out.exists() else 0
    else:
        raise HTTPException(400, "unsupported runtime operation")
    audit(db, principal=p, action=f"runtime.{clean_action}", resource_type="runtime",
          resource_id=p.tenant.id, request=request, extra={
              "ok": result.get("ok"),
              "exit_code": result.get("exit_code"),
              "argv": result.get("argv"),
              "backup_path": result.get("backup_path"),
          })
    db.commit()
    return {"action": clean_action, "hermes_home": str(hermes_home), **result}


@app.get("/api/tenant/isolation")
def tenant_isolation_status(p: Principal = Depends(get_principal), db: Session = Depends(get_db)) -> dict:
    if p.user.role not in (UserRole.tenant_admin, UserRole.system_admin):
        raise HTTPException(403, "tenant isolation status requires admin role")
    runtime = (
        db.query(HermesRuntime)
        .filter(HermesRuntime.tenant_id == p.tenant.id)
        .order_by(HermesRuntime.created_at.desc())
        .first()
    )
    uploads_root = (Path(str(OPENATLAS_HOME)) / "uploads" / p.tenant.id).resolve()
    hermes_home = Path(runtime.hermes_home_path).resolve() if runtime and runtime.hermes_home_path else HERMES_HOME
    session_count = db.query(SessionRecord).filter(SessionRecord.tenant_id == p.tenant.id).count()
    file_count = db.query(FileAsset).filter(FileAsset.tenant_id == p.tenant.id).count()
    employee_count = db.query(DigitalEmployee).filter(DigitalEmployee.tenant_id == p.tenant.id).count()
    skill_count = db.query(SkillPackage).filter(
        (SkillPackage.scope == Scope.global_)
        | (SkillPackage.owner_tenant_id == p.tenant.id)
        | (SkillPackage.owner_user_id == p.user.id)
    ).count()
    memory_count = db.query(MemoryEntry).filter(MemoryEntry.tenant_id == p.tenant.id).count()
    checks = [
        {"name": "Hermes home isolated", "ok": p.tenant.slug in str(hermes_home) or p.tenant.id in str(hermes_home), "value": str(hermes_home)},
        {"name": "Upload root isolated", "ok": p.tenant.id in str(uploads_root), "value": str(uploads_root)},
        {"name": "Session tenant filter", "ok": True, "value": f"{session_count} sessions scoped to tenant"},
        {"name": "Skill scope filter", "ok": True, "value": f"{skill_count} visible skills"},
        {"name": "Memory tenant filter", "ok": True, "value": f"{memory_count} memories scoped to tenant"},
    ]
    return {
        "tenant": {"id": p.tenant.id, "slug": p.tenant.slug, "name": p.tenant.name, "role": p.user.role.value},
        "paths": {
            "openatlas_home": str(OPENATLAS_HOME),
            "hermes_home": str(hermes_home),
            "uploads_root": str(uploads_root),
        },
        "counts": {
            "sessions": session_count,
            "files": file_count,
            "employees": employee_count,
            "skills_visible": skill_count,
            "memories": memory_count,
        },
        "checks": checks,
        "ok": all(c["ok"] for c in checks),
    }


# ── Jobs (Phase 3.5 — own schedule + state; Hermes is read-only) ──────────
def _parse_interval_seconds(expr: str) -> int:
    raw = (expr or "").strip().lower()
    if not raw:
        return 3600
    m = re.match(r"^(\d+)\s*([smhd]?)$", raw)
    if not m:
        raise HTTPException(400, "interval schedule_expr must be like 30s, 15m, 2h, or seconds")
    value = int(m.group(1))
    unit = m.group(2) or "s"
    factor = {"s": 1, "m": 60, "h": 3600, "d": 86400}[unit]
    return max(30, value * factor)


def _compute_next_job_run_at(kind: str, expr: str, *, after: datetime | None = None) -> datetime | None:
    base = after or datetime.utcnow()
    kind = (kind or "cron").strip().lower()
    expr = (expr or "").strip()
    if kind == "once":
        if not expr or expr.lower() in {"now", "once"}:
            return base
        try:
            dt = datetime.fromisoformat(expr.replace("Z", "+00:00"))
            if dt.tzinfo is not None:
                dt = dt.astimezone(timezone.utc).replace(tzinfo=None)
            return dt if dt > base else base
        except Exception as exc:
            raise HTTPException(400, "once schedule_expr must be ISO datetime or now") from exc
    if kind == "interval":
        return base + timedelta(seconds=_parse_interval_seconds(expr))
    if kind != "cron":
        raise HTTPException(400, "schedule_kind must be cron, interval, or once")
    parts = expr.split()
    if len(parts) != 5:
        raise HTTPException(400, "cron schedule_expr must have five fields")
    minute_raw, hour_raw = parts[0], parts[1]
    minute = int(minute_raw) if minute_raw.isdigit() else base.minute
    hour = int(hour_raw) if hour_raw.isdigit() else None
    if not (0 <= minute <= 59):
        raise HTTPException(400, "cron minute must be 0-59")
    if hour is not None and not (0 <= hour <= 23):
        raise HTTPException(400, "cron hour must be 0-23")
    if hour is None:
        candidate = base.replace(minute=minute, second=0, microsecond=0)
        if candidate <= base:
            candidate += timedelta(hours=1)
        return candidate
    candidate = base.replace(hour=hour, minute=minute, second=0, microsecond=0)
    if candidate <= base:
        candidate += timedelta(days=1)
    return candidate


def _refresh_job_next_run(j: Job, *, after: datetime | None = None) -> None:
    j.next_run_at = _compute_next_job_run_at(j.schedule_kind, j.schedule_expr, after=after)


async def _execute_job_once(db: Session, j: Job, *, request: Request | None = None, principal: Principal | None = None) -> None:
    target = await hermes_client.resolve_target(db, j.tenant_id)
    sess = await hermes_client.create_session(target, title=f"job:{j.name}")
    hermes_sid = sess.get("id") or sess.get("session_id")
    if not hermes_sid:
        raise HTTPException(502, "hermes did not return a session id")
    now = datetime.utcnow()
    j.last_run_at = now
    try:
        last_text: list[str] = []
        async for ev in hermes_client.stream_chat(
            target,
            hermes_sid,
            message=j.prompt or "(no prompt)",
        ):
            if ev.get("event") == "assistant.delta":
                d = ev.get("data") or {}
                t = d.get("delta") or d.get("content") or ""
                if isinstance(t, str):
                    last_text.append(t)
        j.last_status = "ok"
        j.last_error = ""
        db.add(SessionRecord(
            tenant_id=j.tenant_id,
            user_id=j.created_by,
            employee_id=j.employee_id,
            hermes_session_id=hermes_sid,
            title=f"[job] {j.name}",
            last_message=("".join(last_text))[:500],
            message_count=1,
            task_status="completed",
        ))
    except Exception as ex:
        j.last_status = "error"
        j.last_error = str(ex)[:1000]
    if request is None:
        if (j.schedule_kind or "").lower() == "once":
            j.status = JobStatus.disabled
            j.next_run_at = None
        else:
            _refresh_job_next_run(j, after=now)
    audit(db, principal=principal, action="job.run", resource_type="job",
          resource_id=j.id, request=request, extra={"status": j.last_status, "scheduled": request is None})


async def _job_scheduler_loop() -> None:
    interval = float(os.environ.get("OPENATLAS_JOB_SCHEDULER_INTERVAL_SECONDS", "30"))
    while True:
        await asyncio.sleep(max(5.0, interval))
        if os.environ.get("OPENATLAS_DISABLE_JOB_SCHEDULER") == "1":
            continue
        now = datetime.utcnow()
        with SessionLocal() as db:
            due = db.query(Job).filter(
                Job.status == JobStatus.active,
                Job.next_run_at != None,  # noqa: E711
                Job.next_run_at <= now,
            ).order_by(Job.next_run_at.asc()).limit(5).all()
            for j in due:
                if j.id in _RUNNING_JOB_IDS:
                    continue
                _RUNNING_JOB_IDS.add(j.id)
                try:
                    await _execute_job_once(db, j)
                    db.commit()
                except Exception as exc:  # noqa: BLE001
                    j.last_status = "error"
                    j.last_error = str(exc)[:1000]
                    try:
                        _refresh_job_next_run(j, after=now)
                    except Exception:
                        j.status = JobStatus.paused
                    db.commit()
                finally:
                    _RUNNING_JOB_IDS.discard(j.id)


def _start_job_scheduler() -> None:
    global _JOB_SCHEDULER_TASK
    if os.environ.get("OPENATLAS_DISABLE_JOB_SCHEDULER") == "1":
        return
    if _JOB_SCHEDULER_TASK and not _JOB_SCHEDULER_TASK.done():
        return
    try:
        _JOB_SCHEDULER_TASK = asyncio.create_task(_job_scheduler_loop())
    except RuntimeError:
        return


def _job_to_dict(j: Job) -> dict:
    return {
        "id": j.id,
        "tenant_id": j.tenant_id,
        "name": j.name,
        "description": j.description,
        "schedule_kind": j.schedule_kind,
        "schedule_expr": j.schedule_expr,
        "employee_id": j.employee_id,
        "skills": [s for s in (j.skills or "").split(",") if s],
        "toolsets": [s for s in (j.toolsets or "").split(",") if s],
        "deliver": j.deliver,
        "prompt": j.prompt,
        "status": j.status.value if hasattr(j.status, "value") else j.status,
        "last_run_at": j.last_run_at.isoformat() if j.last_run_at else None,
        "next_run_at": j.next_run_at.isoformat() if j.next_run_at else None,
        "last_status": j.last_status,
        "last_error": j.last_error,
        "created_at": j.created_at.isoformat() if j.created_at else None,
        "updated_at": j.updated_at.isoformat() if j.updated_at else None,
    }


class JobIn(BaseModel):
    name: str
    description: str = ""
    schedule_kind: str = "cron"
    schedule_expr: str = "0 9 * * *"
    employee_id: str | None = None
    skills: list[str] = []
    toolsets: list[str] = []
    deliver: str = "local"
    prompt: str = ""


class JobPatchIn(BaseModel):
    name: str | None = None
    description: str | None = None
    schedule_kind: str | None = None
    schedule_expr: str | None = None
    employee_id: str | None = None
    skills: list[str] | None = None
    toolsets: list[str] | None = None
    deliver: str | None = None
    prompt: str | None = None


@app.get("/api/jobs")
def list_jobs(
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.5 — list this tenant's local Jobs. (Hermes-side jobs are
    ignored; OpenAtlas owns the schedule. See DELIVERY-PHASE-3.5 §4.)"""
    rows = db.query(Job).filter(Job.tenant_id == p.tenant.id)\
        .order_by(Job.created_at.desc()).all()
    return {"items": [_job_to_dict(j) for j in rows]}


@app.post("/api/jobs")
def create_job(
    body: JobIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    if body.employee_id:
        emp = db.get(DigitalEmployee, body.employee_id)
        if not emp or emp.tenant_id != p.tenant.id:
            raise HTTPException(400, "employee not in this tenant")
    j = Job(
        tenant_id=p.tenant.id, name=body.name, description=body.description,
        schedule_kind=body.schedule_kind, schedule_expr=body.schedule_expr,
        employee_id=body.employee_id,
        skills=",".join(body.skills), toolsets=",".join(body.toolsets),
        deliver=body.deliver, prompt=body.prompt,
        status=JobStatus.active, created_by=p.user.id,
    )
    _refresh_job_next_run(j)
    db.add(j)
    db.flush()
    audit(db, principal=p, action="job.create", resource_type="job",
          resource_id=j.id, request=request,
          extra={"name": body.name, "schedule": body.schedule_expr})
    db.commit()
    db.refresh(j)
    return _job_to_dict(j)


@app.get("/api/jobs/{job_id}")
def get_job(
    job_id: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    j = db.get(Job, job_id)
    if not j or j.tenant_id != p.tenant.id:
        raise HTTPException(404, "job not found")
    return _job_to_dict(j)


@app.patch("/api/jobs/{job_id}")
def patch_job(
    job_id: str, body: JobPatchIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    j = db.get(Job, job_id)
    if not j or j.tenant_id != p.tenant.id:
        raise HTTPException(404, "job not found")
    changes = {}
    for fld in ("name", "description", "schedule_kind", "schedule_expr",
                "employee_id", "deliver", "prompt"):
        v = getattr(body, fld)
        if v is not None:
            setattr(j, fld, v)
            changes[fld] = v
    if body.skills is not None:
        j.skills = ",".join(body.skills); changes["skills"] = body.skills
    if body.toolsets is not None:
        j.toolsets = ",".join(body.toolsets); changes["toolsets"] = body.toolsets
    if not changes:
        raise HTTPException(400, "no fields to update")
    if any(k in changes for k in ("schedule_kind", "schedule_expr")):
        _refresh_job_next_run(j)
    audit(db, principal=p, action="job.update", resource_type="job",
          resource_id=j.id, request=request, extra=changes)
    db.commit()
    return _job_to_dict(j)


@app.delete("/api/jobs/{job_id}")
def delete_job(
    job_id: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    j = db.get(Job, job_id)
    if not j or j.tenant_id != p.tenant.id:
        raise HTTPException(404, "job not found")
    audit(db, principal=p, action="job.delete", resource_type="job",
          resource_id=j.id, request=request, extra={"name": j.name})
    db.delete(j)
    db.commit()
    return {"ok": True, "deleted": job_id}


@app.post("/api/jobs/{job_id}/pause")
def pause_job(
    job_id: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    j = db.get(Job, job_id)
    if not j or j.tenant_id != p.tenant.id:
        raise HTTPException(404, "job not found")
    j.status = JobStatus.paused
    audit(db, principal=p, action="job.pause", resource_type="job",
          resource_id=j.id, request=request)
    db.commit()
    return _job_to_dict(j)


@app.post("/api/jobs/{job_id}/resume")
def resume_job(
    job_id: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    j = db.get(Job, job_id)
    if not j or j.tenant_id != p.tenant.id:
        raise HTTPException(404, "job not found")
    j.status = JobStatus.active
    if not j.next_run_at or j.next_run_at <= datetime.utcnow():
        _refresh_job_next_run(j)
    audit(db, principal=p, action="job.resume", resource_type="job",
          resource_id=j.id, request=request)
    db.commit()
    return _job_to_dict(j)


@app.post("/api/jobs/{job_id}/run")
async def run_job(
    job_id: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.5 — manually trigger a job. Creates a session in the tenant
    Hermes, sends the job prompt, and records the outcome. Returns a
    summary; the actual stream is consumed and discarded (this is a
    background-style trigger, not a UI chat)."""
    j = db.get(Job, job_id)
    if not j or j.tenant_id != p.tenant.id:
        raise HTTPException(404, "job not found")
    await _execute_job_once(db, j, request=request, principal=p)
    db.commit()
    return _job_to_dict(j)


# ── Skill Market ───────────────────────────────────────────────────────────
@app.get("/api/hermes-skills/hub/browse")
def browse_hermes_skills_hub(
    page: int = Query(default=1, ge=1),
    size: int = Query(default=20, ge=1, le=100),
    source: str = Query(default="all"),
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    if source not in {"all", "official", "skills-sh", "well-known", "github", "clawhub", "lobehub", "browse-sh"}:
        raise HTTPException(400, "invalid source")
    hermes_home = _tenant_hermes_home(db, p.tenant.id)
    code = (
        "import json, sys;"
        "from hermes_cli.skills_hub import browse_skills;"
        "print(json.dumps(browse_skills(page=int(sys.argv[1]), page_size=int(sys.argv[2]), source=sys.argv[3]), ensure_ascii=False))"
    )
    return _run_hermes_python_json(hermes_home, code, [str(page), str(size), source], timeout=60)


@app.get("/api/hermes-skills/hub/search")
def search_hermes_skills_hub(
    q: str = Query(min_length=1, max_length=200),
    source: str = Query(default="all"),
    limit: int = Query(default=20, ge=1, le=100),
    only_installable: bool = Query(default=True),
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    if source not in {"all", "official", "skills-sh", "well-known", "github", "clawhub", "lobehub", "browse-sh"}:
        raise HTTPException(400, "invalid source")
    hermes_home = _tenant_hermes_home(db, p.tenant.id)
    code = r"""
import json, sys
from tools.skills_hub import GitHubAuth, create_source_router, unified_search
try:
    from hermes_cli.skills_hub import inspect_skill
except Exception:
    inspect_skill = None
query, source, limit, only_installable = sys.argv[1], sys.argv[2], int(sys.argv[3]), sys.argv[4] == "1"
sources = create_source_router(GitHubAuth())
raw_results = unified_search(query, sources, source_filter=source, limit=max(limit * (3 if only_installable else 1), limit))
items = []
for r in raw_results:
    item = {
        "name": r.name,
        "identifier": r.identifier,
        "source": r.source,
        "trust": r.trust_level,
        "description": r.description,
    }
    if only_installable:
        ok = False
        inspect_error = ""
        if inspect_skill is not None:
            try:
                detail = inspect_skill(r.identifier)
                ok = bool(detail)
                if ok and isinstance(detail, dict):
                    item["installable"] = True
                    item["resolved_name"] = detail.get("name") or detail.get("id") or detail.get("identifier")
                    item["version"] = detail.get("version")
            except Exception as exc:
                inspect_error = str(exc)
        if not ok:
            continue
        if inspect_error:
            item["inspect_error"] = inspect_error[:300]
    items.append(item)
    if len(items) >= limit:
        break
print(json.dumps({"items": items, "only_installable": only_installable}, ensure_ascii=False))
"""
    return _run_hermes_python_json(hermes_home, code, [q, source, str(limit), "1" if only_installable else "0"], timeout=90)


@app.get("/api/hermes-skills/hub/inspect")
def inspect_hermes_skill_hub_item(
    identifier: str = Query(min_length=1, max_length=1024),
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    hermes_home = _tenant_hermes_home(db, p.tenant.id)
    code = (
        "import json, sys;"
        "from hermes_cli.skills_hub import inspect_skill;"
        "result = inspect_skill(sys.argv[1]);"
        "print(json.dumps(result or {}, ensure_ascii=False))"
    )
    result = _run_hermes_python_json(hermes_home, code, [identifier], timeout=60)
    if not result:
        raise HTTPException(404, "skill not found in Hermes Skills Hub")
    return result


@app.post("/api/hermes-skills/hub/install")
async def install_hermes_skill_hub_item(
    body: HermesSkillInstallIn,
    request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    if p.user.role not in (UserRole.tenant_admin, UserRole.system_admin):
        raise HTTPException(403, "only tenant_admin can install Hermes skills")
    hermes_home = _tenant_hermes_home(db, p.tenant.id)
    install_identifier = _normalize_hub_install_identifier(body.identifier, body.source)
    result = _run_hermes_skills_install(
        hermes_home,
        identifier=install_identifier,
        category=body.category.strip(),
        name=body.name.strip(),
        force=body.force,
    )
    sync_stats = {}
    try:
        sync_stats = await _sync_installed_hermes_skills(db, p)
    except Exception as exc:
        sync_stats = {"sync_error": str(exc)[:500]}
    candidates = [body.identifier.strip(), install_identifier, body.name.strip()]
    visible_skill = _find_synced_hermes_skill(db, p.tenant.id, candidates)
    visible = visible_skill is not None
    warning = ""
    if not visible:
        warning = (
            "Hermes install command completed, but the skill is not visible in this "
            "tenant's installed Hermes Skills. It may require a different Hub identifier "
            "or a supported Hermes source."
        )
    audit(db, principal=p, action="skill.install_hermes", resource_type="skill",
          resource_id=getattr(visible_skill, "id", None), request=request,
          extra={
              "identifier": body.identifier,
              "install_identifier": install_identifier,
              "source": body.source,
              "category": body.category,
              "visible": visible,
              "warning": warning,
              **sync_stats,
          })
    db.commit()
    return {
        **result,
        "visible": visible,
        "warning": warning,
        "sync": sync_stats,
        "skill": _skill_to_dict(visible_skill, db) if visible_skill else None,
    }


@app.post("/api/hermes-skills/hub/check")
async def check_hermes_skills(
    body: HermesSkillLifecycleIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes skill check")
    argv = ["skills", "check"]
    if body.name.strip():
        argv.append(body.name.strip())
    result = _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), argv, timeout=120, allow_error=True)
    audit(db, principal=p, action="skill.hermes_check", resource_type="skill",
          resource_id=body.name or None, request=request, extra={"ok": result.get("ok"), "argv": argv})
    db.commit()
    return result


@app.post("/api/hermes-skills/hub/update")
async def update_hermes_skills(
    body: HermesSkillLifecycleIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes skill update")
    argv = ["skills", "update"]
    if body.name.strip():
        argv.append(body.name.strip())
    result = _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), argv, timeout=180, allow_error=True)
    sync_stats = {}
    try:
        sync_stats = await _sync_installed_hermes_skills(db, p)
    except Exception as exc:
        sync_stats = {"sync_error": str(exc)[:500]}
    audit(db, principal=p, action="skill.hermes_update", resource_type="skill",
          resource_id=body.name or None, request=request, extra={"ok": result.get("ok"), "argv": argv, **sync_stats})
    db.commit()
    return {**result, "sync": sync_stats}


@app.post("/api/hermes-skills/hub/audit")
def audit_hermes_skills(
    body: HermesSkillLifecycleIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes skill audit")
    argv = ["skills", "audit"]
    if body.name.strip():
        argv.append(body.name.strip())
    if body.deep:
        argv.append("--deep")
    result = _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), argv, timeout=180, allow_error=True)
    audit(db, principal=p, action="skill.hermes_audit", resource_type="skill",
          resource_id=body.name or None, request=request, extra={"ok": result.get("ok"), "argv": argv})
    db.commit()
    return result


@app.post("/api/hermes-skills/hub/uninstall")
async def uninstall_hermes_skill(
    body: HermesSkillLifecycleIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes skill uninstall")
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "name is required")
    result = _run_hermes_cli(
        _tenant_hermes_home(db, p.tenant.id),
        ["skills", "uninstall", name],
        timeout=120,
        input_text="y\n",
        allow_error=True,
    )
    sync_stats = {}
    try:
        sync_stats = await _sync_installed_hermes_skills(db, p)
    except Exception as exc:
        sync_stats = {"sync_error": str(exc)[:500]}
    audit(db, principal=p, action="skill.hermes_uninstall", resource_type="skill",
          resource_id=name, request=request, extra={"ok": result.get("ok"), **sync_stats})
    db.commit()
    return {**result, "sync": sync_stats}


@app.post("/api/hermes-skills/hub/opt-out")
def opt_out_hermes_skills(
    body: HermesSkillLifecycleIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes bundled skill opt-out")
    argv = ["skills", "opt-out"]
    if body.remove:
        argv.extend(["--remove", "--yes"])
    result = _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), argv, timeout=120, allow_error=True)
    audit(db, principal=p, action="skill.hermes_opt_out", resource_type="skill",
          resource_id=None, request=request, extra={"ok": result.get("ok"), "argv": argv})
    db.commit()
    return result


@app.post("/api/hermes-skills/hub/opt-in")
async def opt_in_hermes_skills(
    body: HermesSkillLifecycleIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes bundled skill opt-in")
    argv = ["skills", "opt-in"]
    if body.sync:
        argv.append("--sync")
    result = _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), argv, timeout=180, allow_error=True)
    sync_stats = {}
    try:
        sync_stats = await _sync_installed_hermes_skills(db, p)
    except Exception as exc:
        sync_stats = {"sync_error": str(exc)[:500]}
    audit(db, principal=p, action="skill.hermes_opt_in", resource_type="skill",
          resource_id=None, request=request, extra={"ok": result.get("ok"), "argv": argv, **sync_stats})
    db.commit()
    return {**result, "sync": sync_stats}


@app.post("/api/hermes-skills/hub/reset")
def reset_hermes_skill(
    body: HermesSkillLifecycleIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes skill reset")
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "name is required")
    argv = ["skills", "reset", name]
    if body.restore:
        argv.extend(["--restore", "--yes"])
    result = _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), argv, timeout=120, allow_error=True)
    audit(db, principal=p, action="skill.hermes_reset", resource_type="skill",
          resource_id=name, request=request, extra={"ok": result.get("ok"), "argv": argv})
    db.commit()
    return result


@app.post("/api/hermes-skills/hub/repair-official")
def repair_official_hermes_skill(
    body: HermesSkillLifecycleIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes official skill repair")
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "name is required")
    argv = ["skills", "repair-official", name]
    if body.restore:
        argv.extend(["--restore", "--yes"])
    result = _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), argv, timeout=180, allow_error=True)
    audit(db, principal=p, action="skill.hermes_repair_official", resource_type="skill",
          resource_id=name, request=request, extra={"ok": result.get("ok"), "argv": argv})
    db.commit()
    return result


@app.post("/api/hermes-skills/hub/snapshot/export")
def export_hermes_skill_snapshot(
    request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes skill snapshot export")
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    out = Path(OPENATLAS_HOME) / "backups" / f"skills-{p.tenant.slug}-{stamp}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    result = _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), ["skills", "snapshot", "export", str(out)], timeout=120, allow_error=True)
    result["snapshot_path"] = str(out)
    result["snapshot_size"] = out.stat().st_size if out.exists() else 0
    audit(db, principal=p, action="skill.hermes_snapshot_export", resource_type="skill",
          resource_id=None, request=request, extra={"ok": result.get("ok"), "path": str(out)})
    db.commit()
    return result


@app.post("/api/hermes-skills/hub/snapshot/import")
async def import_hermes_skill_snapshot(
    body: HermesSkillLifecycleIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes skill snapshot import")
    raw_path = body.path.strip()
    if not raw_path:
        raise HTTPException(400, "path is required")
    in_path = Path(raw_path).expanduser().resolve()
    allowed_roots = [Path(OPENATLAS_HOME).resolve(), _tenant_hermes_home(db, p.tenant.id).resolve()]
    if not in_path.exists() or not any(str(in_path).startswith(str(root)) for root in allowed_roots):
        raise HTTPException(400, "snapshot path is not readable from this tenant workspace")
    argv = ["skills", "snapshot", "import", str(in_path)]
    if body.force:
        argv.append("--force")
    result = _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), argv, timeout=240, allow_error=True)
    sync_stats = {}
    try:
        sync_stats = await _sync_installed_hermes_skills(db, p)
    except Exception as exc:
        sync_stats = {"sync_error": str(exc)[:500]}
    audit(db, principal=p, action="skill.hermes_snapshot_import", resource_type="skill",
          resource_id=None, request=request, extra={"ok": result.get("ok"), "path": str(in_path), **sync_stats})
    db.commit()
    return {**result, "sync": sync_stats}


@app.get("/api/hermes-skills/hub/taps")
def list_hermes_skill_taps(
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes skill taps")
    return _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), ["skills", "tap", "list"], timeout=60, allow_error=True)


@app.post("/api/hermes-skills/hub/taps")
def add_hermes_skill_tap(
    body: HermesSkillLifecycleIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes skill taps")
    repo = (body.url or body.name).strip()
    if not repo:
        raise HTTPException(400, "repo url/name is required")
    result = _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), ["skills", "tap", "add", repo], timeout=120, allow_error=True)
    audit(db, principal=p, action="skill.hermes_tap_add", resource_type="skill",
          resource_id=repo, request=request, extra={"ok": result.get("ok")})
    db.commit()
    return result


@app.delete("/api/hermes-skills/hub/taps/{name}")
def remove_hermes_skill_tap(
    name: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    _require_tenant_admin(p, "Hermes skill taps")
    result = _run_hermes_cli(_tenant_hermes_home(db, p.tenant.id), ["skills", "tap", "remove", name], timeout=120, allow_error=True)
    audit(db, principal=p, action="skill.hermes_tap_remove", resource_type="skill",
          resource_id=name, request=request, extra={"ok": result.get("ok")})
    db.commit()
    return result


@app.get("/api/skill-market")
async def list_market(
    scope: str | None = Query(default=None),
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    sync_error = ""
    try:
        await _sync_installed_hermes_skills(db, p)
    except Exception as exc:
        sync_error = str(exc)[:500]
    q = select(SkillPackage).where(
        (SkillPackage.scope == Scope.global_)
        | ((SkillPackage.scope == Scope.tenant) & (SkillPackage.owner_tenant_id == p.tenant.id))
        | ((SkillPackage.scope == Scope.user) & (SkillPackage.owner_user_id == p.user.id))
        | ((SkillPackage.scope == Scope.employee) & (SkillPackage.owner_tenant_id == p.tenant.id))
    )
    if scope:
        q = q.where(SkillPackage.scope == Scope(scope))
    rows = db.execute(q.order_by(SkillPackage.scope, SkillPackage.name)).scalars().all()
    rows = _dedupe_skill_rows(rows)
    rows.sort(key=lambda s: (
        0 if s.status == "enabled" else 1,
        0 if (getattr(s, "source_ref", "") or "").startswith("hermes:") else 1,
        s.scope.value,
        s.name.lower(),
        s.version,
    ))
    return {"items": [_skill_to_dict(s, db) for s in rows], "hermes_sync_error": sync_error}


@app.post("/api/skill-market/sync-hermes")
async def sync_hermes_skill_market(
    request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    if p.user.role not in (UserRole.tenant_admin, UserRole.system_admin):
        raise HTTPException(403, "only tenant_admin can sync Hermes skills")
    stats = await _sync_installed_hermes_skills(db, p)
    audit(db, principal=p, action="skill.sync_hermes", resource_type="skill",
          resource_id=None, request=request, extra=stats)
    db.commit()
    rows = db.execute(
        select(SkillPackage).where(
            SkillPackage.owner_tenant_id == p.tenant.id,
            SkillPackage.source_ref.like("hermes:%"),
        ).order_by(SkillPackage.category, SkillPackage.name)
    ).scalars().all()
    return {**stats, "items": [_skill_to_dict(s, db) for s in rows]}


@app.get("/api/skill-market/health")
def skill_market_health(
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    if p.user.role not in (UserRole.tenant_admin, UserRole.system_admin):
        raise HTTPException(403, "skill health requires admin role")
    skills = db.execute(
        select(SkillPackage).where(
            (SkillPackage.scope == Scope.global_)
            | ((SkillPackage.owner_tenant_id == p.tenant.id))
            | ((SkillPackage.owner_user_id == p.user.id))
        )
    ).scalars().all()
    items = []
    for s in _dedupe_skill_rows(skills):
        d = _skill_to_dict(s, db)
        items.append({
            "id": d["id"],
            "name": d["name"],
            "slug": d["slug"],
            "version": d["version"],
            "status": d["status"],
            "risk_level": d["risk_level"],
            "bound_employee_count": d["bound_employee_count"],
            "health": d["health"],
        })
    return {"items": sorted(items, key=lambda x: (-x["health"]["failure_rate"], x["name"].lower()))}


@app.get("/api/skill-market/reconcile-hermes")
async def reconcile_hermes_skills(
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    if p.user.role not in (UserRole.tenant_admin, UserRole.system_admin):
        raise HTTPException(403, "skill reconciliation requires admin role")
    sync_error = ""
    try:
        await _sync_installed_hermes_skills(db, p)
    except Exception as exc:
        sync_error = str(exc)[:500]
    target = await hermes_client.resolve_target(db, p.tenant.id)
    hermes_items: list[dict] = []
    try:
        payload = await hermes_client.list_skills(target)
        raw_items = payload.get("data") if isinstance(payload, dict) else []
        hermes_items = raw_items if isinstance(raw_items, list) else []
    except Exception as exc:
        sync_error = sync_error or str(exc)[:500]
    hermes_keys = {
        _hermes_skill_key(str(item.get("name") or item.get("id") or item.get("slug") or "")): item
        for item in hermes_items if isinstance(item, dict)
    }
    openatlas_rows = db.execute(
        select(SkillPackage).where(
            (SkillPackage.scope == Scope.global_)
            | ((SkillPackage.scope == Scope.tenant) & (SkillPackage.owner_tenant_id == p.tenant.id))
            | ((SkillPackage.scope == Scope.user) & (SkillPackage.owner_user_id == p.user.id))
            | ((SkillPackage.scope == Scope.employee) & (SkillPackage.owner_tenant_id == p.tenant.id))
        )
    ).scalars().all()
    openatlas_keys: set[str] = set()
    items: list[dict] = []
    hermes_home = _tenant_hermes_home(db, p.tenant.id)
    for s in _dedupe_skill_rows(openatlas_rows):
        keys = {_hermes_skill_key(s.name), _hermes_skill_key(s.slug), _hermes_skill_key((s.source_ref or "").removeprefix("hermes:"))}
        keys.discard("")
        openatlas_keys |= keys
        md_path, err = _find_hermes_skill_md(hermes_home, s)
        status = "synced" if md_path else "missing_in_hermes"
        if not (s.source_ref or "").startswith("hermes:") and not md_path:
            status = "openatlas_only"
        items.append({
            "openatlas_skill": _skill_to_dict(s, db),
            "status": status,
            "hermes_path": str(md_path) if md_path else "",
            "message": "" if md_path else err,
        })
    for key, item in hermes_keys.items():
        if key and key not in openatlas_keys:
            items.append({
                "openatlas_skill": None,
                "status": "untracked_in_openatlas",
                "hermes_skill": item,
                "message": "Hermes has this skill but OpenAtlas has not registered governance metadata yet.",
            })
    return {"items": items, "hermes_sync_error": sync_error}


@app.post("/api/skill-market")
def create_skill(
    body: SkillIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    # Permission: global only by system_admin; tenant by tenant_admin; user by self
    if body.scope == Scope.global_ and p.user.role != UserRole.system_admin:
        raise HTTPException(403, "only system_admin can publish global skills")
    if body.scope == Scope.tenant and p.user.role not in (UserRole.tenant_admin, UserRole.system_admin):
        raise HTTPException(403, "only tenant_admin can publish tenant skills")
    s = SkillPackage(
        scope=body.scope,
        owner_tenant_id=(p.tenant.id if body.scope == Scope.tenant else None),
        owner_user_id=(p.user.id if body.scope == Scope.user else None),
        name=body.name, slug=body.slug, description=body.description,
        category=body.category, version=body.version,
        visibility=body.visibility, mutable=body.mutable,
        created_by=p.user.id,
    )
    db.add(s)
    audit(db, principal=p, action="skill.create", resource_type="skill",
          resource_id=s.id, request=request, extra={"scope": body.scope.value})
    db.commit()
    db.refresh(s)
    return _skill_to_dict(s, db)


# ── Skill ZIP import (Phase 3.7) ───────────────────────────────────────────
_YAML_FRONT_RE = re.compile(r"^---\s*\n(.*?)\n---\s*\n(.*)$", re.DOTALL)
_ZIP_MAX_BYTES = 10 * 1024 * 1024  # 10 MB cap — a SKILL.md + small scripts is plenty


def _parse_skill_md(content: str) -> dict:
    """Parse YAML frontmatter from a SKILL.md file. Returns {meta: dict, body: str}.

    Accepts a tiny subset of YAML (key: value, key: [a,b], key: {a:b}) — no PyYAML
    dependency, no security surprises. Hermes skills have small simple frontmatters.
    """
    m = _YAML_FRONT_RE.match(content)
    if not m:
        return {"meta": {}, "body": content}
    raw_meta, body = m.group(1), m.group(2)
    meta: dict = {}
    current_list_key: str | None = None
    for line in raw_meta.splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        if line.startswith("  -") and current_list_key:
            # append to current list — single word values
            v = line.strip()[1:].strip()
            if v and (v.startswith('"') or v.startswith("'")) and v[0] == v[-1]:
                v = v[1:-1]
            meta.setdefault(current_list_key, []).append(v)
            continue
        if ":" not in line:
            continue
        k, _, v = line.partition(":")
        k = k.strip()
        v = v.strip()
        if v == "":
            current_list_key = k
            meta[k] = []
            continue
        current_list_key = None
        if v.startswith("[") and v.endswith("]"):
            inner = v[1:-1]
            meta[k] = [
                x.strip().strip('"').strip("'")
                for x in inner.split(",") if x.strip()
            ]
        elif v.startswith('"') and v.endswith('"'):
            meta[k] = v[1:-1]
        elif v.startswith("'") and v.endswith("'"):
            meta[k] = v[1:-1]
        else:
            meta[k] = v
    return {"meta": meta, "body": body}


def _slugify_zip_skill(name: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9._-]+", "-", name).strip("-").lower()
    return s or "skill"


@app.post("/api/skill-market/import")
async def import_skill_zip(
    request: Request,
    file: UploadFile = File(...),
    scope: str = "user",
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    """Upload a SKILL.zip, validate, extract to disk, create SkillPackage.

    Required ZIP structure:
      SKILL.zip
        SKILL.md         (YAML frontmatter with at least: name, version)
        ... (any other files: scripts/, references/, assets/ are kept as-is)

    Behavior:
      - Extracts to OPENATLAS_HOME/tenant-skills/{tenant_slug}/{skill_slug}/
      - scope: 'user' (default, anyone) | 'tenant' (tenant_admin+ only)
      - duplicate (name+version already in this tenant) → 409
    """
    # ── Permission
    if scope == "global":
        if p.user.role != UserRole.system_admin:
            raise HTTPException(403, "only system_admin can import global skills")
    elif scope == "tenant":
        if p.user.role not in (UserRole.tenant_admin, UserRole.system_admin):
            raise HTTPException(403, "only tenant_admin can import tenant-scoped skills")
    elif scope != "user":
        raise HTTPException(400, "scope must be user | tenant | global")
    if not file.filename or not file.filename.lower().endswith(".zip"):
        raise HTTPException(400, "file must be a .zip")

    raw = await file.read()
    if len(raw) > _ZIP_MAX_BYTES:
        raise HTTPException(400, f"zip too large ({len(raw)} > {_ZIP_MAX_BYTES})")
    if len(raw) == 0:
        raise HTTPException(400, "empty file")

    # ── Validate zip structure
    try:
        zf = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile:
        raise HTTPException(400, "not a valid zip file")

    skill_md_names = [n for n in zf.namelist() if n.endswith("SKILL.md")]
    if not skill_md_names:
        zf.close()
        raise HTTPException(400, "zip must contain a SKILL.md file")
    skill_md_name = skill_md_names[0]
    try:
        skill_md_content = zf.read(skill_md_name).decode("utf-8")
    except UnicodeDecodeError:
        zf.close()
        raise HTTPException(400, "SKILL.md is not valid utf-8")

    parsed = _parse_skill_md(skill_md_content)
    meta = parsed["meta"]
    if "name" not in meta:
        zf.close()
        raise HTTPException(400, "SKILL.md frontmatter missing required field: name")
    if "version" not in meta:
        zf.close()
        raise HTTPException(400, "SKILL.md frontmatter missing required field: version")

    skill_name = str(meta["name"]).strip()
    skill_version = str(meta["version"]).strip()
    skill_desc = str(meta.get("description", "")).strip()
    skill_category = str(meta.get("category", "general")).strip() or "general"
    if not skill_name:
        zf.close()
        raise HTTPException(400, "skill name is empty")

    skill_slug = _slugify_zip_skill(skill_name)

    # ── Zip-slip check: every member path must resolve inside target dir
    base_extract = (
        Path(str(OPENATLAS_HOME))
        / "tenant-skills"
        / p.tenant.slug
        / f"{skill_slug}__{skill_version}"
    ).resolve()
    for member in zf.namelist():
        member_path = (base_extract / member).resolve()
        if not str(member_path).startswith(str(base_extract)):
            zf.close()
            raise HTTPException(400, f"zip-slip detected: '{member}' escapes target dir")

    # ── Duplicate check (name+version, scoped to tenant)
    existing = db.execute(
        select(SkillPackage).where(
            SkillPackage.name == skill_name,
            SkillPackage.version == skill_version,
        )
    ).scalar_one_or_none()
    if existing:
        # Also check that this tenant has access / owns the same name+version
        if (existing.owner_tenant_id == p.tenant.id) or (existing.scope == Scope.global_):
            zf.close()
            raise HTTPException(409, f"skill '{skill_name}' v{skill_version} already exists")

    # ── Extract
    base_extract.mkdir(parents=True, exist_ok=True)
    zf.extractall(base_extract)
    zf.close()

    # ── Create SkillPackage
    s = SkillPackage(
        scope=(Scope.global_ if scope == "global" else (Scope.tenant if scope == "tenant" else Scope.user)),
        owner_tenant_id=(p.tenant.id if scope in ("tenant", "user") else None),
        owner_user_id=(p.user.id if scope == "user" else None),
        name=skill_name,
        slug=skill_slug,
        description=skill_desc,
        category=skill_category,
        version=skill_version,
        source_ref=f"zip:{str(base_extract)}",
        visibility=("public" if scope == "global" else ("tenant" if scope == "tenant" else "private")),
        mutable=(scope != "global"),
        status="enabled",
        created_by=p.user.id,
    )
    db.add(s)
    audit(
        db, principal=p, action="skill.import_zip", resource_type="skill",
        resource_id=s.id, request=request,
        extra={"scope": s.scope.value, "version": skill_version, "path": str(base_extract),
               "size": len(raw), "files": sum(1 for _ in Path(base_extract).rglob("*") if _.is_file())},
    )
    db.commit()
    db.refresh(s)
    return {
        **_skill_to_dict(s, db),
        "extracted_path": str(base_extract),
        "extracted_files": sum(1 for m in Path(base_extract).rglob("*") if m.is_file()),
    }


@app.post("/api/skill-market/{skill_id}/bind")
def bind_skill(
    skill_id: str, body: SkillBindIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    s = db.get(SkillPackage, skill_id)
    if not s:
        raise HTTPException(404, "skill not found")
    if s.status != "enabled":
        raise HTTPException(400, "only enabled skills can be bound")
    # Permission: must be visible to caller
    if not _skill_visible_to_principal(s, p):
        raise HTTPException(403, "skill not visible to you")
    if body.target_type not in ("employee", "user"):
        raise HTTPException(400, "target_type must be 'employee' or 'user'")
    if body.target_type == "employee":
        emp = db.get(DigitalEmployee, body.target_id)
        if not emp or emp.tenant_id != p.tenant.id:
            raise HTTPException(404, "employee not found")
    if body.target_type == "user":
        target_user = db.get(User, body.target_id)
        is_admin = p.user.role in (UserRole.tenant_admin, UserRole.system_admin)
        if not target_user or target_user.tenant_id != p.tenant.id:
            raise HTTPException(404, "user not found")
        if not is_admin and target_user.id != p.user.id:
            raise HTTPException(403, "users can only bind skills to themselves")
    existing = db.execute(
        select(SkillBinding).where(
            SkillBinding.tenant_id == p.tenant.id,
            SkillBinding.skill_id == skill_id,
            SkillBinding.target_type == body.target_type,
            SkillBinding.target_id == body.target_id,
        )
    ).scalar_one_or_none()
    if existing:
        existing.enabled = True
        existing.binding_mode = body.binding_mode
        audit(db, principal=p, action="skill.bind.reenable", resource_type="skill",
              resource_id=skill_id, request=request,
              extra={"target_type": body.target_type, "target_id": body.target_id})
        db.commit()
        db.refresh(existing)
        return {"id": existing.id, "binding_mode": existing.binding_mode, "enabled": existing.enabled, "locked": existing.locked}
    b = SkillBinding(
        tenant_id=p.tenant.id, skill_id=skill_id, target_type=body.target_type,
        target_id=body.target_id, binding_mode=body.binding_mode,
        enabled=True, locked=(s.scope == Scope.global_), created_by=p.user.id,
    )
    db.add(b)
    audit(db, principal=p, action="skill.bind", resource_type="skill",
          resource_id=skill_id, request=request,
          extra={"target_type": body.target_type, "target_id": body.target_id})
    db.commit()
    db.refresh(b)
    return {"id": b.id, "binding_mode": b.binding_mode, "enabled": b.enabled, "locked": b.locked}


class SkillUpdateIn(BaseModel):
    name: str | None = None
    description: str | None = None
    category: str | None = None
    version: str | None = None
    visibility: str | None = None
    mutable: bool | None = None
    status: str | None = None  # enabled | disabled | deprecated


def _skill_permission_check(s: SkillPackage, p: Principal, action: str) -> None:
    """Raise 403 if principal cannot mutate the skill."""
    if s.scope == Scope.global_ and p.user.role != UserRole.system_admin:
        raise HTTPException(403, f"only system_admin can {action} global skills")
    if s.scope == Scope.tenant and (
        s.owner_tenant_id != p.tenant.id
        or p.user.role not in (UserRole.tenant_admin, UserRole.system_admin)
    ):
        raise HTTPException(403, f"only owning tenant_admin can {action} this tenant skill")
    if s.scope == Scope.user and s.owner_user_id != p.user.id:
        raise HTTPException(403, f"only the owner can {action} this user skill")
    if s.scope == Scope.employee:
        is_admin = p.user.role in (UserRole.tenant_admin, UserRole.system_admin)
        is_creator = s.created_by == p.user.id
        if s.owner_tenant_id != p.tenant.id or (not is_admin and not is_creator):
            raise HTTPException(403, f"only the creator or tenant admin can {action} this employee skill")


@app.get("/api/skill-market/{skill_id}")
def get_market_skill(
    skill_id: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    s = db.get(SkillPackage, skill_id)
    if not s:
        raise HTTPException(404, "skill not found")
    if not _skill_visible_to_principal(s, p):
        raise HTTPException(404, "skill not found")
    return _skill_to_dict(s, db)


@app.patch("/api/skill-market/{skill_id}")
def update_market_skill(
    skill_id: str, body: SkillUpdateIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    s = db.get(SkillPackage, skill_id)
    if not s:
        raise HTTPException(404, "skill not found")
    _skill_permission_check(s, p, "modify")
    changes = {}
    for field in ("name", "description", "category", "version", "visibility", "mutable", "status"):
        v = getattr(body, field)
        if v is not None:
            setattr(s, field, v)
            changes[field] = v
    if not changes:
        raise HTTPException(400, "no fields to update")
    audit(db, principal=p, action="skill.update", resource_type="skill",
          resource_id=skill_id, request=request, extra=changes)
    db.commit()
    return _skill_to_dict(s, db)


@app.post("/api/skill-market/{skill_id}/publish")
def publish_market_skill(
    skill_id: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.5 — publish a new version. Bumps version + sets status=enabled."""
    s = db.get(SkillPackage, skill_id)
    if not s:
        raise HTTPException(404, "skill not found")
    _skill_permission_check(s, p, "publish")
    s.status = "enabled"
    # bump version: x.y.z -> x.y.(z+1) if it parses
    try:
        parts = s.version.split(".")
        if len(parts) == 3 and all(x.isdigit() for x in parts):
            parts[2] = str(int(parts[2]) + 1)
            s.version = ".".join(parts)
    except Exception:
        pass
    audit(db, principal=p, action="skill.publish", resource_type="skill",
          resource_id=skill_id, request=request, extra={"version": s.version})
    db.commit()
    return _skill_to_dict(s, db)


@app.post("/api/skill-market/{skill_id}/disable")
def disable_market_skill(
    skill_id: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    s = db.get(SkillPackage, skill_id)
    if not s:
        raise HTTPException(404, "skill not found")
    _skill_permission_check(s, p, "disable")
    s.status = "disabled"
    # also disable all bindings
    db.query(SkillBinding).filter(SkillBinding.skill_id == skill_id).update({"enabled": False})
    audit(db, principal=p, action="skill.disable", resource_type="skill",
          resource_id=skill_id, request=request)
    db.commit()
    return _skill_to_dict(s, db)


class SkillForkIn(BaseModel):
    target_scope: Scope = Scope.user  # "user" or "employee"
    target_id: str | None = None  # required if target_scope=employee


@app.post("/api/skill-market/{skill_id}/fork")
def fork_market_skill(
    skill_id: str, body: SkillForkIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.5 / 3.9 — fork a global/tenant skill into a user-owned or
    employee-owned copy. The fork is mutable; the original is not touched.

    If target_scope=employee, target_id is required and a SkillBinding
    is created automatically so the forked skill is immediately available
    to the employee.
    """
    src = db.get(SkillPackage, skill_id)
    if not src:
        raise HTTPException(404, "skill not found")
    # Permission: must be able to see the source
    if src.scope == Scope.tenant and src.owner_tenant_id != p.tenant.id:
        raise HTTPException(404, "skill not found")
    if src.scope == Scope.user and src.owner_user_id != p.user.id:
        raise HTTPException(404, "skill not found")
    if body.target_scope not in (Scope.user, Scope.employee):
        raise HTTPException(400, "fork target must be 'user' or 'employee'")
    # For employee forks, target_id (employee uuid) is required and must be in this tenant
    if body.target_scope == Scope.employee:
        if not body.target_id:
            raise HTTPException(400, "target_id (employee uuid) is required for employee forks")
        emp = db.get(DigitalEmployee, body.target_id)
        if not emp or emp.tenant_id != p.tenant.id:
            raise HTTPException(404, "employee not found")
    new = SkillPackage(
        scope=body.target_scope,
        owner_tenant_id=(p.tenant.id if body.target_scope in (Scope.tenant, Scope.user, Scope.employee) else None),
        owner_user_id=(p.user.id if body.target_scope == Scope.user else None),
        name=f"{src.name} (fork)",
        slug=f"{src.slug}-fork-{p.user.id[:6]}",
        description=src.description,
        category=src.category,
        version="1.0.0",
        visibility="private" if body.target_scope in (Scope.user, Scope.employee) else src.visibility,
        mutable=True,
        status="enabled",
        source_ref=src.id,
        created_by=p.user.id,
    )
    db.add(new)
    db.flush()
    audit(db, principal=p, action="skill.fork", resource_type="skill",
          resource_id=new.id, request=request,
          extra={"source": src.id, "scope": body.target_scope.value})
    # Phase 3.9: if forking to employee, auto-create a SkillBinding so it's immediately usable
    if body.target_scope == Scope.employee and body.target_id:
        binding = SkillBinding(
            tenant_id=p.tenant.id, skill_id=new.id,
            target_type="employee", target_id=body.target_id,
            binding_mode="copied", enabled=True, locked=False,
            created_by=p.user.id,
        )
        db.add(binding)
        db.flush()
        audit(db, principal=p, action="skill.bind", resource_type="skill",
              resource_id=new.id, request=request,
              extra={"target_type": "employee", "target_id": body.target_id,
                     "via": "fork"})
    db.commit()
    db.refresh(new)
    return _skill_to_dict(new, db)


# ── Skill Bindings list (Phase 3.9) ────────────────────────────────────────
@app.get("/api/skill-bindings")
def list_skill_bindings(
    target_type: str | None = Query(None, description="employee | user | tenant"),
    target_id: str | None = Query(None),
    skill_id: str | None = Query(None),
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    """List skill bindings visible to the caller, optionally filtered.

    Used by:
      - EmployeeDetail.tsx  → target_type=employee, target_id={emp_id}
      - Skills.tsx          → (no filter) — show all bindings so user sees what's bound where
      - MemoryCenter-style audits later
    """
    q = select(SkillBinding).where(SkillBinding.tenant_id == p.tenant.id)
    if target_type:
        q = q.where(SkillBinding.target_type == target_type)
    if target_id:
        q = q.where(SkillBinding.target_id == target_id)
    if skill_id:
        q = q.where(SkillBinding.skill_id == skill_id)
    rows = db.execute(q.order_by(SkillBinding.created_at.desc())).scalars().all()
    out = []
    for b in rows:
        skill = db.get(SkillPackage, b.skill_id)
        out.append({
            "id": b.id,
            "skill_id": b.skill_id,
            "skill_name": skill.name if skill else "(deleted skill)",
            "skill_version": skill.version if skill else "?",
            "skill_scope": skill.scope.value if skill else "?",
            "skill_category": skill.category if skill else "?",
            "target_type": b.target_type,
            "target_id": b.target_id,
            "binding_mode": b.binding_mode,
            "enabled": b.enabled,
            "locked": b.locked,
            "created_at": b.created_at.isoformat() if b.created_at else None,
        })
    return {"items": out}


@app.delete("/api/skill-bindings/{binding_id}")
def delete_skill_binding(
    binding_id: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.5 — remove a skill binding. Refuses to delete locked bindings."""
    b = db.get(SkillBinding, binding_id)
    if not b:
        raise HTTPException(404, "binding not found")
    if b.tenant_id != p.tenant.id:
        raise HTTPException(404, "binding not found")
    if b.locked:
        raise HTTPException(400, "binding is locked (inherited from global skill)")
    # permission: target owner or admin
    is_admin = p.user.role in (UserRole.tenant_admin, UserRole.system_admin)
    if not is_admin and b.created_by != p.user.id:
        raise HTTPException(403, "only the binder or an admin can remove this binding")
    audit(db, principal=p, action="skill.unbind", resource_type="binding",
          resource_id=binding_id, request=request, extra={"skill_id": b.skill_id})
    db.delete(b)
    db.commit()
    return {"ok": True, "deleted": binding_id}


# ── Memory Center ──────────────────────────────────────────────────────────
@app.get("/api/memories")
def list_memories(
    scope: str | None = Query(default=None),
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    # All memories the user can READ (global + tenant + own user + employees they own)
    rows = db.execute(
        select(MemoryEntry).where(
            MemoryEntry.tenant_id == p.tenant.id,
            (MemoryEntry.scope == Scope.global_)
            | (MemoryEntry.scope == Scope.tenant)
            | ((MemoryEntry.scope == Scope.user) & (MemoryEntry.owner_user_id == p.user.id))
            | (MemoryEntry.scope == Scope.employee)  # MVP: anyone in tenant sees employee memory metadata
        ).order_by(MemoryEntry.priority.desc(), MemoryEntry.created_at.desc())
    ).scalars().all()
    if scope:
        rows = [m for m in rows if m.scope.value == scope]
    return {"items": [_memory_to_dict(m) for m in rows]}


@app.post("/api/memories")
def create_memory(
    body: MemoryIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    if body.scope == Scope.global_ and p.user.role != UserRole.system_admin:
        raise HTTPException(403, "only system_admin can write global memory")
    if body.scope == Scope.tenant and p.user.role not in (UserRole.tenant_admin, UserRole.system_admin):
        raise HTTPException(403, "only tenant_admin can write tenant memory")
    if body.scope == Scope.employee:
        emp = db.get(DigitalEmployee, body.employee_id) if body.employee_id else None
        if not emp or emp.tenant_id != p.tenant.id:
            raise HTTPException(400, "employee_id required and must be in this tenant")
    m = MemoryEntry(
        tenant_id=p.tenant.id,
        scope=body.scope,
        owner_user_id=(p.user.id if body.scope == Scope.user else None),
        employee_id=body.employee_id,
        title=body.title, content=body.content,
        tags=",".join(body.tags), visibility=body.visibility,
        mutable=body.mutable, priority=body.priority,
        created_by=p.user.id,
    )
    db.add(m)
    audit(db, principal=p, action="memory.create", resource_type="memory",
          resource_id="new", request=request,
          extra={"scope": body.scope.value, "title": body.title})
    db.commit()
    db.refresh(m)
    return _memory_to_dict(m)


@app.patch("/api/memories/{mid}")
def patch_memory(
    mid: str, body: MemoryPatch, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    m = db.get(MemoryEntry, mid)
    if not m or m.tenant_id != p.tenant.id:
        raise HTTPException(404, "memory not found")
    # Permission: only owner OR (admin for tenant/global)
    is_owner = (m.scope == Scope.user and m.owner_user_id == p.user.id)
    is_admin = (m.scope == Scope.tenant and p.user.role in (UserRole.tenant_admin, UserRole.system_admin))
    is_sysadmin = (m.scope == Scope.global_ and p.user.role == UserRole.system_admin)
    is_emp_owner = (m.scope == Scope.employee and m.created_by == p.user.id)
    if not (is_owner or is_admin or is_sysadmin or is_emp_owner):
        raise HTTPException(403, "no permission to edit this memory")
    if not m.mutable:
        raise HTTPException(409, "memory is locked (immutable)")
    for fld in ("title", "content", "tags", "priority", "visibility"):
        val = getattr(body, fld, None)
        if val is not None:
            if fld == "tags":
                m.tags = ",".join(val)
            else:
                setattr(m, fld, val)
    m.version += 1
    audit(db, principal=p, action="memory.update", resource_type="memory",
          resource_id=m.id, request=request)
    db.commit()
    db.refresh(m)
    return _memory_to_dict(m)


@app.post("/api/memories/{mid}/archive")
def archive_memory(
    mid: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    m = db.get(MemoryEntry, mid)
    if not m or m.tenant_id != p.tenant.id:
        raise HTTPException(404, "memory not found")
    m.status = "archived"
    audit(db, principal=p, action="memory.archive", resource_type="memory",
          resource_id=m.id, request=request)
    db.commit()
    return {"ok": True}


# Phase 3.5 — NOTE: /api/memories/effective is registered BEFORE
# /api/memories/{mid} so FastAPI's path matcher routes 'effective' as
# the keyword, not as a memory id.
@app.get("/api/memories/effective")
def effective_memories(
    employee_id: str | None = Query(default=None),
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    eff = resolve_effective_memories(db, tenant_id=p.tenant.id,
                                     user_id=p.user.id, employee_id=employee_id)
    return {"items": eff}


@app.get("/api/memories/{mid}")
def get_memory(
    mid: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    m = db.get(MemoryEntry, mid)
    if not m or m.tenant_id != p.tenant.id:
        raise HTTPException(404, "memory not found")
    return _memory_to_dict(m)


class MemoryForkIn(BaseModel):
    target_scope: Scope  # user | employee
    employee_id: str | None = None


@app.post("/api/memories/{mid}/fork")
def fork_memory(
    mid: str, body: MemoryForkIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.5 — fork a global/tenant memory into a user-owned or
    employee-owned copy. Used for 'this is useful but I want my own version'."""
    src = db.get(MemoryEntry, mid)
    if not src or src.tenant_id != p.tenant.id:
        raise HTTPException(404, "memory not found")
    if body.target_scope not in (Scope.user, Scope.employee):
        raise HTTPException(400, "target_scope must be 'user' or 'employee'")
    if body.target_scope == Scope.employee:
        emp = db.get(DigitalEmployee, body.employee_id) if body.employee_id else None
        if not emp or emp.tenant_id != p.tenant.id:
            raise HTTPException(400, "employee_id required and must be in this tenant")
    new = MemoryEntry(
        tenant_id=p.tenant.id,
        scope=body.target_scope,
        owner_user_id=(p.user.id if body.target_scope == Scope.user else None),
        employee_id=(body.employee_id if body.target_scope == Scope.employee else None),
        title=f"{src.title} (fork)",
        content=src.content,
        tags=src.tags,
        visibility="private" if body.target_scope in (Scope.user, Scope.employee) else src.visibility,
        mutable=True,
        priority=src.priority,
        status="active",
        version=1,
        created_by=p.user.id,
    )
    db.add(new)
    db.flush()
    audit(db, principal=p, action="memory.fork", resource_type="memory",
          resource_id=new.id, request=request,
          extra={"source": src.id, "scope": body.target_scope.value})
    db.commit()
    db.refresh(new)
    return _memory_to_dict(new)


class MemoryBindIn(BaseModel):
    target_type: str  # tenant | user | employee | session
    target_id: str
    injection_mode: str = "always"  # always | on_demand | manual
    priority: int = 50


@app.post("/api/memories/{mid}/bind")
def bind_memory(
    mid: str, body: MemoryBindIn, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.5 — bind a memory to a target. The target determines
    scope: 'tenant' = all users in tenant, 'user' = just me,
    'employee' = bind to one employee, 'session' = bind for one session."""
    m = db.get(MemoryEntry, mid)
    if not m or m.tenant_id != p.tenant.id:
        raise HTTPException(404, "memory not found")
    if body.target_type not in ("tenant", "user", "employee", "session"):
        raise HTTPException(400, "invalid target_type")
    # permission: only memory owner OR admin
    is_owner = (m.scope == Scope.user and m.owner_user_id == p.user.id)
    is_admin = p.user.role in (UserRole.tenant_admin, UserRole.system_admin)
    is_emp_owner = (m.scope == Scope.employee and m.created_by == p.user.id)
    if not (is_owner or is_admin or is_emp_owner):
        raise HTTPException(403, "no permission to bind this memory")
    b = MemoryBinding(
        tenant_id=p.tenant.id, memory_id=mid,
        target_type=body.target_type, target_id=body.target_id,
        injection_mode=body.injection_mode, priority=body.priority,
        enabled=True, locked=False,
    )
    db.add(b)
    audit(db, principal=p, action="memory.bind", resource_type="memory",
          resource_id=mid, request=request,
          extra={"target_type": body.target_type, "target_id": body.target_id})
    db.commit()
    db.refresh(b)
    return {"id": b.id, "target_type": b.target_type, "target_id": b.target_id,
            "injection_mode": b.injection_mode, "priority": b.priority, "enabled": b.enabled}


@app.delete("/api/memory-bindings/{binding_id}")
def delete_memory_binding(
    binding_id: str, request: Request,
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    b = db.get(MemoryBinding, binding_id)
    if not b or b.tenant_id != p.tenant.id:
        raise HTTPException(404, "binding not found")
    if b.locked:
        raise HTTPException(400, "binding is locked (immutable inherited memory)")
    # permission: binding tenant admin or system admin
    if p.user.role not in (UserRole.tenant_admin, UserRole.system_admin):
        raise HTTPException(403, "only tenant_admin can remove memory bindings")
    audit(db, principal=p, action="memory.unbind", resource_type="binding",
          resource_id=binding_id, request=request, extra={"memory_id": b.memory_id})
    db.delete(b)
    db.commit()
    return {"ok": True, "deleted": binding_id}


# ── Dashboard ──────────────────────────────────────────────────────────────
@app.get("/api/dashboard/me")
def dashboard_me(
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    n_emp = db.execute(
        select(DigitalEmployee).where(
            DigitalEmployee.tenant_id == p.tenant.id,
            DigitalEmployee.status == EmployeeStatus.active,
        )
    ).scalars().all()
    n_sessions = db.execute(
        select(SessionRecord).where(
            SessionRecord.tenant_id == p.tenant.id,
            SessionRecord.user_id == p.user.id,
        )
    ).scalars().all()
    n_memory = db.execute(
        select(MemoryEntry).where(
            MemoryEntry.tenant_id == p.tenant.id,
            MemoryEntry.owner_user_id == p.user.id,
            MemoryEntry.status == "active",
        )
    ).scalars().all()
    metrics = _tenant_dashboard_metrics(db, p.tenant.id, user_id=p.user.id)
    return {
        "employees": len(n_emp),
        "sessions": len(n_sessions),
        "memories": len(n_memory),
        **metrics,
        "tenant": {"id": p.tenant.id, "slug": p.tenant.slug, "name": p.tenant.name},
        "user": {"id": p.user.id, "email": p.user.email, "role": p.user.role.value},
    }


@app.get("/api/dashboard/tenant")
def dashboard_tenant(
    p: Principal = Depends(get_principal), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.5 — tenant_admin sees their own tenant metrics.
    system_admin sees the tenant they're currently bound to.
    """
    tid = p.tenant.id
    users = db.query(User).filter(User.tenant_id == tid).count()
    emps = db.query(DigitalEmployee).filter(
        DigitalEmployee.tenant_id == tid,
        DigitalEmployee.status == EmployeeStatus.active,
    ).count()
    sessions = db.query(SessionRecord).filter(
        SessionRecord.tenant_id == tid,
    ).count()
    memories = db.query(MemoryEntry).filter(
        MemoryEntry.tenant_id == tid, MemoryEntry.status == "active",
    ).count()
    skills = db.query(SkillPackage).filter(
        (SkillPackage.scope == Scope.global_)
        | ((SkillPackage.scope == Scope.tenant) & (SkillPackage.owner_tenant_id == tid))
        | ((SkillPackage.scope == Scope.user) & (SkillPackage.owner_user_id == p.user.id))
    ).count()
    runtime = (
        db.query(HermesRuntime)
        .filter(HermesRuntime.tenant_id == tid)
        .order_by(HermesRuntime.created_at.desc())
        .first()
    )
    metrics = _tenant_dashboard_metrics(db, tid)
    skill_usage = []
    for b in db.query(SkillBinding).filter(SkillBinding.tenant_id == tid, SkillBinding.enabled == True).all():  # noqa: E712
        skill = db.get(SkillPackage, b.skill_id)
        if skill:
            skill_usage.append({
                "id": skill.id,
                "name": skill.name,
                "scope": skill.scope.value,
                "target_type": b.target_type,
                "target_id": b.target_id,
            })
    session_rows = db.query(SessionRecord).filter(SessionRecord.tenant_id == tid, SessionRecord.archived == False).all()  # noqa: E712
    stale_sessions = sum(1 for s in session_rows if _is_stale_running_session(s))
    running_sessions = sum(1 for s in session_rows if (s.task_status or "") == "running")
    context_sessions = {
        r.session_id for r in db.query(ContextInjection.session_id).filter(ContextInjection.tenant_id == tid).distinct().all()
    }
    artifact_sessions = {
        r.session_id for r in db.query(TaskArtifact.session_id).filter(TaskArtifact.tenant_id == tid, TaskArtifact.archived == False).distinct().all()  # noqa: E712
    }
    active_session_count = max(1, len(session_rows))
    context_coverage = round(len(context_sessions) / active_session_count, 3)
    artifact_coverage = round(len(artifact_sessions) / active_session_count, 3)
    risk_items: list[dict[str, Any]] = []
    score = 100
    runtime_ok = runtime and str(runtime.status) in {"RuntimeStatus.running", "running", "starting"}
    if not runtime_ok:
        score -= 18
        risk_items.append({"code": "runtime_not_running", "severity": "critical", "message": "Hermes Runtime 未处于 running/starting 状态。"})
    if metrics["failure_rate"] > 0.08:
        score -= 14
        risk_items.append({"code": "high_task_failure_rate", "severity": "warning", "message": f"任务失败率 {metrics['failure_rate']:.1%} 偏高。"})
    if stale_sessions:
        score -= min(24, stale_sessions * 8)
        risk_items.append({"code": "stale_sessions", "severity": "critical", "message": f"{stale_sessions} 个运行中会话超过 30 分钟没有更新。"})
    if metrics["skill_failure_rate"] > 0.08:
        score -= 12
        risk_items.append({"code": "skill_failure_rate", "severity": "warning", "message": f"Skill 失败率 {metrics['skill_failure_rate']:.1%} 偏高。"})
    if context_coverage < 0.5 and len(session_rows) >= 3:
        score -= 10
        risk_items.append({"code": "low_context_coverage", "severity": "warning", "message": "多数会话缺少文件/Skill/记忆注入追踪。"})
    if artifact_coverage < 0.25 and len(session_rows) >= 3:
        score -= 8
        risk_items.append({"code": "low_artifact_coverage", "severity": "info", "message": "会话交付物覆盖率偏低，任务闭环感不足。"})
    maturity = {
        "score": max(0, min(100, score)),
        "level": "production_ready" if score >= 90 else ("beta" if score >= 75 else "pilot"),
        "risk_items": risk_items,
        "signals": {
            "runtime_ok": bool(runtime_ok),
            "running_sessions": running_sessions,
            "stale_sessions": stale_sessions,
            "context_coverage": context_coverage,
            "artifact_coverage": artifact_coverage,
            "skill_failure_rate": metrics["skill_failure_rate"],
            "task_failure_rate": metrics["failure_rate"],
        },
    }
    return {
        "tenant": {"id": tid, "slug": p.tenant.slug, "name": p.tenant.name,
                   "plan": p.tenant.plan, "max_sessions": p.tenant.max_sessions,
                   "max_employees": p.tenant.max_employees},
        "users": users,
        "employees": emps,
        "sessions": sessions,
        "memories": memories,
        "skills_available": skills,
        "token_usage": {
            "input": metrics["input_tokens"],
            "output": metrics["output_tokens"],
            "total": metrics["total_tokens"],
        },
        "files": {"count": metrics["files"], "bytes": metrics["file_bytes"], "expired": metrics["expired_files"]},
        "artifacts": {"count": metrics["artifacts"], "kinds": metrics["artifact_kinds"]},
        "task_status_counts": metrics["task_status_counts"],
        "conversation_usage": {
            "messages": metrics["messages"],
            "input_chars": metrics["input_chars"],
            "output_chars": metrics["output_chars"],
        },
        "failure_rate": metrics["failure_rate"],
        "skill_health": {
            "runs": metrics["skill_runs"],
            "total_runs": metrics["skill_runs_total"],
            "window": metrics["skill_run_window"],
            "failures": metrics["skill_failures"],
            "failure_rate": metrics["skill_failure_rate"],
        },
        "skill_usage": skill_usage[:50],
        "gateway": {
            "status": runtime.status.value if runtime and hasattr(runtime.status, "value") else (runtime.status if runtime else "stopped"),
            "port": runtime.port if runtime else None,
            "pid": runtime.pid if runtime else None,
            "hermes_home": runtime.hermes_home_path if runtime else None,
            "health_checked_at": runtime.health_checked_at.isoformat() if runtime and runtime.health_checked_at else None,
        },
        "maturity": maturity,
    }


@app.get("/api/dashboard/system")
def dashboard_system(
    p: Principal = Depends(require_system_admin), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.5 — system_admin only. Cross-tenant aggregate."""
    tenants = db.query(Tenant).count()
    active_tenants = db.query(Tenant).filter(Tenant.status == TenantStatus.active).count()
    users = db.query(User).count()
    emps = db.query(DigitalEmployee).filter(
        DigitalEmployee.status == EmployeeStatus.active,
    ).count()
    sessions = db.query(SessionRecord).count()
    memories = db.query(MemoryEntry).filter(MemoryEntry.status == "active").count()
    skills = db.query(SkillPackage).count()
    bindings = db.query(SkillBinding).count()
    files = db.query(FileAsset).count()
    file_bytes = sum(int(f.size or 0) for f in db.query(FileAsset).all())
    messages = db.query(MessageRecord).all()
    input_tokens = sum(int(m.input_tokens or 0) for m in messages)
    output_tokens = sum(int(m.output_tokens or 0) for m in messages)
    runtimes = db.query(HermesRuntime).all()
    running_gw = sum(1 for r in runtimes if r.status in ("running", "starting"))
    error_gw = sum(1 for r in runtimes if r.status == "error")
    crashed_gw = sum(1 for r in runtimes if r.status == "crashed")
    # skill usage top 5
    from collections import Counter
    skill_counter = Counter(b.skill_id for b in db.query(SkillBinding).all())
    top_skill_ids = [sid for sid, _ in skill_counter.most_common(5)]
    top_skills = []
    for sid in top_skill_ids:
        s = db.get(SkillPackage, sid)
        if s:
            top_skills.append({"id": sid, "name": s.name, "scope": s.scope.value, "bindings": skill_counter[sid]})
    return {
        "tenants_total": tenants,
        "tenants_active": active_tenants,
        "users_total": users,
        "employees_total": emps,
        "sessions_total": sessions,
        "memories_total": memories,
        "skills_total": skills,
        "skill_bindings_total": bindings,
        "files_total": files,
        "file_bytes_total": file_bytes,
        "token_usage": {
            "input": input_tokens,
            "output": output_tokens,
            "total": input_tokens + output_tokens,
        },
        "gateways": {
            "total": len(runtimes),
            "running": running_gw,
            "error": error_gw,
            "crashed": crashed_gw,
            "stopped": len(runtimes) - running_gw - error_gw - crashed_gw,
        },
        "top_skills": top_skills,
    }


@app.get("/api/audit")
def list_audit(
    p: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
    limit: int = 50,
    action: str | None = None,
    resource_type: str | None = None,
    resource_id: str | None = None,
    user_id: str | None = None,
    employee_id: str | None = None,
    skill_id: str | None = None,
    file_id: str | None = None,
    session_id: str | None = None,
    since: str | None = None,
    until: str | None = None,
    q: str | None = None,
    fmt: str | None = None,
):
    """审计日志: spec §3.5 要求, 记录创建员工/删除员工/创建会话/聊天/任务/Skill 引入/Skill 绑定/Memory 修改/Memory 注入。
    支持 ?action=&resource_type=&user_id=&since=&until=&q= 过滤。
    system_admin 默认跨租户;tenant_admin 仅本租户。
    fmt=csv 触发下载。"""
    from fastapi.responses import StreamingResponse
    import io
    import csv as csv_mod
    from datetime import datetime as _dt

    if p.user.role.value not in ("tenant_admin", "system_admin"):
        raise HTTPException(403, "audit log requires admin role")

    stmt = select(AuditLog)
    if p.user.role.value == "tenant_admin":
        stmt = stmt.where(AuditLog.tenant_id == p.tenant.id)
    if action:
        stmt = stmt.where(AuditLog.action == action)
    if resource_type:
        stmt = stmt.where(AuditLog.resource_type == resource_type)
    if resource_id:
        stmt = stmt.where(AuditLog.resource_id == resource_id)
    if user_id:
        stmt = stmt.where(AuditLog.user_id == user_id)
    for trace_id in (employee_id, skill_id, file_id, session_id):
        if trace_id:
            pat = f"%{trace_id}%"
            stmt = stmt.where(or_(AuditLog.resource_id == trace_id, AuditLog.extra.ilike(pat)))
    if since:
        try:
            since_dt = _dt.fromisoformat(since.replace("Z", "+00:00"))
            stmt = stmt.where(AuditLog.created_at >= since_dt)
        except Exception:
            raise HTTPException(400, f"since must be ISO-8601, got: {since!r}")
    if until:
        try:
            until_dt = _dt.fromisoformat(until.replace("Z", "+00:00"))
            stmt = stmt.where(AuditLog.created_at <= until_dt)
        except Exception:
            raise HTTPException(400, f"until must be ISO-8601, got: {until!r}")
    if q:
        pat = f"%{q}%"
        stmt = stmt.where(or_(
            AuditLog.action.ilike(pat),
            AuditLog.resource_type.ilike(pat),
            AuditLog.resource_id.ilike(pat),
            AuditLog.user_id.ilike(pat),
            AuditLog.extra.ilike(pat),
        ))
    stmt = stmt.order_by(AuditLog.created_at.desc()).limit(min(limit, 500) if fmt != "csv" else 5000)
    rows = db.execute(stmt).scalars().all()

    out = [
        {
            "id": r.id,
            "tenant_id": r.tenant_id,
            "action": r.action,
            "resource_type": r.resource_type,
            "resource_id": r.resource_id,
            "user_id": r.user_id,
            "request_id": r.request_id,
            "ip": r.ip,
            "user_agent": r.user_agent,
            "metadata": json.loads(r.extra or "{}"),
            "created_at": r.created_at.isoformat() if r.created_at else None,
        }
        for r in rows
    ]

    if fmt == "csv":
        buf = io.StringIO()
        if out:
            w = csv_mod.DictWriter(buf, fieldnames=list(out[0].keys()), extrasaction="ignore")
            w.writeheader()
            for row in out:
                row2 = {**row, "metadata": json.dumps(row["metadata"], ensure_ascii=False)}
                w.writerow(row2)
        else:
            buf.write("id,tenant_id,action,resource_type,resource_id,user_id,created_at\n")
        buf.seek(0)
        return StreamingResponse(
            iter([buf.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="openatlas-audit-{_dt.utcnow().strftime("%Y%m%dT%H%M%SZ")}.csv"'},
        )
    return out


# ── Phase 2/3: Multi-tenant admin ───────────────────────────────────────────
# (deps defined near top: require_system_admin, require_tenant_or_system_admin, db_query_tenant)


class TenantCreateIn(BaseModel):
    slug: str = Field(min_length=2, max_length=64, pattern=r"^[a-z0-9-]+$")
    name: str = Field(min_length=1, max_length=128)
    admin_email: str = Field(min_length=3, max_length=255)
    admin_password: str = Field(min_length=4, max_length=128)
    plan: str = "starter"
    max_sessions: int | None = None
    max_employees: int | None = None


class TenantUpdateIn(BaseModel):
    name: str | None = None
    status: str | None = None
    plan: str | None = None
    max_sessions: int | None = None
    max_employees: int | None = None


class OrgUnitIn(BaseModel):
    tenant_id: str | None = None
    parent_id: str | None = None
    name: str = Field(min_length=1, max_length=128)
    code: str = Field(default="", max_length=64)
    description: str = ""
    status: str = "active"
    sort_order: int = 0


class OrgUnitPatchIn(BaseModel):
    parent_id: str | None = None
    name: str | None = Field(default=None, max_length=128)
    code: str | None = Field(default=None, max_length=64)
    description: str | None = None
    status: str | None = None
    sort_order: int | None = None


class AdminUserCreateIn(BaseModel):
    tenant_id: str | None = None
    org_unit_id: str | None = None
    email: str = Field(min_length=3, max_length=255)
    username: str = Field(min_length=1, max_length=64)
    password: str = Field(min_length=4, max_length=128)
    role: str = "user"
    is_active: bool = True


class AdminUserPatchIn(BaseModel):
    org_unit_id: str | None = None
    username: str | None = Field(default=None, max_length=64)
    password: str | None = Field(default=None, min_length=4, max_length=128)
    role: str | None = None
    is_active: bool | None = None


class AdminUserMoveIn(BaseModel):
    user_ids: list[str] = Field(default_factory=list)
    org_unit_id: str | None = None


class RolePermissionPatchIn(BaseModel):
    tenant_id: str | None = None
    role: str
    capability: str
    allowed: bool


def _iso(v: Any) -> str | None:
    return v.isoformat() if v else None


def _role(v: Any) -> str:
    return v.value if hasattr(v, "value") else str(v)


ROLE_CAPABILITIES = [
    {"key": "tenant.manage", "label": "租户配置", "group": "治理"},
    {"key": "runtime.manage", "label": "Hermes Runtime", "group": "治理"},
    {"key": "org.manage", "label": "组织管理", "group": "身份"},
    {"key": "user.manage", "label": "用户管理", "group": "身份"},
    {"key": "employee.manage", "label": "员工管理", "group": "业务"},
    {"key": "skill.manage", "label": "Skill 管理", "group": "业务"},
    {"key": "memory.manage", "label": "记忆管理", "group": "业务"},
    {"key": "file.manage", "label": "文件管理", "group": "资源"},
    {"key": "session.view_all", "label": "查看会话范围", "group": "资源"},
    {"key": "audit.view", "label": "审计日志", "group": "审计"},
]

ROLE_PERMISSION_DEFAULTS: dict[str, set[str]] = {
    "system_admin": {item["key"] for item in ROLE_CAPABILITIES},
    "tenant_admin": {
        "runtime.manage",
        "org.manage",
        "user.manage",
        "employee.manage",
        "skill.manage",
        "memory.manage",
        "file.manage",
        "session.view_all",
        "audit.view",
    },
    "user": {"employee.manage", "memory.manage", "file.manage"},
}


def _tenant_ids_for_admin(p: Principal, tenant_id: str | None = None) -> list[str] | None:
    if p.user.role.value == "system_admin":
        return [tenant_id] if tenant_id else None
    if tenant_id and tenant_id != p.tenant.id:
        raise HTTPException(403, "tenant_admin can only access own tenant")
    return [p.tenant.id]


def _assert_tenant_admin_scope(p: Principal, tenant_id: str) -> None:
    if p.user.role.value == "system_admin":
        return
    if p.user.role.value == "tenant_admin" and p.tenant.id == tenant_id:
        return
    raise HTTPException(403, "admin can only access own tenant")


def _serialize_org_unit(db: Session, org: OrganizationUnit) -> dict:
    return {
        "id": org.id,
        "tenant_id": org.tenant_id,
        "parent_id": org.parent_id,
        "name": org.name,
        "code": org.code,
        "description": org.description,
        "status": org.status,
        "sort_order": org.sort_order,
        "user_count": db.query(User).filter(User.org_unit_id == org.id).count(),
        "child_count": db.query(OrganizationUnit).filter(OrganizationUnit.parent_id == org.id).count(),
        "created_at": _iso(org.created_at),
        "updated_at": _iso(org.updated_at),
    }


def _serialize_admin_user(db: Session, user: User) -> dict:
    tenant = db.get(Tenant, user.tenant_id)
    org = db.get(OrganizationUnit, user.org_unit_id) if user.org_unit_id else None
    return {
        "id": user.id,
        "tenant_id": user.tenant_id,
        "tenant_name": tenant.name if tenant else "",
        "org_unit_id": user.org_unit_id,
        "org_unit_name": org.name if org else "",
        "email": user.email,
        "username": user.username,
        "role": _role(user.role),
        "is_active": user.is_active,
        "created_at": _iso(user.created_at),
        "last_login_at": _iso(user.last_login_at),
    }


def _serialize_admin_tenant(db: Session, tenant: Tenant) -> dict:
    runtime = (
        db.query(HermesRuntime)
        .filter(HermesRuntime.tenant_id == tenant.id)
        .order_by(HermesRuntime.created_at.desc())
        .first()
    )
    return {
        "id": tenant.id,
        "slug": tenant.slug,
        "name": tenant.name,
        "status": tenant.status.value if hasattr(tenant.status, "value") else tenant.status,
        "plan": tenant.plan,
        "max_sessions": tenant.max_sessions,
        "max_employees": tenant.max_employees,
        "user_count": db.query(User).filter(User.tenant_id == tenant.id).count(),
        "org_unit_count": db.query(OrganizationUnit).filter(OrganizationUnit.tenant_id == tenant.id).count(),
        "employee_count": db.query(DigitalEmployee).filter(DigitalEmployee.tenant_id == tenant.id).count(),
        "session_count": db.query(SessionRecord).filter(SessionRecord.tenant_id == tenant.id).count(),
        "runtime": {
            "status": runtime.status if runtime else "none",
            "port": runtime.port if runtime else None,
            "gateway_base_url": runtime.gateway_base_url if runtime else "",
            "hermes_home": runtime.hermes_home_path if runtime else "",
            "health_checked_at": _iso(runtime.health_checked_at) if runtime else None,
        },
        "created_at": _iso(tenant.created_at),
        "updated_at": _iso(tenant.updated_at),
    }


def _serialize_role_permissions(db: Session, tenant_id: str | None = None) -> dict:
    overrides = db.query(RolePermissionOverride)
    if tenant_id:
        overrides = overrides.filter(RolePermissionOverride.tenant_id == tenant_id)
    else:
        overrides = overrides.filter(RolePermissionOverride.tenant_id.is_(None))
    override_map = {(r.role, r.capability): r for r in overrides.all()}
    roles = []
    for role in ("system_admin", "tenant_admin", "user"):
        capabilities = []
        for item in ROLE_CAPABILITIES:
            key = item["key"]
            override = override_map.get((role, key))
            default_allowed = key in ROLE_PERMISSION_DEFAULTS.get(role, set())
            allowed = override.allowed if override else default_allowed
            capabilities.append({
                **item,
                "allowed": allowed,
                "default_allowed": default_allowed,
                "overridden": override is not None,
                "updated_at": _iso(override.updated_at) if override else None,
            })
        roles.append({"role": role, "label": {"system_admin": "系统管理员", "tenant_admin": "租户管理员", "user": "普通用户"}[role], "capabilities": capabilities})
    return {"tenant_id": tenant_id, "capabilities": ROLE_CAPABILITIES, "roles": roles}


@app.get("/api/admin/identity/overview")
def admin_identity_overview(
    tenant_id: str | None = Query(default=None),
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    tids = _tenant_ids_for_admin(p, tenant_id)
    users_q = db.query(User)
    orgs_q = db.query(OrganizationUnit)
    tenants_q = db.query(Tenant)
    if tids is not None:
        users_q = users_q.filter(User.tenant_id.in_(tids))
        orgs_q = orgs_q.filter(OrganizationUnit.tenant_id.in_(tids))
        tenants_q = tenants_q.filter(Tenant.id.in_(tids))
    return {
        "tenants": tenants_q.count(),
        "users": users_q.count(),
        "active_users": users_q.filter(User.is_active == True).count(),  # noqa: E712
        "org_units": orgs_q.count(),
        "tenant_role": p.user.role.value,
        "current_tenant_id": p.tenant.id,
    }


@app.get("/api/admin/org-units")
def admin_list_org_units(
    tenant_id: str | None = Query(default=None),
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    tids = _tenant_ids_for_admin(p, tenant_id)
    q = db.query(OrganizationUnit)
    if tids is not None:
        q = q.filter(OrganizationUnit.tenant_id.in_(tids))
    rows = q.order_by(OrganizationUnit.sort_order.asc(), OrganizationUnit.created_at.asc()).all()
    return {"items": [_serialize_org_unit(db, row) for row in rows]}


@app.post("/api/admin/org-units")
def admin_create_org_unit(
    body: OrgUnitIn,
    request: Request,
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    tenant_id = body.tenant_id if p.user.role.value == "system_admin" and body.tenant_id else p.tenant.id
    _assert_tenant_admin_scope(p, tenant_id)
    if body.parent_id:
        parent = db.get(OrganizationUnit, body.parent_id)
        if not parent or parent.tenant_id != tenant_id:
            raise HTTPException(404, "parent org unit not found")
    code = body.code.strip() or f"ORG-{_uuid.uuid4().hex[:6].upper()}"
    if db.query(OrganizationUnit).filter(OrganizationUnit.tenant_id == tenant_id, OrganizationUnit.code == code).first():
        raise HTTPException(409, "org code already exists in tenant")
    org = OrganizationUnit(
        tenant_id=tenant_id,
        parent_id=body.parent_id,
        name=body.name.strip(),
        code=code,
        description=body.description,
        status=body.status,
        sort_order=body.sort_order,
    )
    db.add(org)
    db.flush()
    audit(db, principal=p, action="org.create", resource_type="organization_unit",
          resource_id=org.id, request=request, extra={"name": org.name, "code": org.code})
    db.commit()
    return _serialize_org_unit(db, org)


@app.patch("/api/admin/org-units/{oid}")
def admin_update_org_unit(
    oid: str,
    body: OrgUnitPatchIn,
    request: Request,
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    org = db.get(OrganizationUnit, oid)
    if not org:
        raise HTTPException(404, "org unit not found")
    _assert_tenant_admin_scope(p, org.tenant_id)
    if body.parent_id == oid:
        raise HTTPException(400, "parent cannot be self")
    changes = {}
    if "parent_id" in body.model_fields_set:
        if body.parent_id:
            parent = db.get(OrganizationUnit, body.parent_id)
            if not parent or parent.tenant_id != org.tenant_id:
                raise HTTPException(404, "parent org unit not found")
        org.parent_id = body.parent_id
        changes["parent_id"] = body.parent_id
    for field in ("name", "description", "status", "sort_order"):
        if field in body.model_fields_set:
            v = getattr(body, field)
            if v is not None:
                setattr(org, field, v.strip() if isinstance(v, str) else v)
                changes[field] = v
    if "code" in body.model_fields_set and body.code is not None:
        code = body.code.strip() or org.code
        hit = db.query(OrganizationUnit).filter(
            OrganizationUnit.tenant_id == org.tenant_id,
            OrganizationUnit.code == code,
            OrganizationUnit.id != oid,
        ).first()
        if hit:
            raise HTTPException(409, "org code already exists in tenant")
        org.code = code
        changes["code"] = code
    if not changes:
        raise HTTPException(400, "no fields to update")
    org.updated_at = datetime.now(timezone.utc)
    audit(db, principal=p, action="org.update", resource_type="organization_unit",
          resource_id=org.id, request=request, extra=changes)
    db.commit()
    return _serialize_org_unit(db, org)


@app.delete("/api/admin/org-units/{oid}")
def admin_delete_org_unit(
    oid: str,
    request: Request,
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    org = db.get(OrganizationUnit, oid)
    if not org:
        raise HTTPException(404, "org unit not found")
    _assert_tenant_admin_scope(p, org.tenant_id)
    if db.query(OrganizationUnit).filter(OrganizationUnit.parent_id == oid).first():
        raise HTTPException(400, "cannot delete org unit with children")
    if db.query(User).filter(User.org_unit_id == oid).first():
        raise HTTPException(400, "cannot delete org unit with users")
    db.delete(org)
    audit(db, principal=p, action="org.delete", resource_type="organization_unit",
          resource_id=oid, request=request, extra={"name": org.name})
    db.commit()
    return {"ok": True}


@app.get("/api/admin/users")
def admin_list_users(
    tenant_id: str | None = Query(default=None),
    org_unit_id: str | None = Query(default=None),
    q: str | None = Query(default=None),
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    tids = _tenant_ids_for_admin(p, tenant_id)
    query = db.query(User)
    if tids is not None:
        query = query.filter(User.tenant_id.in_(tids))
    if org_unit_id:
        query = query.filter(User.org_unit_id == org_unit_id)
    if q:
        like = f"%{q.strip()}%"
        query = query.filter(or_(User.email.like(like), User.username.like(like)))
    rows = query.order_by(User.created_at.desc()).all()
    return {"items": [_serialize_admin_user(db, row) for row in rows]}


@app.post("/api/admin/users")
def admin_create_user(
    body: AdminUserCreateIn,
    request: Request,
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    tenant_id = body.tenant_id if p.user.role.value == "system_admin" and body.tenant_id else p.tenant.id
    _assert_tenant_admin_scope(p, tenant_id)
    role = body.role if body.role in {"tenant_admin", "user", "system_admin"} else "user"
    if role == "system_admin" and p.user.role.value != "system_admin":
        raise HTTPException(403, "tenant_admin cannot create system_admin")
    if body.org_unit_id:
        org = db.get(OrganizationUnit, body.org_unit_id)
        if not org or org.tenant_id != tenant_id:
            raise HTTPException(404, "org unit not found")
    if db.query(User).filter(User.tenant_id == tenant_id, User.email == body.email).first():
        raise HTTPException(409, "user email already exists in tenant")
    user = User(
        tenant_id=tenant_id,
        org_unit_id=body.org_unit_id,
        email=body.email.strip(),
        username=body.username.strip(),
        password_hash=hash_password(body.password),
        role=UserRole(role),
        is_active=body.is_active,
    )
    db.add(user)
    db.flush()
    audit(db, principal=p, action="user.create", resource_type="user",
          resource_id=user.id, request=request, extra={"email": user.email, "role": role})
    db.commit()
    return _serialize_admin_user(db, user)


@app.get("/api/admin/users/{uid}/detail")
def admin_get_user_detail(
    uid: str,
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    user = db.get(User, uid)
    if not user:
        raise HTTPException(404, "user not found")
    _assert_tenant_admin_scope(p, user.tenant_id)
    sessions = db.query(SessionRecord).filter(SessionRecord.user_id == uid)
    files = db.query(FileAsset).filter(FileAsset.user_id == uid)
    memories = db.query(MemoryEntry).filter(MemoryEntry.owner_user_id == uid)
    skill_bindings = db.query(SkillBinding).filter(SkillBinding.target_type == "user", SkillBinding.target_id == uid)
    audits = (
        db.query(AuditLog)
        .filter(AuditLog.user_id == uid)
        .order_by(AuditLog.created_at.desc())
        .limit(8)
        .all()
    )
    recent_sessions = (
        sessions.order_by(SessionRecord.updated_at.desc())
        .limit(6)
        .all()
    )
    return {
        "user": _serialize_admin_user(db, user),
        "counts": {
            "sessions": sessions.count(),
            "files": files.count(),
            "memories": memories.count(),
            "skill_bindings": skill_bindings.count(),
            "audit_events": db.query(AuditLog).filter(AuditLog.user_id == uid).count(),
        },
        "recent_sessions": [
            {
                "id": s.id,
                "title": s.title or "新会话",
                "task_status": s.task_status,
                "message_count": s.message_count,
                "updated_at": _iso(s.updated_at),
            }
            for s in recent_sessions
        ],
        "recent_audits": [
            {
                "id": a.id,
                "action": a.action,
                "resource_type": a.resource_type,
                "resource_id": a.resource_id,
                "created_at": _iso(a.created_at),
                "metadata": json.loads(a.extra or "{}"),
            }
            for a in audits
        ],
    }


@app.patch("/api/admin/users/{uid}")
def admin_update_user(
    uid: str,
    body: AdminUserPatchIn,
    request: Request,
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    user = db.get(User, uid)
    if not user:
        raise HTTPException(404, "user not found")
    _assert_tenant_admin_scope(p, user.tenant_id)
    changes = {}
    if "org_unit_id" in body.model_fields_set:
        if body.org_unit_id:
            org = db.get(OrganizationUnit, body.org_unit_id)
            if not org or org.tenant_id != user.tenant_id:
                raise HTTPException(404, "org unit not found")
        user.org_unit_id = body.org_unit_id
        changes["org_unit_id"] = body.org_unit_id
    if body.username is not None:
        user.username = body.username.strip()
        changes["username"] = user.username
    if body.password:
        user.password_hash = hash_password(body.password)
        changes["password"] = "***"
    if body.role is not None:
        if body.role == "system_admin" and p.user.role.value != "system_admin":
            raise HTTPException(403, "tenant_admin cannot assign system_admin")
        if body.role not in {"system_admin", "tenant_admin", "user"}:
            raise HTTPException(400, "invalid role")
        user.role = UserRole(body.role)
        changes["role"] = body.role
    if body.is_active is not None:
        if user.id == p.user.id and body.is_active is False:
            raise HTTPException(400, "cannot disable yourself")
        user.is_active = body.is_active
        changes["is_active"] = body.is_active
    if not changes:
        raise HTTPException(400, "no fields to update")
    audit(db, principal=p, action="user.update", resource_type="user",
          resource_id=user.id, request=request, extra=changes)
    db.commit()
    return _serialize_admin_user(db, user)


@app.post("/api/admin/users/move-org")
def admin_move_users_org(
    body: AdminUserMoveIn,
    request: Request,
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    user_ids = [uid for uid in dict.fromkeys(body.user_ids) if uid]
    if not user_ids:
        raise HTTPException(400, "user_ids required")
    target_org = None
    if body.org_unit_id:
        target_org = db.get(OrganizationUnit, body.org_unit_id)
        if not target_org:
            raise HTTPException(404, "target org unit not found")
        _assert_tenant_admin_scope(p, target_org.tenant_id)
    users = db.query(User).filter(User.id.in_(user_ids)).all()
    if len(users) != len(user_ids):
        raise HTTPException(404, "some users not found")
    tenant_ids = {u.tenant_id for u in users}
    if len(tenant_ids) != 1:
        raise HTTPException(400, "cannot move users across tenants")
    tenant_id = next(iter(tenant_ids))
    _assert_tenant_admin_scope(p, tenant_id)
    if target_org and target_org.tenant_id != tenant_id:
        raise HTTPException(400, "target org belongs to another tenant")
    for user in users:
        user.org_unit_id = body.org_unit_id
    audit(db, principal=p, action="user.move_org", resource_type="user",
          resource_id=",".join(user_ids), request=request,
          extra={"count": len(users), "org_unit_id": body.org_unit_id})
    db.commit()
    return {"ok": True, "moved": len(users), "org_unit_id": body.org_unit_id}


@app.delete("/api/admin/users/{uid}")
def admin_delete_user(
    uid: str,
    request: Request,
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    """Deprovision a user without destroying historical audit/session data."""
    user = db.get(User, uid)
    if not user:
        raise HTTPException(404, "user not found")
    _assert_tenant_admin_scope(p, user.tenant_id)
    if user.id == p.user.id:
        raise HTTPException(400, "cannot deprovision yourself")
    user.is_active = False
    audit(db, principal=p, action="user.deprovision", resource_type="user",
          resource_id=user.id, request=request, extra={"email": user.email})
    db.commit()
    return {"ok": True, "user": _serialize_admin_user(db, user)}


@app.get("/api/admin/tenants")
def admin_list_tenants(
    p: Principal = Depends(require_identity_admin), db: Session = Depends(get_db),
) -> dict:
    q = db.query(Tenant)
    if p.user.role.value != "system_admin":
        q = q.filter(Tenant.id == p.tenant.id)
    rows = q.order_by(Tenant.created_at.desc()).all()
    return {"items": [_serialize_admin_tenant(db, t) for t in rows]}


@app.get("/api/admin/tenants/{tid}/isolation")
def admin_tenant_isolation(
    tid: str,
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    tenant = db.get(Tenant, tid)
    if not tenant:
        raise HTTPException(404, "tenant not found")
    _assert_tenant_admin_scope(p, tid)
    runtime = (
        db.query(HermesRuntime)
        .filter(HermesRuntime.tenant_id == tid)
        .order_by(HermesRuntime.created_at.desc())
        .first()
    )
    hermes_home = runtime.hermes_home_path if runtime else str(_tenant_home(tenant.slug))
    upload_root = str(OPENATLAS_HOME / "uploads" / tid)
    checks = [
        {"key": "tenant_scope", "label": "数据库租户范围", "status": "ok", "detail": f"tenant_id={tid}"},
        {"key": "hermes_home", "label": "Hermes Home 独立目录", "status": "ok" if hermes_home else "warn", "detail": hermes_home},
        {"key": "upload_root", "label": "文件上传根目录", "status": "ok", "detail": upload_root},
        {"key": "runtime_port", "label": "Gateway 端口隔离", "status": "ok" if runtime and runtime.port else "warn", "detail": f":{runtime.port}" if runtime and runtime.port else "未登记"},
        {"key": "auth_boundary", "label": "API 权限边界", "status": "ok", "detail": "系统管理员全局；租户管理员仅本租户；普通用户仅个人资源"},
    ]
    return {
        "tenant": _serialize_admin_tenant(db, tenant),
        "paths": {
            "hermes_home": hermes_home,
            "upload_root": upload_root,
            "runtime_base_url": runtime.gateway_base_url if runtime else "",
        },
        "counts": {
            "users": db.query(User).filter(User.tenant_id == tid).count(),
            "org_units": db.query(OrganizationUnit).filter(OrganizationUnit.tenant_id == tid).count(),
            "employees": db.query(DigitalEmployee).filter(DigitalEmployee.tenant_id == tid).count(),
            "sessions": db.query(SessionRecord).filter(SessionRecord.tenant_id == tid).count(),
            "files": db.query(FileAsset).filter(FileAsset.tenant_id == tid).count(),
            "memories": db.query(MemoryEntry).filter(MemoryEntry.tenant_id == tid).count(),
            "skills": db.query(SkillPackage).filter(
                or_(SkillPackage.owner_tenant_id == tid, SkillPackage.scope == Scope.global_)
            ).count(),
        },
        "checks": checks,
    }


@app.get("/api/admin/permission-matrix")
def admin_permission_matrix(
    tenant_id: str | None = Query(default=None),
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    tids = _tenant_ids_for_admin(p, tenant_id)
    effective_tenant_id = tenant_id or (p.tenant.id if p.user.role.value != "system_admin" else None)
    if tids is not None and effective_tenant_id not in tids:
        raise HTTPException(403, "admin can only access own tenant")
    return _serialize_role_permissions(db, effective_tenant_id)


@app.patch("/api/admin/permission-matrix")
def admin_patch_permission_matrix(
    body: RolePermissionPatchIn,
    request: Request,
    p: Principal = Depends(require_identity_admin),
    db: Session = Depends(get_db),
) -> dict:
    if body.role not in {"system_admin", "tenant_admin", "user"}:
        raise HTTPException(400, "invalid role")
    if body.capability not in {item["key"] for item in ROLE_CAPABILITIES}:
        raise HTTPException(400, "invalid capability")
    tenant_id = body.tenant_id or (None if p.user.role.value == "system_admin" else p.tenant.id)
    if tenant_id:
        _assert_tenant_admin_scope(p, tenant_id)
    if body.role == "system_admin" and p.user.role.value != "system_admin":
        raise HTTPException(403, "tenant_admin cannot change system_admin permissions")
    row = db.query(RolePermissionOverride).filter(
        RolePermissionOverride.tenant_id == tenant_id,
        RolePermissionOverride.role == body.role,
        RolePermissionOverride.capability == body.capability,
    ).first()
    if not row:
        row = RolePermissionOverride(
            tenant_id=tenant_id,
            role=body.role,
            capability=body.capability,
            allowed=body.allowed,
            updated_by=p.user.id,
        )
        db.add(row)
    else:
        row.allowed = body.allowed
        row.updated_by = p.user.id
        row.updated_at = datetime.now(timezone.utc)
    audit(db, principal=p, action="permission.update", resource_type="role_permission",
          resource_id=f"{body.role}:{body.capability}", request=request,
          extra={"tenant_id": tenant_id, "allowed": body.allowed})
    db.commit()
    return _serialize_role_permissions(db, tenant_id)


@app.post("/api/admin/tenants")
def admin_create_tenant(
    body: TenantCreateIn, request: Request,
    p: Principal = Depends(require_system_admin), db: Session = Depends(get_db),
) -> dict:
    """Create a new tenant. Auto-creates a tenant_admin user.

    Phase 2: also auto-provisions a HermesRuntime row + tenant home directory,
    but does NOT start the gateway process — admin must call /start separately.
    """
    if db.execute(select(Tenant).where(Tenant.slug == body.slug)).scalar_one_or_none():
        raise HTTPException(409, f"tenant slug '{body.slug}' already exists")
    if db.execute(select(User).where(User.email == body.admin_email)).scalar_one_or_none():
        raise HTTPException(409, f"user email '{body.admin_email}' already exists")

    tenant = Tenant(
        slug=body.slug, name=body.name,
        status=TenantStatus.active, plan=body.plan,
        max_sessions=body.max_sessions,
        max_employees=body.max_employees,
    )
    db.add(tenant)
    db.flush()
    admin = User(
        tenant_id=tenant.id,
        email=body.admin_email,
        username=body.admin_email.split("@")[0],
        password_hash=hash_password(body.admin_password),
        role=UserRole.tenant_admin,
        is_active=True,
    )
    db.add(admin)
    # Reserve the next available port (Phase 2: simple linear allocator)
    port = _next_hermes_port(db)
    api_key = f"openatlas-{body.slug}-key-{_uuid.uuid4().hex[:8]}"
    # Phase 3.1 — encrypt the API key at rest. The plaintext is only
    # echoed back ONCE in the create response.
    api_key_cipher = encrypt(api_key)
    runtime = HermesRuntime(
        tenant_id=tenant.id,
        runtime_type=RuntimeType.process,
        gateway_base_url=f"http://127.0.0.1:{port}",
        api_key=api_key_cipher,            # Phase 1 legacy field (mirror, encrypted)
        api_key_encrypted=api_key_cipher,  # Phase 3 primary; stored encrypted
        hermes_home_path=str(_tenant_home(body.slug)),
        port=port,
        status=RuntimeStatus.stopped,
    )
    db.add(runtime)
    audit(db, principal=p, action="tenant.create", resource_type="tenant",
          resource_id=tenant.id, request=request,
          extra={"slug": body.slug, "admin_email": body.admin_email, "port": port})
    db.commit()
    # Create the tenant home directory tree on disk (idempotent)
    _tenant_home(body.slug).mkdir(parents=True, exist_ok=True)
    return {
        "id": tenant.id,
        "slug": tenant.slug,
        "name": tenant.name,
        "admin_email": admin.email,
        "admin_password": body.admin_password,  # only returned at create time
        "runtime": {
            "id": runtime.id,
            "port": port,
            "status": runtime.status.value if hasattr(runtime.status, "value") else runtime.status,
            "hermes_home": runtime.hermes_home_path,
        },
    }


@app.get("/api/admin/tenants/{tid}/hermes-runtime")
def admin_get_runtime(
    tid: str, p: Principal = Depends(require_tenant_or_system_admin), db: Session = Depends(get_db),
) -> dict:
    tenant = db.get(Tenant, tid)
    if not tenant:
        raise HTTPException(404, "tenant not found")
    runtime = (
        db.query(HermesRuntime).filter(HermesRuntime.tenant_id == tid)
        .order_by(HermesRuntime.created_at.desc()).first()
    )
    if not runtime:
        return {"tenant_id": tid, "runtime": None, "message": "no runtime registered"}
    # Phase 3.1 — decrypt for display preview; show "***" suffix.
    _plain = decrypt(runtime.api_key_encrypted)
    _preview = (_plain[:12] + "***") if _plain else "***"
    return {
        "tenant_id": tid,
        "tenant_slug": tenant.slug if tenant else None,
        "runtime": {
            "id": runtime.id,
            "port": runtime.port,
            "status": runtime.status.value if hasattr(runtime.status, "value") else runtime.status,
            "gateway_base_url": runtime.gateway_base_url,
            "hermes_home": runtime.hermes_home_path,
            "api_key_preview": _preview,
            "health_checked_at": runtime.health_checked_at.isoformat() if runtime.health_checked_at else None,
            "created_at": runtime.created_at.isoformat() if runtime.created_at else None,
        }
    }


@app.post("/api/admin/tenants/{tid}/hermes-runtime/start")
def admin_start_runtime(
    tid: str, request: Request,
    p: Principal = Depends(require_tenant_or_system_admin), db: Session = Depends(get_db),
) -> dict:
    """Spawn a per-tenant Hermes gateway process in the background.

    Idempotent: if already running, no-op.
    """
    import subprocess
    runtime = (
        db.query(HermesRuntime).filter(HermesRuntime.tenant_id == tid)
        .order_by(HermesRuntime.created_at.desc()).first()
    )
    if not runtime:
        raise HTTPException(404, "no runtime registered for this tenant")
    pid_path = _tenant_pid_path(runtime.hermes_home_path)
    # Already running?
    if pid_path.exists():
        try:
            old_pid = int(pid_path.read_text().strip())
            from app.services.process_utils import is_pid_alive
            if is_pid_alive(old_pid):
                _api_key_plain = decrypt(runtime.api_key_encrypted)
                auth_ok, auth_note = _runtime_auth_probe_sync(runtime.gateway_base_url, _api_key_plain)
                if auth_ok:
                    runtime.pid = old_pid
                    runtime.status = RuntimeStatus.running
                    runtime.health_checked_at = datetime.now(timezone.utc)
                    audit(db, principal=p, action="runtime.start", resource_type="runtime",
                          resource_id=runtime.id, request=request, extra={"noop": True, "pid": old_pid, "auth": auth_note})
                    db.commit()
                    return {"ok": True, "already_running": True, "pid": old_pid,
                            "port": runtime.port, "base_url": runtime.gateway_base_url}
                from app.services.process_utils import terminate_pid
                terminate_pid(old_pid, grace=5.0)
                pid_path.unlink(missing_ok=True)
                audit(db, principal=p, action="runtime.restart.auth_mismatch", resource_type="runtime",
                      resource_id=runtime.id, request=request, extra={"pid": old_pid, "auth": auth_note})
                db.flush()
        except (ProcessLookupError, ValueError):
            pid_path.unlink(missing_ok=True)
    # Spawn via start_tenant.sh (8-layer isolation guard)
    tenant = db.get(Tenant, tid)
    script = _start_tenant_script()
    log_path = _tenant_log_path(runtime.hermes_home_path)
    # Spawn via start_tenant.sh (8-layer isolation guard).
    # Use bash -c with KEY=VAL inline so the child shell sees the exports.
    # We MUST avoid 'shell=True' on Popen (we use redirect for stdout/stderr).
    # So we craft a single argv: env ... bash -c "<full command string>".
    # Phase 3.1 — decrypt the API key for env injection to the child shell.
    _api_key_plain = decrypt(runtime.api_key_encrypted)
    cmd_str = (
        f"HERMES_HOME={runtime.hermes_home_path} "
        f"API_SERVER_HOST=127.0.0.1 "
        f"API_SERVER_PORT={runtime.port} "
        f"API_SERVER_KEY={_api_key_plain} "
        f"OPENATLAS_HOME={str(OPENATLAS_HOME)} "
        f"OPENATLAS_HERMES_AGENT_ROOT={str(OPENATLAS_HERMES_AGENT_ROOT)} "
        f"OPENATLAS_TENANT={tenant.slug} "
        f"bash {script} {runtime.hermes_home_path}"
    )
    env_args = [
        "env", "-u", "ALL_PROXY", "-u", "all_proxy",
        "-u", "HTTP_PROXY", "-u", "HTTPS_PROXY",
        "-u", "http_proxy", "-u", "https_proxy",
        "-u", "SOCKS_PROXY", "-u", "socks_proxy",
        "bash", "-c", cmd_str,
    ]
    log_path.parent.mkdir(parents=True, exist_ok=True)
    fh = open(log_path, "ab")
    proc = subprocess.Popen(
        env_args, stdout=fh, stderr=fh, stdin=subprocess.DEVNULL,
        start_new_session=True, close_fds=True,
    )
    runtime.status = RuntimeStatus.starting
    runtime.health_checked_at = datetime.now(timezone.utc)
    audit(db, principal=p, action="runtime.start", resource_type="runtime",
          resource_id=runtime.id, request=request, extra={"pid": proc.pid, "port": runtime.port})
    db.commit()
    return {"ok": True, "spawned": True, "pid": proc.pid,
            "port": runtime.port, "base_url": runtime.gateway_base_url,
            "log": str(log_path)}


@app.post("/api/admin/tenants/{tid}/hermes-runtime/stop")
def admin_stop_runtime(
    tid: str, request: Request,
    p: Principal = Depends(require_tenant_or_system_admin), db: Session = Depends(get_db),
) -> dict:
    import os as _os, signal
    runtime = (
        db.query(HermesRuntime).filter(HermesRuntime.tenant_id == tid)
        .order_by(HermesRuntime.created_at.desc()).first()
    )
    if not runtime:
        raise HTTPException(404, "no runtime registered for this tenant")
    pid_path = _tenant_pid_path(runtime.hermes_home_path)
    if not pid_path.exists():
        runtime.status = RuntimeStatus.stopped
        db.commit()
        return {"ok": True, "already_stopped": True}
    try:
        pid = int(pid_path.read_text().strip())
        from app.services.process_utils import terminate_pid
        terminate_pid(pid, grace=5.0)
    except ProcessLookupError:
        pass
    pid_path.unlink(missing_ok=True)
    runtime.status = RuntimeStatus.stopped
    audit(db, principal=p, action="runtime.stop", resource_type="runtime",
          resource_id=runtime.id, request=request, extra={})
    db.commit()
    return {"ok": True, "stopped": True}


@app.post("/api/admin/tenants/{tid}/hermes-runtime/healthcheck")
async def admin_runtime_healthcheck(
    tid: str, request: Request,
    p: Principal = Depends(require_tenant_or_system_admin), db: Session = Depends(get_db),
) -> dict:
    """Probe the tenant's gateway /health and update runtime.status."""
    from app.services.hermes_client import _client  # trust_env=False (SOCKS5-safe)
    runtime = (
        db.query(HermesRuntime).filter(HermesRuntime.tenant_id == tid)
        .order_by(HermesRuntime.created_at.desc()).first()
    )
    if not runtime:
        raise HTTPException(404, "no runtime registered for this tenant")
    try:
        async with _client(timeout=5) as c:
            # Phase 3.1 — decrypt before using as bearer token
            _api_key = decrypt(runtime.api_key_encrypted)
            r = await c.get(f"{runtime.gateway_base_url}/health",
                            headers={"Authorization": f"Bearer {_api_key}"})
            r.raise_for_status()
            body = r.json()
            auth = await c.get(f"{runtime.gateway_base_url}/v1/models",
                               headers={"Authorization": f"Bearer {_api_key}"})
            auth.raise_for_status()
            pid = None
            pid_path = _tenant_pid_path(runtime.hermes_home_path)
            if pid_path.exists():
                try:
                    pid = int(pid_path.read_text().strip())
                except ValueError:
                    pid = None
            runtime.pid = pid or _lsof_listen_pid(runtime.port) or runtime.pid
            runtime.status = RuntimeStatus.running
            runtime.health_checked_at = datetime.now(timezone.utc)
            audit(db, principal=p, action="runtime.healthcheck", resource_type="runtime",
                  resource_id=runtime.id, request=request, extra={"status": "ok", "auth": "ok", "pid": runtime.pid})
            db.commit()
            return {"ok": True, "running": True, "gateway": body, "auth": "ok", "pid": runtime.pid, "port": runtime.port}
    except Exception as ex:
        runtime.status = RuntimeStatus.error
        runtime.health_checked_at = datetime.now(timezone.utc)
        db.commit()
        return {"ok": False, "running": False, "error": str(ex), "port": runtime.port}


@app.patch("/api/admin/tenants/{tid}")
def admin_update_tenant(
    tid: str, body: TenantUpdateIn, request: Request,
    p: Principal = Depends(require_system_admin), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.4 — update tenant name/plan/resource-limits.

    Distinguishes 'field omitted' (no change) from 'field is null' (clear).
    Use `model_fields_set` to detect which were explicitly sent.
    """
    tenant = db.get(Tenant, tid)
    if not tenant:
        raise HTTPException(404, "tenant not found")
    changes = {}
    for field in ("name", "plan", "max_sessions", "max_employees"):
        if field in body.model_fields_set:
            v = getattr(body, field)
            setattr(tenant, field, v)
            changes[field] = v
    if "status" in body.model_fields_set and body.status is not None:
        if body.status not in {"active", "suspended", "archived"}:
            raise HTTPException(400, "invalid tenant status")
        tenant.status = TenantStatus(body.status)
        changes["status"] = body.status
    if not changes:
        raise HTTPException(400, "no fields to update")
    tenant.updated_at = datetime.utcnow()
    audit(db, principal=p, action="tenant.update", resource_type="tenant",
          resource_id=tid, request=request, extra=changes)
    db.commit()
    return _serialize_admin_tenant(db, tenant)


@app.delete("/api/admin/tenants/{tid}")
def admin_delete_tenant(
    tid: str, request: Request,
    p: Principal = Depends(require_system_admin), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.5 — delete a tenant + all rows referencing it.

    Kills any running gateway process for this tenant, removes the
    tenant home directory, and cascades user/employee/session/memory/
    skill rows. The system tenant `demo` is PROTECTED.
    """
    tenant = db.get(Tenant, tid)
    if not tenant:
        raise HTTPException(404, "tenant not found")
    if tenant.slug == "demo":
        raise HTTPException(400, "cannot delete the demo tenant")

    # 1) Stop the runtime if running
    runtime = (
        db.query(HermesRuntime)
        .filter(HermesRuntime.tenant_id == tid)
        .order_by(HermesRuntime.created_at.desc())
        .first()
    )
    if runtime and runtime.pid and _pid_alive(runtime.pid):
        try:
            from app.services.process_utils import terminate_pid
            terminate_pid(runtime.pid, grace=2.0)
        except ProcessLookupError:
            pass

    # 2) Cascade-delete referencing rows
    for model, col in [
        (User, "tenant_id"),
        (OrganizationUnit, "tenant_id"),
        (HermesRuntime, "tenant_id"),
        (DigitalEmployee, "tenant_id"),
        (SessionRecord, "tenant_id"),
        (AuditLog, "tenant_id"),
    ]:
        db.query(model).filter(getattr(model, col) == tid).delete(synchronize_session=False)
    # Skill/Memory rows that may have owner_tenant_id set
    from app.db.models import SkillPackage, MemoryEntry, MemoryBinding, SkillBinding
    db.query(SkillPackage).filter(SkillPackage.owner_tenant_id == tid).delete(synchronize_session=False)
    db.query(MemoryEntry).filter(MemoryEntry.tenant_id == tid).delete(synchronize_session=False)
    db.query(MemoryBinding).filter(MemoryBinding.tenant_id == tid).delete(synchronize_session=False)
    db.query(SkillBinding).filter(SkillBinding.tenant_id == tid).delete(synchronize_session=False)

    # 3) Audit BEFORE deleting tenant itself (audit has FK)
    audit(db, principal=p, action="tenant.delete", resource_type="tenant",
          resource_id=tid, request=request, extra={"slug": tenant.slug})
    db.commit()
    db.delete(tenant)
    db.commit()

    # 4) Remove tenant home directory from disk (best effort)
    import shutil
    home = OPENATLAS_HOME / "hermes-tenants" / tenant.slug
    try:
        if home.exists():
            shutil.rmtree(home, ignore_errors=True)
    except Exception:
        pass
    return {"ok": True, "deleted": tid, "slug": tenant.slug}


def _pid_alive(pid: int | None) -> bool:
    from app.services.process_utils import is_pid_alive
    return is_pid_alive(pid)


@app.get("/api/admin/supervisor")
def admin_supervisor_status(
    p: Principal = Depends(require_system_admin), db: Session = Depends(get_db),
) -> dict:
    """Phase 3.3 — show supervisor in-memory state + per-runtime restart history."""
    from app.services.supervisor import (
        _last_attempt, _consecutive_failures, _recent_restarts, _MAX_RESTART_PER_HOUR,
    )
    rows = db.query(HermesRuntime).all()
    out = []
    for r in rows:
        out.append({
            "tenant_id": r.tenant_id,
            "port": r.port,
            "status": r.status.value if hasattr(r.status, "value") else r.status,
            "pid": r.pid,
            "last_attempt_at": _last_attempt.get(r.id),
            "consecutive_failures": _consecutive_failures.get(r.id, 0),
            "restarts_last_hour": len(_recent_restarts.get(r.id, [])),
            "max_restarts_per_hour": _MAX_RESTART_PER_HOUR,
        })
    return {"items": out}


# ── Phase 2 helpers ─────────────────────────────────────────────────────────

def _next_hermes_port(db) -> int:
    """Allocate the next available port starting at 58643 (58642 reserved for demo default)."""
    used = {r.port for r in db.query(HermesRuntime).all()}
    p = 58643
    while p in used or _port_in_use(p):
        p += 1
        if p > 58999:
            raise RuntimeError("no available hermes port under 59000")
    return p


def _port_in_use(port: int) -> bool:
    import socket
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        return s.connect_ex(("127.0.0.1", port)) == 0


def _tenant_home(slug: str) -> Path:
    return OPENATLAS_HOME / "hermes-tenants" / slug / ".hermes"


def _tenant_pid_path(hermes_home: str) -> Path:
    return Path(hermes_home) / "openatlas-gateway.pid"


def _tenant_log_path(hermes_home: str) -> Path:
    return OPENATLAS_HOME / "logs" / f"hermes-tenant-{Path(hermes_home).parent.name}.log"


def _start_tenant_script() -> Path:
    return Path(OPENATLAS_HOME).parent / "Desktop" / "Atlasagent" / "openatlas" / "scripts" / "start_tenant.sh"


def _lsof_listen_pid(port: int) -> int | None:
    try:
        out = subprocess.run(
            ["lsof", "-nP", f"-iTCP:{port}", "-sTCP:LISTEN", "-F", "p"],
            capture_output=True, text=True, timeout=3,
        )
        for line in out.stdout.splitlines():
            if line.startswith("p"):
                return int(line[1:])
    except Exception:
        return None
    return None


def _runtime_auth_probe_sync(base_url: str, api_key: str, *, timeout: float = 5) -> tuple[bool, str]:
    """Verify the gateway accepts the runtime key on an authenticated endpoint."""
    import urllib.error
    import urllib.request

    url = f"{base_url.rstrip('/')}/v1/models"
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {api_key}"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return 200 <= resp.status < 500, f"HTTP {resp.status}"
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        if exc.code in (401, 403):
            return False, f"HTTP {exc.code}: {detail[:200]}"
        return False, f"HTTP {exc.code}: {detail[:200]}"
    except Exception as exc:
        return False, str(exc)
